"""The NDC (spec-driven statement) "entries in system" export.

What a user downloads to check an upload landed, so the cases worth guarding are: the sheet
carries the Open screen's columns in order, an amount becomes a number but an unparseable
one is kept rather than zeroed, the Billing words match the worklist, the Summary accounts
for what the type skips by design, and uploaded text never turns into a live formula.

No DB, no network — statement_export is pure.

Run:  python -m unittest discover -s tests      (from backend/)
"""

import os
import sys
import tempfile
import unittest
from decimal import Decimal

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from openpyxl import load_workbook  # noqa: E402

from app.services import statement_export as ex  # noqa: E402

COLUMNS = [
    {"header": "Document No", "field": "document_no"},
    {"header": "TXN Type", "field": "txn_type"},
    {"header": "Passenger Name", "field": "passenger_name"},
    {"header": "Total Fare", "field": "total_fare", "kind": "money"},
]
SUMMARY = [{"field": "total_fare", "label": "Total Fare"}]
ROW_FILTER = {"header": "TXN Type", "exclude": ["FREE_SEAT", "UNPAID_BOOKING"],
              "note": "Unpaid holds and free seats are not transactions."}


NDC = {"extra_headers": ("Line",), "with_billing": True}
API = {"extra_headers": ("Category", "Kind"), "with_billing": True}


def entry(id_, fare, billing="no_party", line="Ticket", **data):
    return ex.Entry(record={"id": id_, "document_no": f"098{id_:010d}", "txn_type": "PAID_BOOKING",
                            "total_fare": fare, **data}, billing=billing, extras={"Line": line})


def facts(**kw):
    return ex.UploadFacts(**{"type_label": "NDC", "source_file": "AirIndia.xlsx", "stored_rows": 3, **kw})


class ColumnTests(unittest.TestCase):

    def test_open_screen_columns_then_billing_then_id(self):
        base = ["Document No", "TXN Type", "Passenger Name", "Total Fare"]
        ndc = [c.header for c in ex.entry_columns(COLUMNS, **NDC)]
        self.assertEqual(ndc, base + ["Line", "Billing", "System ID"])
        api = [c.header for c in ex.entry_columns(COLUMNS, **API)]
        self.assertEqual(api, base + ["Category", "Kind", "Billing", "System ID"])
        # Billing not set up (or no billing flow): no billing columns at all.
        self.assertEqual([c.header for c in ex.entry_columns(COLUMNS, extra_headers=("Line",))],
                         base + ["System ID"])

    def test_amounts_become_numbers_but_unparseable_ones_are_kept(self):
        vals = ex.entry_values(entry(7, "1,23,456.50"), COLUMNS, **NDC)
        self.assertEqual(vals[3], Decimal("123456.50"))
        # "N/A" stays visible as the file's text — never a silent 0.
        self.assertEqual(ex.entry_values(entry(8, "N/A"), COLUMNS, **NDC)[3], "N/A")
        # A document number is text, so its leading zeros survive.
        self.assertEqual(vals[0], "0980000000007")

    def test_billing_uses_the_worklist_words_and_extras_line_up(self):
        e = ex.Entry(record={"id": 1, "total_fare": "10"}, billing="not_billable",
                     extras={"Category": "Hotel", "Kind": ex.KIND_LABEL["pending"]})
        self.assertEqual(ex.entry_values(e, COLUMNS, **API)[-4:],
                         ["Hotel", "Payment pending", "Not billable", 1])
        self.assertEqual(ex.LINE_LABEL["latched"], "Rolled up")


class RowCheckTests(unittest.TestCase):

    def test_ndc_explains_what_it_skips(self):
        msg = ex.row_check(facts(stored_rows=58, row_filter=ROW_FILTER))
        self.assertIn("58 entries are stored", msg)
        self.assertIn("FREE_SEAT, UNPAID_BOOKING", msg)

    def test_types_that_record_lines_read_compare_them(self):
        self.assertIn("All 40 lines", ex.row_check(facts(stored_rows=40, file_rows=40)))
        self.assertTrue(ex.row_check(facts(stored_rows=38, file_rows=40)).startswith("2 of the 40"))


class WorkbookTests(unittest.TestCase):

    def setUp(self):
        fd, self.path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)

    def tearDown(self):
        if os.path.exists(self.path):
            os.remove(self.path)

    def test_round_trip(self):
        entries = [
            entry(1, "100.50", billing="ready", passenger_name="=HYPERLINK(\"http://x\")"),
            entry(2, "-25.25", billing="no_party"),
            entry(3, "5", billing="latched", line="Rolled up"),
        ]
        ex.write_workbook(self.path, COLUMNS, entries,
                          facts(row_filter=ROW_FILTER, filters=[("TXN Type", "PAID_BOOKING")],
                                billing={"resolution_status": "resolved", "billable_rows": 2,
                                         "resolved_rows": 1, "projected_tickets": 0}),
                          SUMMARY, **NDC)

        wb = load_workbook(self.path)
        self.assertEqual(wb.sheetnames, [ex.ENTRIES_TITLE, ex.SUMMARY_TITLE])
        ws = wb[ex.ENTRIES_TITLE]
        self.assertEqual(ws.max_row, 4)                                  # header + 3 entries
        headers = [c.value for c in ws[1]]
        name_cell = ws.cell(row=2, column=headers.index("Passenger Name") + 1)
        self.assertEqual(name_cell.data_type, "s")                      # text, not a formula
        self.assertEqual(name_cell.value, "=HYPERLINK(\"http://x\")")
        self.assertEqual(ws.cell(row=2, column=headers.index("Total Fare") + 1).value, 100.5)

        flat = {}
        for r in wb[ex.SUMMARY_TITLE].iter_rows(values_only=True):
            vals = [v for v in r if v is not None]
            if len(vals) == 2:
                flat.setdefault(vals[0], vals[1])
        self.assertEqual(flat["Entries in this export (filtered)"], 3)
        self.assertEqual(flat["TXN Type"], "PAID_BOOKING")
        self.assertEqual(flat["Total Fare"], 80.25)
        self.assertEqual(flat["Billable lines"], 2)
        self.assertEqual(flat["Ready to send"], 1)
        self.assertEqual(flat["Needs a party"], 1)
        self.assertEqual(flat["On its ticket"], 1)
        # The Line breakdown sits beside the billing-state one.
        self.assertEqual(flat["Ticket"], 2)
        self.assertEqual(flat["Rolled up"], 1)

    def _summary_text(self) -> list[str]:
        return [str(r[0]) for r in load_workbook(self.path)[ex.SUMMARY_TITLE].iter_rows(values_only=True)
                if r and r[0]]

    def test_a_type_without_billing_has_no_billing_blocks(self):
        ex.write_workbook(self.path, COLUMNS, [entry(1, "10")], facts(stored_rows=1))
        self.assertFalse(any(t.startswith("Billing") for t in self._summary_text()))

    def test_billing_not_set_up_says_so_instead_of_labelling_every_row(self):
        # What the endpoint does for an upload still showing "Set up billing": no billing
        # columns, and the Summary says why.
        ex.write_workbook(self.path, COLUMNS, [ex.Entry(record={"id": 1, "total_fare": "10"})],
                          facts(stored_rows=1, billing={}), extra_headers=("Category", "Kind"))
        ws = load_workbook(self.path)[ex.ENTRIES_TITLE]
        self.assertNotIn("Billing", [c.value for c in ws[1]])
        self.assertTrue(any(t.startswith("Billing has not been set up") for t in self._summary_text()))
        self.assertFalse(facts(billing={"resolution_status": "none"}).billing_set_up)
        self.assertTrue(facts(billing={"resolution_status": "resolved"}).billing_set_up)


class FileNameTests(unittest.TestCase):

    def test_extension_dropped_and_unsafe_characters_replaced(self):
        self.assertEqual(ex.file_name("AirIndiaLccStatement.xlsx"),
                         "AirIndiaLccStatement - entries in system.xlsx")
        self.assertEqual(ex.file_name("a/b:c.TXT"), "a_b_c - entries in system.xlsx")
        self.assertEqual(ex.file_name(""), "upload - entries in system.xlsx")


if __name__ == "__main__":
    unittest.main()
