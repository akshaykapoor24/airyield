"""Customer and Corporate Billing's "Download Tickets XLS" — what each row of the sheet must say.

No DB. Everything asserted here is the pure row builder in services/party_ticket_export;
its `export_rows` loader is exercised against the live database instead. Three things
about the rows are not cosmetic:

  * a BILLED ticket shows its saved invoice line, not a recalculation — the sheet and the
    PDF must not disagree after the corporate's markup changes;
  * an UNBILLED ticket is priced the way Sold Tickets previews it;
  * every row adds up: BASE FARE + OTHERS + XXLN is the billed base, on a refund too.

Run:  python -m unittest discover backend/tests
"""

import os
import sys
import unittest
from datetime import date
from types import SimpleNamespace as N

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from app.services import party_ticket_export as pte  # noqa: E402
from app.services.party_ticket_export import (  # noqa: E402
    HEADERS, _flight_label, _route_points, build_workbook, corporate_name, customer_name,
)


def ticket_row(t, party, *, bill_to=None, **kw):
    """The service's row, with BILL TO defaulting to the party's name as a corporate router passes it."""
    return pte.ticket_row(t, party, bill_to=bill_to or corporate_name(party), **kw)


def ticket(**kw):
    base = dict(
        id=1, product_category="air", pax_count=1, pax_name=None, first_name=None, last_name=None,
        total_amt=None, sell_fare=None, booking_fee_sell=None, seat_selection=None,
        excess_baggage=None, meals=None, can_charge=None, comm_sell=None,
        calculated_incentive=None, tds_sell=None, ticket_date=None, travel_dt=None,
        departure_datetime=None, sector=None, flight_no=None, airlines_code=None,
        air_pnr=None, gds_pnr=None, booking_ref=None, segments=None, service_details=None,
        airline_name=None, booking_agency_name=None,
    )
    base.update(kw)
    return N(**base)


def corporate(**kw):
    base = dict(company="RELIOBRIX CONSULTING", first_name=None, last_name=None,
                markup_type="fixed", markup_value=300, category_markups=None, billing_type="agency")
    base.update(kw)
    return N(**base)


def row_dict(row):
    """Keyed by header; the two TOs and two TOTALs get a suffix so both survive."""
    seen, out = {}, {}
    for header, value in zip(HEADERS, row):
        n = seen.get(header, 0)
        seen[header] = n + 1
        out[header if n == 0 else f"{header} 2"] = value
    return out


class SampleRowTests(unittest.TestCase):
    """The two air rows of the sheet the business sent, reproduced end to end."""

    def test_row_with_others_cancellation_and_commission(self):
        t = ticket(pax_name="VIVEK AGARWAL", total_amt=12993, booking_fee_sell=700, can_charge=350,
                   comm_sell=3139, ticket_date="2026-09-05", travel_dt="2026-09-12",
                   sector="HYD-DEL", flight_no="711", airlines_code="6E", air_pnr="EG1W6U",
                   airline_name="INDIGO")
        r = row_dict(ticket_row(t, corporate(), line=None, bill_no=None, interstate=False))

        self.assertEqual(r["Type"], "AIR")
        self.assertIsNone(r["BILL NO"])
        self.assertEqual(r["BILL TO"], "RELIOBRIX CONSULTING")
        self.assertEqual(r["SUPPLIER"], "INDIGO")
        self.assertEqual(r["BOOKING DATE"], date(2026, 9, 5))
        self.assertEqual((r["FROM"], r["TO"], r["TO 2"]), ("HYD", "DEL", None))
        self.assertEqual(r["PNR"], "EG1W6U")
        self.assertEqual(r["FLIGHT NO / HOTEL NAME"], "6E711")
        self.assertEqual(r["DATE FROM"], date(2026, 9, 12))
        self.assertIsNone(r["DATE TO"])

        self.assertEqual(r["BASE FARE"], 11943)
        self.assertEqual(r["OTHERS"], 700)
        self.assertEqual(r["XXLN CHARGES"], 350)
        self.assertIsNone(r["ADDITIONAL MARKUP"])
        self.assertEqual(r["TOTAL"], 12993)
        self.assertEqual(r["SERVICE CHARGE"], 300)
        self.assertEqual(r["TOTAL 2"], 13293)
        self.assertEqual(r["GST CHARGED"], 54)          # agency: 18% of the markup alone
        self.assertEqual(r["TOTAL AMOUNT"], 13347)
        self.assertEqual(r["GROSS PURCHASE"], 12993)
        self.assertEqual(r["COM"], 3139)
        self.assertIsNone(r["TDS"])
        self.assertEqual(r["NET PURCHASE"], 9854)
        self.assertEqual(r["BILL AMOUNT"], 13347)

    def test_two_passengers_double_a_fixed_markup_and_mark_the_name(self):
        t = ticket(pax_name="SHISHIR VASUDEO DESAI", pax_count=2, total_amt=12956,
                   sector="BOM-BLR", flight_no="5071", airlines_code="6E")
        r = row_dict(ticket_row(t, corporate(), line=None, bill_no=None, interstate=True))
        self.assertEqual(r["NAME OF PASSANGER"], "SHISHIR VASUDEO DESAI^2")
        self.assertEqual(r["SERVICE CHARGE"], 600)
        self.assertEqual(r["GST CHARGED"], 108)
        self.assertEqual(r["BILL AMOUNT"], 13664)
        self.assertEqual(r["FLIGHT NO / HOTEL NAME"], "6E5071")


class PricingTests(unittest.TestCase):

    def test_reseller_is_taxed_on_gross_plus_markup(self):
        c = corporate(billing_type="reseller", markup_type="percentage", markup_value=5)
        r = row_dict(ticket_row(ticket(total_amt=1000), c, line=None, bill_no=None, interstate=None))
        self.assertEqual(r["SERVICE CHARGE"], 50)
        self.assertEqual(r["GST CHARGED"], 189)
        self.assertEqual(r["BILL AMOUNT"], 1239)

    def test_no_billing_type_charges_no_gst(self):
        c = corporate(billing_type=None)
        r = row_dict(ticket_row(ticket(total_amt=1000), c, line=None, bill_no=None, interstate=False))
        self.assertIsNone(r["GST CHARGED"])
        self.assertEqual(r["BILL AMOUNT"], 1300)

    def test_billed_ticket_reads_its_invoice_line_not_todays_markup(self):
        line = {"ticket_id": 1, "base_amount": 1000, "markup_amount": 300, "additional_markup": 100,
                "discount": 50, "gst_amount": 45, "total": 1395, "pax_count": 1}
        # The corporate has since moved to a different markup; the bill must not.
        c = corporate(markup_value=999)
        r = row_dict(ticket_row(ticket(total_amt=1000), c, line=line, bill_no="GT/26-27/0027",
                                interstate=False))
        self.assertEqual(r["BILL NO"], "GT/26-27/0027")
        self.assertEqual(r["ADDITIONAL MARKUP"], 100)
        self.assertEqual(r["TOTAL"], 1100)
        self.assertEqual(r["SERVICE CHARGE"], 300)
        self.assertEqual(r["TOTAL 2"], 1400)
        self.assertEqual(r["GST CHARGED"], 45)
        self.assertEqual(r["TOTAL AMOUNT"], 1445)
        self.assertEqual(r["BILL AMOUNT"], 1395)        # net of the bill's discount

    def test_bill_to_names_the_party_the_bill_was_raised_to(self):
        line = {"ticket_id": 1, "base_amount": 1000, "markup_amount": 50, "additional_markup": 0,
                "discount": 0, "gst_amount": 0, "total": 1050}
        r = row_dict(ticket_row(ticket(total_amt=1000), corporate(company="fareqube"), line=line,
                                bill_no="GT/26-27/0019", interstate=False, bill_to="Jatin lk wasnik"))
        self.assertEqual(r["BILL TO"], "Jatin lk wasnik")
        unbilled = row_dict(ticket_row(ticket(total_amt=1000), corporate(company="fareqube"), line=None,
                                       bill_no=None, interstate=False))
        self.assertEqual(unbilled["BILL TO"], "fareqube")

    def test_a_customer_is_billed_by_name_and_a_corporate_by_company(self):
        # The invoice's own rule (billing_pdf._bill_to_lines): a person, not their employer.
        employee = N(first_name="UMESH", last_name="SHARMA", company="fareqube")
        self.assertEqual(customer_name(employee), "UMESH SHARMA")
        self.assertEqual(corporate_name(N(first_name=None, last_name=None, company="fareqube")), "fareqube")
        self.assertEqual(customer_name(N(first_name=None, last_name=None, company="Acme")), "Acme")

    def test_customer_markup_prices_an_unbilled_customer_line(self):
        person = N(first_name="Jatin", last_name=None, company="fareqube", markup_type="percentage",
                   markup_value=10, category_markups=None, billing_type="reseller")
        r = row_dict(pte.ticket_row(ticket(total_amt=1000), person, line=None, bill_no=None,
                                    interstate=False, bill_to=customer_name(person)))
        self.assertEqual(r["BILL TO"], "Jatin")
        self.assertEqual(r["SERVICE CHARGE"], 100)
        self.assertEqual(r["GST CHARGED"], 198)
        self.assertEqual(r["BILL AMOUNT"], 1298)

    def test_refund_with_a_cancellation_charge_still_adds_up(self):
        t = ticket(total_amt=-3500, can_charge=1500)
        r = row_dict(ticket_row(t, corporate(), line=None, bill_no=None, interstate=False))
        self.assertEqual(r["BASE FARE"], -5000)
        self.assertEqual(r["XXLN CHARGES"], 1500)
        self.assertEqual(r["TOTAL"], -3500)
        self.assertEqual(r["BASE FARE"] + r["XXLN CHARGES"], r["GROSS PURCHASE"])
        self.assertEqual(r["SERVICE CHARGE"], -300)
        self.assertEqual(r["BILL AMOUNT"], -3854)

    def test_commission_falls_back_to_the_deal_engine_and_tds_raises_net(self):
        t = ticket(total_amt=10000, calculated_incentive=500, tds_sell=25)
        r = row_dict(ticket_row(t, corporate(), line=None, bill_no=None, interstate=False))
        self.assertEqual(r["COM"], 500)
        self.assertEqual(r["TDS"], 25)
        self.assertEqual(r["NET PURCHASE"], 9525)


class ItineraryTests(unittest.TestCase):

    def test_route_points(self):
        self.assertEqual(_route_points("COK-BLR / BLR-CCU"), ["COK", "BLR", "CCU"])
        self.assertEqual(_route_points("DEL-STV / STV-DEL"), ["DEL", "STV", "DEL"])
        self.assertEqual(_route_points("DEL-BOM / GOI-DEL"), ["DEL", "BOM", "GOI", "DEL"])
        self.assertEqual(_route_points("MOPA Airport/Hyderabad"), ["MOPA Airport", "Hyderabad"])
        self.assertEqual(_route_points(None), [])

    def test_open_jaw_keeps_every_point_in_the_last_to(self):
        t = ticket(total_amt=1, sector="DEL-BOM / GOI-DEL")
        r = row_dict(ticket_row(t, corporate(), line=None, bill_no=None, interstate=False))
        self.assertEqual((r["FROM"], r["TO"], r["TO 2"]), ("DEL", "BOM", "GOI-DEL"))

    def test_flight_label(self):
        self.assertEqual(_flight_label("6398 / 6453", "6E"), "6E6398 / 6E6453")
        self.assertEqual(_flight_label("AI423", "AI"), "AI423")
        self.assertEqual(_flight_label("S5-211-S5-213", None), "S5-211-S5-213")
        self.assertIsNone(_flight_label("", "6E"))

    def test_round_trip_shows_the_return_date(self):
        t = ticket(total_amt=1, sector="DEL-STV / STV-DEL", travel_dt="2026-08-19",
                   segments=[{"route": "DEL-STV", "dep_date": "19-Aug-2026 06:10:00 AM"},
                             {"route": "STV-DEL", "dep_date": "20-Aug-2026 05:40:00 PM"}])
        r = row_dict(ticket_row(t, corporate(), line=None, bill_no=None, interstate=False))
        self.assertEqual(r["DATE FROM"], date(2026, 8, 19))
        self.assertEqual(r["DATE TO"], date(2026, 8, 20))

    def test_same_day_connection_has_no_date_to(self):
        t = ticket(total_amt=1, sector="COK-BLR / BLR-CCU", travel_dt="2026-09-27",
                   segments=[{"dep_date": "27-Sep-2026 06:05:00 PM"}, {"dep_date": "27-Sep-2026 08:20:00 PM"}])
        r = row_dict(ticket_row(t, corporate(), line=None, bill_no=None, interstate=False))
        self.assertIsNone(r["DATE TO"])

    def test_hotel_line(self):
        t = ticket(product_category="hotel", pax_name="ANKUR SHUKLA", total_amt=1545,
                   booking_ref="NH24229512007220", booking_agency_name="MakeMyTrip",
                   travel_dt="2026-08-26",
                   service_details={"category": "hotel", "label": "Hotel", "title": "Hotel Deviram Palace",
                                    "route": "AGRA", "check_in": "2026-08-26", "check_out": "2026-08-28"})
        r = row_dict(ticket_row(t, corporate(), line=None, bill_no=None, interstate=False))
        self.assertEqual(r["Type"], "HOTEL")
        self.assertEqual(r["SUPPLIER"], "MakeMyTrip")
        self.assertEqual((r["FROM"], r["TO"]), ("AGRA", None))
        self.assertEqual(r["PNR"], "NH24229512007220")
        self.assertEqual(r["FLIGHT NO / HOTEL NAME"], "Hotel Deviram Palace")
        self.assertEqual(r["DATE FROM"], date(2026, 8, 26))
        self.assertEqual(r["DATE TO"], date(2026, 8, 28))

    def test_train_line_uses_the_irctc_pnr_and_route(self):
        t = ticket(product_category="train", total_amt=2000, booking_ref="TBOB123",
                   service_details={"label": "Train", "title": "22221 NZM RAJDHANI", "route": "CSMT → NZM",
                                    "journey_date": "2026-08-02", "pnr": "4512345678"})
        r = row_dict(ticket_row(t, corporate(), line=None, bill_no=None, interstate=False))
        self.assertEqual((r["FROM"], r["TO"]), ("CSMT", "NZM"))
        self.assertEqual(r["PNR"], "4512345678")
        self.assertEqual(r["FLIGHT NO / HOTEL NAME"], "22221 NZM RAJDHANI")
        self.assertEqual(r["DATE FROM"], date(2026, 8, 2))


class WorkbookTests(unittest.TestCase):

    def test_sheet_has_the_header_one_row_per_ticket_and_real_dates(self):
        from io import BytesIO  # noqa: F401
        from openpyxl import load_workbook

        t = ticket(total_amt=12993, ticket_date="2026-09-05", sector="HYD-DEL")
        rows = [ticket_row(t, corporate(), line=None, bill_no=None, interstate=False)] * 3
        ws = load_workbook(build_workbook(rows)).active
        self.assertEqual([c.value for c in ws[1]], HEADERS)
        self.assertEqual(ws.max_row, 4)
        self.assertEqual(ws["E2"].value.date(), date(2026, 9, 5))
        self.assertEqual(ws["E2"].number_format, "DD-MMM-YY")
        self.assertEqual(ws["AA2"].value, 13347)
        self.assertEqual(ws.freeze_panes, "A2")


if __name__ == "__main__":
    unittest.main()
