"""Deal-sheet PDF parsing: merged cells, the deterministic fast path, bare airline codes.

Two layers:
  * Pure tests (always run) pin each rule on hand-built rows and fake tables.
  * Sample-sheet tests read the real supplier PDFs (Lords, Akbar, Mirror PCC, Vicky)
    and pin what each must parse to. The PDFs are supplier documents and are NOT in
    the repo; point DEAL_SHEET_DIR at a folder holding them, or they are looked for in
    ~/Downloads and ~/OneDrive/Desktop/FareQube Docs/TriumphTravel. Missing → skipped.

Run:  python -m unittest discover -s tests      (from backend/)
"""

import os
import sys
import unittest
from datetime import date
from types import SimpleNamespace

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from app.services.ai_deal_extraction import (  # noqa: E402
    AIDealExtractionService, Stats, _simple_cell, deterministic_compact, expand_row,
    parse_period_window, parse_validity_window,
)
from app.services.airline_resolver import AirlineIndex, bare_designator, RESOLVED  # noqa: E402
from app.services.pdf_table_rows import (  # noqa: E402
    ColumnMap, DocRow, _is_repeat_header, build_colmap, fill_merged_cells, is_noise_row,
)


def row(n, cells, cm, section=None):
    from app.services.pdf_table_rows import _row_expects_deals
    return DocRow(n=n, page=0, cells=tuple(cells), section=section,
                  raw="\t".join(cells), expects_deals=_row_expects_deals(tuple(cells), cm))


# Vicky's column order: rates run BUS, P.ECO, ECO — the reverse of most sheets.
VICKY_HEADER = ("Airline", "BUS/FIRST", "P.ECO", "ECO", "Validity", "Remarks")


class HeaderTest(unittest.TestCase):
    def test_bus_first_is_business(self):
        cm = build_colmap(VICKY_HEADER)
        self.assertEqual((cm.bus, cm.pey, cm.eco, cm.validity), (1, 2, 3, 4))

    def test_partial_header_repeat(self):
        header = ("Airline Name", "Airline Code", "Cabin Type", "Commission")
        self.assertTrue(_is_repeat_header(("Airline Name", "Airline Code", "", "Commission"), header))
        self.assertFalse(_is_repeat_header(("Air India", "AI", "Eco", "1%"), header))

    def test_repeated_banner_is_noise(self):
        cm = build_colmap(VICKY_HEADER)
        self.assertTrue(is_noise_row(("", "DEAL", "DEAL", "DEAL", "", ""), cm))


def fake_table(grid):
    """grid: rows of (text, bbox|None). The first row has every cell (the grid row)."""
    rows = []
    for r in grid:
        boxes = [b for _, b in r]
        present = [b for b in boxes if b is not None]
        y0 = min(b[1] for b in present)
        y1 = max(b[3] for b in present)
        rows.append(SimpleNamespace(cells=boxes, bbox=(0, y0, 100, y1)))
    return SimpleNamespace(rows=rows), [[t for t, _ in r] for r in grid]


class MergeTest(unittest.TestCase):
    # columns at x: 0-10 airline, 10-20 BUS, 20-30 PEY, 30-40 ECO, 40-50 validity
    def test_across_rate_columns_and_down(self):
        table, texts = fake_table([
            [("Airline", (0, 0, 10, 10)), ("BUS", (10, 0, 20, 10)), ("PEY", (20, 0, 30, 10)),
             ("ECO", (30, 0, 40, 10)), ("Validity", (40, 0, 50, 10))],
            # VS: 2.75% spans BUS+PEY, 1.75% under ECO
            [("VS", (0, 10, 10, 20)), ("2.75%", (10, 10, 30, 20)), (None, None),
             ("1.75%", (30, 10, 40, 20)), ("Dec'26", (40, 10, 50, 20))],
            # AZ: name and validity span two rows; each row's rate spans all classes
            [("AZ", (0, 20, 10, 40)), ("2.75%", (10, 20, 40, 30)), (None, None), (None, None),
             ("Dec'26", (40, 20, 50, 40))],
            [(None, None), ("2%", (10, 30, 40, 40)), (None, None), (None, None), (None, None)],
        ])
        fill_merged_cells(table, texts, rate_cols={1, 2, 3})
        self.assertEqual(texts[1][1:4], ["2.75%", "2.75%", "1.75%"])
        self.assertEqual(texts[2][1:4], ["2.75%"] * 3)
        self.assertEqual(texts[3], ["AZ", "2%", "2%", "2%", "Dec'26"])

    def test_banner_never_spreads_across(self):
        table, texts = fake_table([
            [("A", (0, 0, 10, 10)), ("B", (10, 0, 20, 10)), ("C", (20, 0, 30, 10))],
            [("BANNER", (0, 10, 30, 20)), (None, None), (None, None)],
        ])
        fill_merged_cells(table, texts, rate_cols={1, 2})
        self.assertEqual(texts[1], ["BANNER", None, None])


class FastPathTest(unittest.TestCase):
    def test_simple_cells(self):
        self.assertEqual(_simple_cell("2.75%(B+YQ)"), (True, {"p": 2.75, "cc": "B+YQ"}))
        self.assertEqual(_simple_cell("2% B+YR+YQ")[1]["cc"], "B+YQ+YR")
        self.assertEqual(_simple_cell("1.25% Basic")[1], {"p": 1.25, "cc": "B"})
        self.assertEqual(_simple_cell("0.0%")[1]["p"], 0.0)
        self.assertEqual(_simple_cell("NIL"), (True, None))
        for messy in ("CNP+2%", "3+5%", "3% IATA", "NET + SC 50", "1.25%-STD-(B+YQ), 1.5%"):
            self.assertFalse(_simple_cell(messy)[0], messy)

    def test_vicky_row_needs_no_model(self):
        cm = build_colmap(VICKY_HEADER)
        r = row(9, ["MH", "10.5%", "8.5% (B)", "8.5% (B)", "FLOWN TILL 31ST Dec'26", ""], cm)
        compact = deterministic_compact(r, cm)
        self.assertIsNotNone(compact)
        deals = expand_row(compact, r, Stats(), cm)
        got = {d["incentive_data"]["PLB"]["class"]: d["incentive_data"]["PLB"]["incentiveAmtPct"] for d in deals}
        self.assertEqual(got, {"Economy": 8.5, "Premium": 8.5, "Business": 10.5})
        self.assertEqual(deals[0]["contract_valid_to"], "2026-12-31")
        # No period column: the validity cell is both windows' end, and travel's start is
        # left to the review page (the contract start, once Valid From is known).
        plb = deals[0]["incentive_data"]["PLB"]
        self.assertEqual((plb["travelValidFrom"], plb["travelValidTo"]), (None, "2026-12-31"))
        self.assertEqual(deals[0]["airline_code"], "MH")       # the name cell IS the code

    def test_messy_row_goes_to_model(self):
        cm = build_colmap(VICKY_HEADER)
        r = row(4, ["BA", "CNP+2.5% (B+YQ)", "CNP+2.5% (B+YQ)", "CNP+2% (B+YQ)", "Dec'26", ""], cm)
        self.assertIsNone(deterministic_compact(r, cm))

    def test_cabin_format(self):
        cm = build_colmap(("Airline Name", "Airline Code", "Cabin Type", "Commission",
                           "Applicable on", "Travel Validity", "Remarks"))
        r = row(1, ["Air Astana", "KC", "All", "1.18%", "Base Fare + YQ", "31 Dec'2026", ""], cm)
        deals = expand_row(deterministic_compact(r, cm), r, Stats(), cm)
        self.assertEqual(len(deals), 3)
        self.assertEqual(deals[0]["incentive_data"]["PLB"]["payoutCalcCols"], "Basic + YQ")


class BareCodeTest(unittest.TestCase):
    def setUp(self):
        self.idx = AirlineIndex.from_rows([
            (1, "VIRGIN ATLANTIC", "VS"), (2, "AIR FRANCE", "AF"), (3, "KLM ROYAL DUTCH AIRLINES", "KL"),
            (4, "DELTA AIRLINES", "DL"), (5, "MYANMAR AIRWAYS INTERNATIONAL", "8M"),
        ])

    def test_bare_designator(self):
        self.assertEqual(bare_designator("VS"), "VS")
        for not_code in ("Vs", "Virgin Atlantic", "AF/KLM/DL", "123"):
            self.assertIsNone(bare_designator(not_code), not_code)

    def test_code_in_name_cell_resolves(self):
        m = self.idx.resolve("VS", None)
        self.assertEqual((m.status, m.canonical_name), (RESOLVED, "VIRGIN ATLANTIC"))

    def test_multi_carrier_is_not_near_matched(self):
        self.assertNotEqual(self.idx.resolve("AF/KLM/DL", None).canonical_name, "KLM ROYAL DUTCH AIRLINES")
        self.assertEqual(len(self.idx.split_carriers("AF/KLM/DL", None)), 3)


class SpreadsheetTest(unittest.TestCase):
    """A vendor-comparison workbook: %-formatted rates, a Source column, a TYPE column."""

    def workbook(self):
        import io
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Week", "A/L", "TYPE", "Code", "Source", "IATA", "PLB(ECO)", "PLB(PRE ECO)",
                   "PLB(BUS)", "Applicable", "Travel Validity", "Remarks"])
        ws.append(["23-30SEP", "AC", None, 14, "AKBAR", 0, 0.0125, 0.0175, 0.02, "B+YQ", "Till 31DEC'26", "x"])
        ws.append(["23-30SEP", "AI", "INTL", 98, "AKBAR", 0, 0.01, 0.015, 0.0275, "B+YQ", "TKTG Till 31DEC'26", ""])
        ws.append(["23-30SEP", "AC", None, 14, "BRIGHTSUN", 0, 0.0325, 0.0325, 0.035, "B+YQ", "Till 31DEC'26", ""])
        for row in ws.iter_rows(min_row=2):
            for c in row[5:9]:
                c.number_format = "0.00%"
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_week_is_the_contract_and_validity_the_travel_window(self):
        from app.services.pdf_table_rows import build_layout_from_sheet
        L = build_layout_from_sheet(self.workbook(), "vendor.xlsx")
        self.assertEqual(L.colmap.period, 0)
        r = L.rows[0]                                         # 23-30SEP · Till 31DEC'26
        deals = expand_row(deterministic_compact(r, L.colmap), r, Stats(), L.colmap)
        week_from, week_to = parse_period_window("23-30SEP")
        plb = deals[0]["incentive_data"]["PLB"]
        self.assertEqual((deals[0]["contract_valid_from"], deals[0]["contract_valid_to"]), (week_from, week_to))
        self.assertEqual((plb["validFrom"], plb["validTo"]), (week_from, week_to))
        self.assertEqual((plb["travelValidFrom"], plb["travelValidTo"]), (week_from, "2026-12-31"))

    def test_reads_rates_and_keeps_one_supplier(self):
        from app.services.deal_supplier_check import supplier_cell_matches
        from app.services.pdf_table_rows import build_layout_from_sheet
        L = build_layout_from_sheet(
            self.workbook(), "vendor.xlsx",
            lambda cell: supplier_cell_matches(cell, "Akbar Travels of India Private Limited"))
        self.assertEqual(len(L.rows), 2)
        self.assertEqual(L.skipped_suppliers, {"BRIGHTSUN": 1})
        self.assertIsNone(L.colmap.code)                       # "Code" is the numeric stock code
        ac, ai = L.rows
        self.assertEqual(ac.cells[L.colmap.eco], "1.25%")       # 0.0125 with a % format
        self.assertEqual(ai.cells[L.colmap.airline], "AI (INTL)")
        deals = expand_row(deterministic_compact(ai, L.colmap), ai, Stats(), L.colmap)
        self.assertEqual([d["incentive_data"]["PLB"]["incentiveAmtPct"] for d in deals], [1.0, 1.5, 2.75])
        self.assertEqual(deals[0]["incentive_data"]["PLB"]["flightType"], "International")
        self.assertEqual(deals[0]["airline_code"], "AI")

    def test_no_filter_keeps_everyone(self):
        from app.services.pdf_table_rows import build_layout_from_sheet
        self.assertEqual(len(build_layout_from_sheet(self.workbook(), "vendor.xlsx").rows), 3)

    def test_deals_carry_their_source_cell(self):
        from app.services.pdf_table_rows import build_layout_from_sheet
        L = build_layout_from_sheet(self.workbook(), "vendor.xlsx")
        got = [d["supplier_cell"] for r in L.rows
               for d in expand_row(deterministic_compact(r, L.colmap), r, Stats(), L.colmap)]
        self.assertEqual(got, ["AKBAR"] * 6 + ["BRIGHTSUN"] * 3)

    def test_pdf_style_sheet_without_source_column_carries_none(self):
        cm = build_colmap(VICKY_HEADER)
        r = row(9, ["MH", "10.5%", "8.5% (B)", "8.5% (B)", "Dec'26", ""], cm)
        deals = expand_row(deterministic_compact(r, cm), r, Stats(), cm)
        self.assertTrue(all(d["supplier_cell"] is None for d in deals))


def _xlsx(sheets):
    """{title: [rows]} -> .xlsx bytes. A cell given as ("", fill) is formatted but empty."""
    import io
    import openpyxl
    from openpyxl.styles import PatternFill
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for ri, values in enumerate(rows, start=1):
            for ci, v in enumerate(values, start=1):
                if isinstance(v, tuple):
                    ws.cell(row=ri, column=ci).fill = PatternFill("solid", fgColor="FFA500")
                elif v is not None:
                    ws.cell(row=ri, column=ci, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


VENDOR_HEADER = ["A/L", "Source", "IATA", "PLB(ECO)", "PLB(PRE ECO)", "PLB(BUS)",
                 "Applicable", "Travel Validity", "Remarks"]


class VendorWorkbookTest(unittest.TestCase):
    """The weekly multi-supplier workbook: a GDS tab, an LCC tab, a Service Charge tab."""

    def book(self, header=VENDOR_HEADER, gds_extra=()):
        orange = ("", "fill")
        return _xlsx({
            # The header row's fill runs two columns past the table, as Excel sheets do.
            "GDS": [header + [orange, orange],
                    ["AA", "AKBAR", "0%", "2%", "2%", "2%", "B+YR", "Till 31DEC'26", "x"],
                    ["AF", "BRIGHTSUN", "0%", "2.25%", "2.75%", "3%", "B+YR", "01JAN'27 Onwards", ""],
                    ["EK", "TSC", "0%", "1350", "2700", "2700", "FLAT", "Till 15DEC'26", "flat"],
                    *gds_extra],
            "LCC": [VENDOR_HEADER, ["6E", "RIYA", "0%", "1%", "", "1%", "B", "Till 31DEC'26", ""]],
            "Service Charge": [["A/L", "Source", "Charge"], ["AA", "AKBAR", "250"]],
        })

    def test_formatting_past_the_table_does_not_hide_its_header(self):
        from app.services.pdf_table_rows import build_layout_from_sheet
        L = build_layout_from_sheet(self.book(), "vendor.xlsx")
        self.assertEqual((L.sheet_title, len(L.rows), L.header), ("GDS", 3, tuple(VENDOR_HEADER)))

    def test_other_deal_tabs_are_named(self):
        from app.services.pdf_table_rows import build_layout_from_sheet
        self.assertEqual(build_layout_from_sheet(self.book(), "vendor.xlsx").unread_deal_sheets, ["LCC"])

    def _extract(self, content, **kw):
        import asyncio
        import io
        from fastapi import UploadFile
        return asyncio.run(AIDealExtractionService.extract(
            UploadFile(file=io.BytesIO(content), filename="vendor.xlsx"), **kw))

    def test_all_suppliers_keeps_every_row_without_a_model_call(self):
        res = self._extract(self.book(), all_suppliers=True)
        by_src = {}
        for d in res["deals"]:
            by_src.setdefault(d["supplier_cell"], []).append(d)
        self.assertEqual(sorted(by_src), ["AKBAR", "BRIGHTSUN", "TSC"])
        self.assertEqual(len(by_src["AKBAR"]), 3)
        # EK's FLAT amounts are no percentage: one fill-in row, still TSC's.
        self.assertEqual([d["needs_manual"] for d in by_src["TSC"]], [True])
        self.assertIn("'LCC' tab also has a deal table", res["warning"])

    def test_picked_supplier_still_keeps_only_its_rows(self):
        res = self._extract(self.book(), supplier_name="Akbar Travels of India Pvt Ltd")
        self.assertEqual({d["supplier_cell"] for d in res["deals"]}, {"AKBAR"})
        self.assertIn("left out", res["warning"])

    def test_all_suppliers_needs_a_source_column(self):
        header = [h for h in VENDOR_HEADER if h != "Source"]
        content = _xlsx({"GDS": [header, ["AA", "0%", "2%", "2%", "2%", "B+YR", "Till 31DEC'26", ""]]})
        res = self._extract(content, all_suppliers=True)
        self.assertEqual(res["deals"], [])
        self.assertIn("Source / Supplier column", res["warning"])

    def test_onwards_opens_the_deal(self):
        res = self._extract(self.book(), all_suppliers=True)
        af = [d for d in res["deals"] if d["supplier_cell"] == "BRIGHTSUN"]
        self.assertEqual({(d["contract_valid_from"], d["contract_valid_to"]) for d in af}, {("2027-01-01", None)})
        self.assertEqual(af[0]["incentive_data"]["PLB"]["validFrom"], "2027-01-01")
        plb = af[0]["incentive_data"]["PLB"]
        self.assertEqual((plb["travelValidFrom"], plb["travelValidTo"]), ("2027-01-01", None))

    def test_onwards_on_a_week_sheet_opens_travel_not_the_contract(self):
        header = ("Week", "A/L", "Source", "IATA", "PLB(ECO)", "PLB(PRE ECO)", "PLB(BUS)",
                  "Applicable", "Travel Validity", "Remarks")
        cm = build_colmap(header)
        r = row(1, ["01-07OCT'26", "AF", "BRIGHTSUN", "0%", "2.25%", "2.75%", "3%", "B+YR",
                    "01JAN'27 Onwards", ""], cm)
        deals = expand_row(deterministic_compact(r, cm), r, Stats(), cm)
        plb = deals[0]["incentive_data"]["PLB"]
        self.assertEqual((deals[0]["contract_valid_from"], deals[0]["contract_valid_to"]),
                         ("2026-10-01", "2026-10-07"))
        self.assertEqual((plb["travelValidFrom"], plb["travelValidTo"]), ("2027-01-01", None))


class PeriodWindowTest(unittest.TestCase):
    TODAY = date(2026, 10, 25)

    def test_week_cells(self):
        for cell, want in (
            ("01-07OCT", ("2026-10-01", "2026-10-07")),
            ("01-07OCT'26", ("2026-10-01", "2026-10-07")),
            ("23-30SEP26", ("2026-09-23", "2026-09-30")),
            ("28SEP-04OCT", ("2026-09-28", "2026-10-04")),
            ("01OCT-07OCT", ("2026-10-01", "2026-10-07")),
            ("01 Oct 2026 - 07 Oct 2026", ("2026-10-01", "2026-10-07")),
            ("01/10/2026 - 07/10/2026", ("2026-10-01", "2026-10-07")),
        ):
            with self.subTest(cell):
                self.assertEqual(parse_period_window(cell, self.TODAY), want)

    def test_missing_year_is_the_nearest_one(self):
        self.assertEqual(parse_period_window("29DEC-04JAN", date(2027, 1, 2)), ("2026-12-29", "2027-01-04"))
        self.assertEqual(parse_period_window("29DEC-04JAN", date(2026, 12, 20)), ("2026-12-29", "2027-01-04"))

    def test_not_a_period(self):
        for cell in (None, "", "Till 31DEC'26", "W.E.F 07OCT'26", "01JAN'27 Onwards", "Week 40"):
            with self.subTest(cell):
                self.assertEqual(parse_period_window(cell, self.TODAY), (None, None))


class ValidityWindowTest(unittest.TestCase):
    def test_till_is_an_end(self):
        self.assertEqual(parse_validity_window("TKTG Till 31DEC'26"), (None, "2026-12-31"))
        self.assertEqual(parse_validity_window("O/B Till 31DEC'26"), (None, "2026-12-31"))
        self.assertEqual(parse_validity_window("Till Futher Notice"), (None, None))

    def test_onwards_is_a_start(self):
        self.assertEqual(parse_validity_window("01JAN'27 Onwards"), ("2027-01-01", None))
        self.assertEqual(parse_validity_window("Travel 01JAN'27 onwards till 31MAR'27"),
                         ("2027-01-01", "2027-03-31"))
        self.assertEqual(parse_validity_window(None), (None, None))


# ── the real sheets ──────────────────────────────────────────────────────────

def _find(name):
    home = os.path.expanduser("~")
    for d in (os.environ.get("DEAL_SHEET_DIR"), os.path.join(home, "Downloads"),
              os.path.join(home, "OneDrive", "Desktop", "FareQube Docs", "TriumphTravel")):
        if d and os.path.isfile(os.path.join(d, name)):
            return os.path.join(d, name)
    return None


SHEETS = {
    # name: (rows, deal rows, rows the fast path reads, deals it yields)
    "LORDS DEAL 18 MAR.pdf": (115, 111, 106, 296),
    "AKBAR TRAVELS COSOLIDATOR PLB SHEET EFF 7 March 2026.pdf": (39, 34, 29, 84),
    "Mirror PCC Deal Sheet Effective March 24th.pdf": (80, 79, 79, 191),
    "Deal Sheet Effective 03Oct26.pdf": (60, 60, 58, 174),
}


class SampleSheetTest(unittest.TestCase):
    def _layout(self, name):
        path = _find(name)
        if not path:
            self.skipTest(f"{name} not available")
        from app.services.pdf_table_rows import build_layout
        with open(path, "rb") as f:
            return build_layout(f.read())

    def test_counts(self):
        for name, (rows, expects, local, deals) in SHEETS.items():
            with self.subTest(name):
                path = _find(name)
                if not path:
                    continue
                L = self._layout(name)
                got_local, got_deals = 0, 0
                for r in L.rows:
                    c = deterministic_compact(r, L.colmap)
                    if c is not None:
                        got_local += 1
                        got_deals += len(expand_row(c, r, Stats(), L.colmap))
                self.assertEqual(
                    (len(L.rows), sum(r.expects_deals for r in L.rows), got_local, got_deals),
                    (rows, expects, local, deals),
                )

    def test_vicky_merged_rates(self):
        L = self._layout("Deal Sheet Effective 03Oct26.pdf")
        cm = L.colmap
        by_name = {}
        for r in L.rows:
            by_name.setdefault(r.cells[cm.airline], []).append(r)

        def rates(name, i=0):
            r = by_name[name][i]
            return r.cells[cm.bus], r.cells[cm.pey], r.cells[cm.eco]

        self.assertEqual(rates("VS"), ("2.75%(B+YQ)", "2.75%(B+YQ)", "1.75%(B+YQ)"))
        self.assertEqual(rates("MH"), ("10.5%", "8.5% (B)", "8.5% (B)"))
        self.assertEqual(rates("SK"), ("2% (B+YQ)", "2% (B+YQ)", "0.0%"))
        self.assertEqual(rates("AZ", 1), ("2% (B+YQ+YR)",) * 3)    # the continuation row
        self.assertEqual(by_name["AZ"][1].cells[cm.validity], "FLOWN TILL 31ST Dec'26")


if __name__ == "__main__":
    unittest.main()
