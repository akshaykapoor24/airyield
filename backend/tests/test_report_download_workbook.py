"""Report download workbook writer — the file must open, read as text where it should, and split.

No DB. Every test writes a real .xlsx through ReportWorkbook (openpyxl write-only) and
reads it back with openpyxl's normal loader, because the failure modes that matter only
show up in the saved package:

  * sheets created lazily and in any order come out by rank (Summary written last, second);
  * header row, frozen panes and auto filter on every data sheet and split part;
  * XML-illegal characters (C0 controls, lone surrogates, U+FFFE) are stripped so the
    package still loads;
  * uploaded text such as "=HYPERLINK(…)" or "#N/A" is stored as text, never as a formula
    or error cell, and "-285.72" text is not turned into a number;
  * a sheet past max_data_rows continues on "<title> (2)" with the header repeated;
  * money / date / datetime number formats; value conversions; discard() cleans up.

Run:  python -m unittest test_report_download_workbook   (from backend/tests)
"""

import os
import shutil
import sys
import tempfile
import unittest
import zipfile
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from unittest import mock

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from openpyxl import load_workbook  # noqa: E402

from app.services.report_download import workbook as wbmod  # noqa: E402
from app.services.report_download.columns import COMBINED_COLUMNS, COMBINED_HEADERS  # noqa: E402
from app.services.report_download.workbook import (  # noqa: E402
    MAX_CELL_CHARS, Block, ColSpec, ReportWorkbook, excel_value,
)

COMBINED_SPECS = [ColSpec(c.header, c.kind, c.width) for c in COMBINED_COLUMNS]
TWO_COLS = [ColSpec("Ref"), ColSpec("Amount", "money", 12)]


class _TempDirCase(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="report-wb-test-")
        self.addCleanup(shutil.rmtree, self.dir, ignore_errors=True)

    def path(self, name="report.xlsx"):
        return os.path.join(self.dir, name)

    def save_and_load(self, wb, name="report.xlsx"):
        wb.save()
        return load_workbook(self.path(name))


def _values(ws):
    return [[c.value for c in row] for row in ws.iter_rows()]


class SheetOrderTest(_TempDirCase):
    def test_summary_written_last_is_ranked_second(self):
        wb = ReportWorkbook(self.path())
        wb.ensure_sheet("bsp", "BSP Detailed", TWO_COLS, 4)
        wb.ensure_sheet("tp-api", "TP API", TWO_COLS, 18)
        wb.ensure_sheet("combined", "Combined", COMBINED_SPECS, 3)
        wb.append("bsp", ["bsp:1", 10])
        wb.append("combined", ["x"] * len(COMBINED_SPECS))
        wb.append("bsp", ["bsp:2", 20])            # back to an earlier sheet
        wb.write_blocks("readme", "Read Me", 1, [Block("Report", None, [["Name", "Test"]])])
        wb.write_blocks("summary", "Summary", 2, [Block("Totals", ["Rows"], [[3]])])
        loaded = self.save_and_load(wb)
        self.assertEqual(loaded.sheetnames, ["Read Me", "Summary", "Combined", "BSP Detailed", "TP API"])
        self.assertEqual(_values(loaded["BSP Detailed"]), [["Ref", "Amount"], ["bsp:1", 10], ["bsp:2", 20]])
        self.assertEqual(wb.row_counts(), {"bsp": 2, "tp-api": 0, "combined": 1, "readme": 1, "summary": 1})

    def test_interleaved_appends_keep_each_sheet_in_order(self):
        wb = ReportWorkbook(self.path())
        wb.ensure_sheet("a", "A", [ColSpec("n", "int")], 5)
        wb.ensure_sheet("b", "B", [ColSpec("n", "int")], 4)
        for i in range(50):
            wb.append("a" if i % 3 else "b", [i])
        loaded = self.save_and_load(wb)
        self.assertEqual(loaded.sheetnames, ["B", "A"])
        self.assertEqual([r[0] for r in _values(loaded["B"])[1:]], [i for i in range(50) if i % 3 == 0])
        self.assertEqual([r[0] for r in _values(loaded["A"])[1:]], [i for i in range(50) if i % 3])

    def test_empty_workbook_gets_a_no_data_sheet(self):
        loaded = self.save_and_load(ReportWorkbook(self.path()))
        self.assertEqual(loaded.sheetnames, ["No data"])

    def test_header_only_sheet_is_kept(self):
        wb = ReportWorkbook(self.path())
        wb.ensure_sheet("combined", "Combined", TWO_COLS, 3)
        loaded = self.save_and_load(wb)
        self.assertEqual(_values(loaded["Combined"]), [["Ref", "Amount"]])
        self.assertEqual(loaded["Combined"].auto_filter.ref, "A1:B1")


class DataSheetLayoutTest(_TempDirCase):
    def test_header_freeze_filter_widths_and_style(self):
        wb = ReportWorkbook(self.path())
        wb.ensure_sheet("combined", "Combined", COMBINED_SPECS, 3)
        for i in range(2):
            wb.append("combined", [f"r{i}"] + [None] * (len(COMBINED_SPECS) - 1))
        ws = self.save_and_load(wb)["Combined"]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        self.assertEqual(header, list(COMBINED_HEADERS))
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws.auto_filter.ref, "A1:BJ3")          # 62 columns, header + 2 rows
        self.assertEqual(ws.column_dimensions["C"].width, COMBINED_COLUMNS[2].width)
        first = ws["A1"]
        self.assertTrue(first.font.b)
        self.assertTrue(first.font.color.rgb.endswith("FFFFFF"))
        self.assertEqual(first.fill.fill_type, "solid")
        self.assertTrue(first.fill.fgColor.rgb.endswith("1E3A5F"))

    def test_unstyled_header_is_plain(self):
        wb = ReportWorkbook(self.path(), styled=False)
        wb.ensure_sheet("s", "S", TWO_COLS, 4)
        ws = self.save_and_load(wb)["S"]
        self.assertFalse(ws["A1"].font.b)
        self.assertEqual(ws.freeze_panes, "A2")

    def test_ensure_sheet_is_idempotent(self):
        wb = ReportWorkbook(self.path())
        wb.ensure_sheet("s", "S", TWO_COLS, 4)
        wb.ensure_sheet("s", "Other title", [ColSpec("x")], 9)
        self.assertTrue(wb.has_sheet("s"))
        self.assertFalse(wb.has_sheet("t"))
        wb.append("s", ["a", 1])
        self.assertEqual(self.save_and_load(wb).sheetnames, ["S"])

    def test_misuse_is_rejected(self):
        wb = ReportWorkbook(self.path())
        wb.ensure_sheet("s", "S", TWO_COLS, 4)
        with self.assertRaises(ValueError):
            wb.append("s", ["only one"])
        with self.assertRaises(KeyError):
            wb.append("missing", ["a", 1])
        with self.assertRaises(ValueError):
            wb.ensure_sheet("t", "s", TWO_COLS, 5)             # titles are case-insensitive
        with self.assertRaises(ValueError):
            wb.ensure_sheet("u", "U", [ColSpec("x", "percent")], 6)
        with self.assertRaises(ValueError):
            wb.ensure_sheet("v", "V", [], 7)
        with self.assertRaises(ValueError):
            ReportWorkbook(self.path("x.xlsx"), max_data_rows=0)
        wb.discard()

    def test_title_sanitised_and_truncated(self):
        wb = ReportWorkbook(self.path(), max_data_rows=1)
        wb.ensure_sheet("cta", "LCC CTA/BTA", TWO_COLS, 15)
        long_title = "A very long sheet title that exceeds"
        wb.ensure_sheet("long", long_title, TWO_COLS, 16)
        wb.append("long", ["a", 1])
        wb.append("long", ["b", 2])
        titles = wb.sheet_titles()
        self.assertEqual(titles["cta"], ["LCC CTA-BTA"])
        self.assertEqual(titles["long"][0], long_title[:31])
        self.assertTrue(titles["long"][1].endswith(" (2)"))
        self.assertLessEqual(len(titles["long"][1]), 31)
        self.assertEqual(self.save_and_load(wb).sheetnames, titles["cta"] + titles["long"])


class CellSafetyTest(_TempDirCase):
    def test_illegal_characters_stripped_and_file_reloads(self):
        wb = ReportWorkbook(self.path())
        wb.ensure_sheet("s", "S", [ColSpec("a"), ColSpec("b"), ColSpec("c"), ColSpec("d")], 4)
        wb.append("s", ["bell\x07here", "not\ufffechar", "lone\ud800surrogate", "\uffff\x00"])
        ws = self.save_and_load(wb)["S"]
        self.assertEqual(_values(ws)[1], ["bellhere", "notchar", "lonesurrogate", None])

    def test_formula_and_error_literals_stay_text(self):
        wb = ReportWorkbook(self.path())
        wb.ensure_sheet(
            "s", "S", [ColSpec("=cmd|' /C calc'!A0"), ColSpec("b"), ColSpec("c", "money"), ColSpec("d")], 4,
        )
        hyperlink = '=HYPERLINK("http://evil.example","click")'
        wb.append("s", [hyperlink, "#N/A", "-285.72", "="])
        path = self.path()
        loaded = self.save_and_load(wb)
        ws = loaded["S"]
        self.assertEqual(ws["A1"].value, "=cmd|' /C calc'!A0")
        self.assertEqual(ws["A1"].data_type, "s")
        self.assertTrue(ws["A1"].font.b)                        # still styled as a header
        self.assertEqual((ws["A2"].value, ws["A2"].data_type), (hyperlink, "s"))
        self.assertEqual((ws["B2"].value, ws["B2"].data_type), ("#N/A", "s"))
        self.assertEqual((ws["C2"].value, ws["C2"].data_type), ("-285.72", "s"))
        self.assertEqual(ws["C2"].number_format, "General")      # text in a money column
        self.assertEqual((ws["D2"].value, ws["D2"].data_type), ("=", "s"))
        with zipfile.ZipFile(path) as zf:
            sheet_xml = zf.read("xl/worksheets/sheet1.xml").decode("utf-8")
        self.assertNotIn("<f>", sheet_xml)
        self.assertNotIn('t="e"', sheet_xml)

    def test_long_string_truncated(self):
        wb = ReportWorkbook(self.path())
        wb.ensure_sheet("s", "S", [ColSpec("a")], 4)
        wb.append("s", ["x" * 40_000])
        ws = self.save_and_load(wb)["S"]
        self.assertEqual(len(ws["A2"].value), MAX_CELL_CHARS)

    def test_block_sheet_values_are_neutralised_too(self):
        wb = ReportWorkbook(self.path())
        wb.write_blocks("readme", "Read Me", 1, [Block("=title", ["=h"], [["=v", "ok\x01"]])])
        ws = self.save_and_load(wb)["Read Me"]
        for coord in ("A1", "A2", "A3"):
            self.assertEqual(ws[coord].data_type, "s", coord)
        self.assertEqual(ws["B3"].value, "ok")


class SplitTest(_TempDirCase):
    def test_max_data_rows_splits_into_parts_with_headers(self):
        wb = ReportWorkbook(self.path(), max_data_rows=3)
        wb.ensure_sheet("combined", "Combined", TWO_COLS, 3)
        wb.ensure_sheet("bsp", "BSP Detailed", TWO_COLS, 4)
        wb.write_blocks("summary", "Summary", 2, [Block(None, ["Rows"], [[7]])])
        for i in range(7):
            wb.append("combined", [f"c{i}", i])
        wb.append("bsp", ["b0", 1])
        self.assertEqual(wb.sheet_titles()["combined"], ["Combined", "Combined (2)", "Combined (3)"])
        self.assertEqual(wb.row_counts()["combined"], 7)
        loaded = self.save_and_load(wb)
        self.assertEqual(
            loaded.sheetnames, ["Summary", "Combined", "Combined (2)", "Combined (3)", "BSP Detailed"],
        )
        expected = {"Combined": [0, 1, 2], "Combined (2)": [3, 4, 5], "Combined (3)": [6]}
        for title, amounts in expected.items():
            ws = loaded[title]
            rows = _values(ws)
            with self.subTest(title=title):
                self.assertEqual(rows[0], ["Ref", "Amount"])
                self.assertEqual([r[1] for r in rows[1:]], amounts)
                self.assertEqual(ws.freeze_panes, "A2")
                self.assertEqual(ws.auto_filter.ref, f"A1:B{len(amounts) + 1}")

    def test_exact_multiple_does_not_open_an_empty_part(self):
        wb = ReportWorkbook(self.path(), max_data_rows=2)
        wb.ensure_sheet("s", "S", TWO_COLS, 4)
        for i in range(4):
            wb.append("s", [str(i), i])
        self.assertEqual(wb.sheet_titles()["s"], ["S", "S (2)"])
        self.assertEqual(self.save_and_load(wb).sheetnames, ["S", "S (2)"])


class NumberFormatTest(_TempDirCase):
    SPECS = [
        ColSpec("Money", "money"), ColSpec("Date", "date"), ColSpec("When", "datetime"),
        ColSpec("Count", "int"), ColSpec("Text"),
    ]

    def test_styled_formats(self):
        wb = ReportWorkbook(self.path())
        wb.ensure_sheet("s", "S", self.SPECS, 4)
        wb.append("s", [Decimal("1234.50"), date(2026, 7, 1), datetime(2026, 7, 1, 10, 30), 2, date(2026, 7, 2)])
        wb.append("s", [-285.72, None, None, None, None])
        ws = self.save_and_load(wb)["S"]
        self.assertEqual((ws["A2"].value, ws["A2"].number_format), (1234.5, "#,##0.00"))
        self.assertEqual(ws["B2"].number_format, "DD-MMM-YYYY")
        self.assertEqual(ws["B2"].value, datetime(2026, 7, 1))
        self.assertEqual(ws["C2"].number_format, "DD-MMM-YYYY HH:MM")
        self.assertEqual(ws["C2"].value, datetime(2026, 7, 1, 10, 30))
        self.assertEqual((ws["D2"].value, ws["D2"].number_format), (2, "General"))
        self.assertTrue(ws["E2"].is_date)                       # a date in a text column is still a date
        self.assertEqual((ws["A3"].value, ws["A3"].number_format), (-285.72, "#,##0.00"))
        self.assertIsNone(ws["B3"].value)

    def test_unstyled_keeps_real_dates_without_custom_formats(self):
        wb = ReportWorkbook(self.path(), styled=False)
        wb.ensure_sheet("s", "S", self.SPECS, 4)
        wb.append("s", [Decimal("10"), date(2026, 7, 1), datetime(2026, 7, 1, 10, 30), 1, "t"])
        ws = self.save_and_load(wb)["S"]
        self.assertEqual((ws["A2"].value, ws["A2"].number_format), (10, "General"))
        # exactly the formats openpyxl gives a bare date / datetime
        self.assertEqual((ws["B2"].value, ws["B2"].number_format), (datetime(2026, 7, 1), "yyyy-mm-dd"))
        self.assertEqual(ws["C2"].number_format, "yyyy-mm-dd h:mm:ss")
        self.assertEqual(ws["E2"].value, "t")

    def test_block_kinds_apply_formats_and_layout(self):
        wb = ReportWorkbook(self.path())
        wb.write_blocks("summary", "Summary", 2, [
            Block("Net by currency", ["Currency", "Net"], [["INR", Decimal("100.5")], ["USD", 7]], ["text", "money"]),
            Block("Flags", ["Flag", "Rows"], [["DATE_UNREADABLE", 3]]),
        ])
        ws = self.save_and_load(wb)["Summary"]
        self.assertEqual(_values(ws), [
            ["Net by currency", None], ["Currency", "Net"], ["INR", 100.5], ["USD", 7],
            [None, None], ["Flags", None], ["Flag", "Rows"], ["DATE_UNREADABLE", 3],
        ])
        self.assertTrue(ws["A1"].font.b)
        self.assertTrue(ws["A2"].fill.fgColor.rgb.endswith("1E3A5F"))
        self.assertEqual(ws["B3"].number_format, "#,##0.00")
        self.assertEqual(ws["B8"].number_format, "General")
        self.assertIsNone(ws.freeze_panes)
        self.assertEqual(wb.row_counts(), {"summary": 3})

    def test_write_blocks_once_per_key(self):
        wb = ReportWorkbook(self.path())
        wb.write_blocks("readme", "Read Me", 1, [])
        with self.assertRaises(ValueError):
            wb.write_blocks("readme", "Read Me", 1, [])
        with self.assertRaises(ValueError):
            wb.write_blocks("other", "Other", 1, [Block(None, None, [["x"]], ["percent"])])
        with self.assertRaises(KeyError):
            wb.append("readme", ["x"])
        self.assertEqual(self.save_and_load(wb).sheetnames, ["Read Me"])


class ExcelValueTest(unittest.TestCase):
    def test_conversions(self):
        self.assertEqual(excel_value(True), "Yes")
        self.assertEqual(excel_value(False), "No")
        self.assertEqual(excel_value(Decimal("12.50")), 12.5)
        self.assertIsInstance(excel_value(Decimal("12.50")), float)
        self.assertEqual(excel_value(["a", None, "", "b"]), "a, b")
        self.assertEqual(excel_value(("x", 1, Decimal("2.50"))), "x, 1, 2.50")
        self.assertEqual(excel_value({"code": "YQ", "amount": 10}), '{"code": "YQ", "amount": 10}')
        self.assertEqual(excel_value({"b", "a"}), "a, b")
        self.assertIsNone(excel_value(None))
        self.assertIsNone(excel_value(""))
        self.assertIsNone(excel_value([]))
        self.assertEqual(excel_value(5), 5)
        self.assertEqual(excel_value(12345678901234567), "12345678901234567")
        self.assertEqual(excel_value(float("nan")), "NaN")
        self.assertEqual(excel_value(Decimal("-Infinity")), "-Infinity")
        self.assertEqual(excel_value(b"caf\xc3\xa9"), "café")
        self.assertEqual(excel_value("  spaced  "), "  spaced  ")

    def test_dates_pass_and_aware_datetimes_become_naive_utc(self):
        d = date(2026, 7, 1)
        self.assertIs(excel_value(d), d)
        naive = datetime(2026, 7, 1, 10, 0)
        self.assertIs(excel_value(naive), naive)
        ist = timezone(timedelta(hours=5, minutes=30))
        self.assertEqual(excel_value(datetime(2026, 7, 1, 10, 0, tzinfo=ist)), datetime(2026, 7, 1, 4, 30))

    def test_strings_cleaned_and_truncated(self):
        self.assertEqual(excel_value("a\x0bb\x0cc\ud83dd"), "abcd")
        self.assertEqual(excel_value("tab\tnew\nline\r"), "tab\tnew\nline\r")
        self.assertEqual(len(excel_value("é" * 40_000)), MAX_CELL_CHARS)
        self.assertEqual(len(excel_value(["y" * 40_000])), MAX_CELL_CHARS)


class LifecycleTest(_TempDirCase):
    def _writer_files(self, wb):
        return [ws._writer.out for ws in wb._wb._sheets if ws._writer is not None]

    def test_discard_closes_writers_and_removes_temp_files(self):
        wb = ReportWorkbook(self.path(), max_data_rows=2)
        wb.ensure_sheet("s", "S", TWO_COLS, 4)
        for i in range(5):                                     # parts 1-2 closed early, part 3 open
            wb.append("s", [str(i), i])
        wb.ensure_sheet("never", "Never appended", TWO_COLS, 5)
        temp_files = self._writer_files(wb)
        self.assertEqual(len(temp_files), 4)
        self.assertTrue(all(os.path.exists(p) for p in temp_files))
        wb.discard()
        wb.discard()                                           # repeatable
        self.assertFalse(any(os.path.exists(p) for p in temp_files))
        self.assertFalse(os.path.exists(self.path()))
        with self.assertRaises(RuntimeError):
            wb.append("s", ["x", 1])
        with self.assertRaises(RuntimeError):
            wb.save()

    def test_discard_after_save_keeps_the_file(self):
        wb = ReportWorkbook(self.path())
        wb.ensure_sheet("s", "S", TWO_COLS, 4)
        wb.append("s", ["a", 1])
        wb.save()
        wb.discard()
        self.assertEqual(load_workbook(self.path()).sheetnames, ["S"])
        with self.assertRaises(RuntimeError):
            wb.save()

    def test_failed_save_removes_partial_file_and_discard_cleans_up(self):
        wb = ReportWorkbook(self.path())
        wb.ensure_sheet("s", "S", TWO_COLS, 4)
        wb.append("s", ["a", 1])
        temp_files = self._writer_files(wb)

        class BrokenWriter:
            def __init__(self, workbook, archive):
                archive.writestr("partial.txt", "x")

            def save(self):
                raise OSError("disk full")

        with mock.patch.object(wbmod, "ExcelWriter", BrokenWriter):
            with self.assertRaises(OSError):
                wb.save()
        self.assertFalse(os.path.exists(self.path()))
        wb.discard()
        self.assertFalse(any(os.path.exists(p) for p in temp_files))


if __name__ == "__main__":
    unittest.main()
