"""Report download Summary sheet — totals an accountant can tie out, never double counted.

No DB. Rows are plain dicts shaped like Combined rows and BSP rows are SimpleNamespace
views, so the accumulator is tested exactly as the builder streams into it:

  * per-currency totals take "Yes" rows only, are never added across currencies, and carry
    a subtotal per (Category, Currency) plus a grand total per currency;
  * the BSP tie-out buckets transaction_amount by settlement section, shows the variance,
    OK within ±1.00, and "n/a" for partial scope or a missing statement total;
  * sign counters, remittance identity counts, exposure buckets, linking %, flags;
  * to_json() stays JSON-safe and small; the blocks render through ReportWorkbook.

Run:  python -m unittest test_report_download_summary   (from backend/tests)
"""

import json
import os
import shutil
import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal
from types import SimpleNamespace

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from openpyxl import load_workbook  # noqa: E402

from app.services.report_download import columns as C  # noqa: E402
from app.services.report_download import summary as S  # noqa: E402
from app.services.report_download.summary import SummaryAccumulator  # noqa: E402
from app.services.report_download.workbook import ReportWorkbook  # noqa: E402

D = Decimal


def row(**kw):
    r = {k: None for k in C.COMBINED_KEYS}
    r.update(category="BSP", source_type="BSP", currency="INR", counts_in_net=C.NET_YES, data_flags=[])
    r.update(kw)
    return r


def block(acc, title):
    found = [b for b in acc.blocks() if b.title == title]
    return found[0] if found else None


def find(rows, *prefix):
    return [r for r in rows if list(r[:len(prefix)]) == list(prefix)]


def statement(**kw):
    base = dict(
        file_name="BSP_AUG.pdf", statement_name=None, period_from=date(2026, 8, 1), period_to=date(2026, 8, 15),
        gt_issues=D("1000.00"), gt_refunds=D("-200.00"), gt_debit_memos=D("50.00"), gt_credit_memos=D("-30.00"),
        gt_std_comm=D("-10.00"), gt_sup_comm=D("0"), gt_tax_on_comm=D("1.80"), gt_balance_payable=D("810.00"),
        gt_doc_count=4,
    )
    base.update(kw)
    return SimpleNamespace(**base)


def bsp_row(txn, *, ttype="TKTT", fop="CA", std=None, supp=None, toc=None, bal=None, raw=None):
    return SimpleNamespace(
        transaction_amount=txn, standard_commission_amount=std, supplier_discount_amount=supp,
        tax_on_commission=toc, balance_payable=bal, transaction_type=ttype, form_of_payment=fop,
        raw_data=raw,
    )


class NetBlockTests(unittest.TestCase):
    def setUp(self):
        self.acc = SummaryAccumulator()
        a = self.acc.add_combined
        a(row(base_fare=D("80"), total_taxes=D("20"), gross_amount=D("100"), commission=D("-5"), net_payable=D("95")))
        a(row(base_fare=D("40"), gross_amount=D("50"), net_payable=D("48.50")))
        a(row(counts_in_net=C.NET_SUPERSEDED, gross_amount=D("100"), net_payable=D("95")))
        a(row(currency="USD", gross_amount=D("10"), net_payable=D("9")))
        a(row(source_type="ADM", counts_in_net=C.NET_MEMO_IN_BILLING, net_payable=D("300")))
        a(row(category="Third Party", source_type="Third Party GDS", gross_amount=D("1000"), net_payable=D("990")))
        self.rows = block(self.acc, S.T_NET).rows

    def test_counted_only_and_not_counted_columns(self):
        (bsp_inr,) = find(self.rows, "BSP", "BSP", "INR")
        # Rows, counted, base, taxes, gross, commission, net, not counted, net not counted
        self.assertEqual(bsp_inr[3:], [3, 2, D("120"), D("20"), D("150"), D("-5"), D("143.50"), 1, D("95")])

    def test_blank_money_stays_blank(self):
        (adm,) = find(self.rows, "BSP", "ADM", "INR")
        self.assertEqual(adm[3:5], [1, 0])
        self.assertIsNone(adm[9])                       # no counted net → blank, not 0
        self.assertEqual(adm[11], D("300"))
        (usd,) = find(self.rows, "BSP", "BSP", "USD")
        self.assertIsNone(usd[5])                       # no base fare given

    def test_subtotals_per_category_and_currency(self):
        (sub_inr,) = find(self.rows, "BSP", "Subtotal", "INR")
        self.assertEqual(sub_inr[3:], [4, 2, D("120"), D("20"), D("150"), D("-5"), D("143.50"), 2, D("395")])
        (sub_usd,) = find(self.rows, "BSP", "Subtotal", "USD")
        self.assertEqual(sub_usd[9], D("9"))
        (tp,) = find(self.rows, "Third Party", "Subtotal", "INR")
        self.assertEqual(tp[9], D("990"))

    def test_grand_total_per_currency_never_mixed(self):
        (g_inr,) = find(self.rows, "All categories", "Grand total", "INR")
        (g_usd,) = find(self.rows, "All categories", "Grand total", "USD")
        self.assertEqual(g_inr[9], D("1133.50"))
        self.assertEqual(g_usd[9], D("9"))
        self.assertEqual(len(find(self.rows, "All categories")), 2)

    def test_order_category_then_currency_inr_first(self):
        labels = [(r[0], r[1], r[2]) for r in self.rows]
        self.assertEqual(labels, [
            ("BSP", "BSP", "INR"), ("BSP", "ADM", "INR"), ("BSP", "Subtotal", "INR"),
            ("BSP", "BSP", "USD"), ("BSP", "Subtotal", "USD"),
            ("Third Party", "Third Party GDS", "INR"), ("Third Party", "Subtotal", "INR"),
            ("All categories", "Grand total", "INR"), ("All categories", "Grand total", "USD"),
        ])

    def test_money_kinds(self):
        b = block(self.acc, S.T_NET)
        self.assertEqual(len(b.kinds), len(b.headers))
        self.assertEqual([h for h, k in zip(b.headers, b.kinds) if k == "money"],
                         ["Base Fare", "Total Taxes", "Gross", "Commission", "Net Payable", "Net Payable not counted"])

    def test_by_transaction_type(self):
        acc = SummaryAccumulator()
        acc.add_combined(row(transaction_type=C.REFUND, gross_amount=D("-100"), net_payable=D("-90")))
        acc.add_combined(row(transaction_type=C.SALE, gross_amount=D("500"), net_payable=D("450")))
        acc.add_combined(row(transaction_type=C.SALE, gross_amount=D("7"), net_payable=D("7"),
                             counts_in_net=C.NET_SUPERSEDED))
        rows = block(acc, S.T_BY_TYPE).rows
        self.assertEqual(rows, [
            ["BSP", "INR", C.SALE, 2, 1, D("500"), D("450")],
            ["BSP", "INR", C.REFUND, 1, 1, D("-100"), D("-90")],
        ])


class EmptyTests(unittest.TestCase):
    def test_empty_accumulator_has_only_first_block_with_no_rows(self):
        blocks = SummaryAccumulator().blocks()
        self.assertEqual([b.title for b in blocks], [S.T_NET])
        self.assertEqual(blocks[0].rows[0][0], "No rows")
        self.assertEqual(len(blocks[0].rows[0]), len(S.NET_HEADERS))

    def test_empty_to_json(self):
        js = SummaryAccumulator().to_json()
        self.assertEqual(js["rows"], 0)
        self.assertIsNone(js["tie_out_ok"])
        json.dumps(js)


class TieOutTests(unittest.TestCase):
    def feed(self, acc, sid="S1"):
        acc.add_bsp_row(sid, bsp_row(D("600"), std=D("-6"), toc=D("1.08"), bal=D("595.08")), "ISSUES")
        acc.add_bsp_row(sid, bsp_row(D("400"), std=D("-4"), toc=D("0.72"), bal=D("396.72")), "ISSUE")
        acc.add_bsp_row(sid, bsp_row(D("-200"), ttype="RFND", fop="CC", bal=D("-200")), "REFUNDS")
        # section None → raw_data: settlement_section wins over the printed section
        acc.add_bsp_row(sid, bsp_row(D("48"), ttype="CANX", bal=D("48"),
                                     raw={"section": "ISSUES", "settlement_section": "DEBIT MEMOS"}), None)

    def lines(self, acc):
        return {r[3]: r for r in block(acc, S.T_TIE_OUT).rows}

    def test_variance_and_ok(self):
        acc = SummaryAccumulator()
        acc.set_bsp_statement("S1", statement(), None, "full")
        self.feed(acc)
        lines = self.lines(acc)
        self.assertEqual(lines["Issues"][4:], [D("1000"), D("1000.00"), D("0.00"), "Yes"])
        self.assertEqual(lines["Refunds"][4:], [D("-200"), D("-200.00"), D("0.00"), "Yes"])
        # 48 vs 50: variance −2.00 is beyond the ±1.00 tolerance
        self.assertEqual(lines["Debit memos"][4:], [D("48"), D("50.00"), D("-2.00"), "No"])
        # nothing in a CREDIT section: report Σ 0 vs −30
        self.assertEqual(lines["Credit memos"][4:], [D("0"), D("-30.00"), D("30.00"), "No"])
        self.assertEqual(lines["Std comm"][4:], [D("-10"), D("-10.00"), D("0.00"), "Yes"])
        self.assertEqual(lines["Tax on comm"][6:], [D("0.00"), "Yes"])
        self.assertEqual(lines["Balance payable"][4], D("839.80"))
        self.assertEqual(lines["Balance payable"][7], "No")
        self.assertEqual(lines["Doc count"][4:], [4, 4, 0, "Yes"])
        first = lines["Issues"]
        self.assertEqual(first[:3], ["BSP_AUG.pdf", "01-Aug-2026 to 15-Aug-2026", "Whole statement"])
        self.assertFalse(acc.to_json()["tie_out_ok"])

    def test_within_tolerance_is_ok(self):
        acc = SummaryAccumulator()
        acc.set_bsp_statement("S1", statement(gt_issues=D("999.10")), None, "full")
        self.feed(acc)
        self.assertEqual(self.lines(acc)["Issues"][6:], [D("0.90"), "Yes"])

    def test_doc_count_has_no_tolerance(self):
        acc = SummaryAccumulator()
        acc.set_bsp_statement("S1", statement(gt_doc_count=5), None, "full")
        self.feed(acc)
        self.assertEqual(self.lines(acc)["Doc count"][6:], [-1, "No"])

    def test_partial_scope_is_na(self):
        acc = SummaryAccumulator()
        acc.set_bsp_statement("S1", statement(), None, "partial")
        self.feed(acc)
        lines = self.lines(acc)
        self.assertTrue(all(r[7] == "n/a" for r in lines.values()))
        self.assertEqual(lines["Debit memos"][6], D("-2.00"))      # variance still shown
        self.assertIn("Partial", lines["Issues"][2])
        self.assertIsNone(acc.to_json()["tie_out_ok"])

    def test_missing_statement_total_is_na(self):
        acc = SummaryAccumulator()
        acc.set_bsp_statement("S1", statement(gt_refunds=None, gt_doc_count=None), None, "full")
        self.feed(acc)
        lines = self.lines(acc)
        self.assertEqual(lines["Refunds"][5:], [None, None, "n/a"])
        self.assertEqual(lines["Doc count"][5:], [None, None, "n/a"])

    def test_all_ok_json(self):
        acc = SummaryAccumulator()
        acc.set_bsp_statement("S1", statement(
            gt_debit_memos=D("48"), gt_credit_memos=D("0"), gt_balance_payable=D("839.80")), None, "full")
        self.feed(acc)
        self.assertTrue(acc.to_json()["tie_out_ok"])

    def test_no_statements_omits_block(self):
        acc = SummaryAccumulator()
        acc.add_combined(row(net_payable=D("1")))
        self.assertIsNone(block(acc, S.T_TIE_OUT))

    def test_section_keywords(self):
        self.assertEqual(S.section_line("ISSUES"), "issues")
        self.assertEqual(S.section_line("refund"), "refunds")
        self.assertEqual(S.section_line("AGENT DEBIT MEMOS"), "debit_memos")
        self.assertEqual(S.section_line("CREDIT"), "credit_memos")
        self.assertIsNone(S.section_line("TOTALS"))
        self.assertIsNone(S.section_line(None))

    def test_float_header_values(self):
        acc = SummaryAccumulator()
        acc.set_bsp_statement("S1", statement(gt_issues=1000.0), None, "full")
        self.feed(acc)
        self.assertEqual(self.lines(acc)["Issues"][7], "Yes")

    def test_summary_vs_detailed_block(self):
        detail = {
            "issues": {"summary": "1000.00", "detail": "1000.00", "variance": "0.00", "ok": True},
            "refunds": {"summary": "-195.00", "detail": "-200.00", "variance": "5.00", "ok": False},
        }
        sm = SimpleNamespace(file_name="SUMMARY.pdf", period_from=date(2026, 8, 1), period_to=date(2026, 8, 15),
                             match_status="mismatch", match_detail=detail)
        acc = SummaryAccumulator()
        acc.set_bsp_statement("S1", statement(), sm, "full")
        acc.set_bsp_statement("S2", statement(file_name="B.pdf"),
                              SimpleNamespace(file_name="S2.pdf", match_status="pending", match_detail=None), "full")
        rows = block(acc, S.T_SUMMARY_VS_DETAILED).rows
        self.assertEqual(rows[0], ["SUMMARY.pdf", "BSP_AUG.pdf", "01-Aug-2026 to 15-Aug-2026", "mismatch",
                                   "Issues", D("1000.00"), D("1000.00"), D("0.00"), "Yes"])
        self.assertEqual(rows[1][4:], ["Refunds", D("-195.00"), D("-200.00"), D("5.00"), "No"])
        self.assertEqual(rows[2][:5], ["S2.pdf", "B.pdf", "01-Aug-2026 to 15-Aug-2026", "pending", None])

    def test_no_summary_headers_omits_block(self):
        acc = SummaryAccumulator()
        acc.set_bsp_statement("S1", statement(), None, "full")
        self.assertIsNone(block(acc, S.T_SUMMARY_VS_DETAILED))


class SignAndIdentityTests(unittest.TestCase):
    def test_bsp_sign_counters_by_txn_and_fop(self):
        acc = SummaryAccumulator()
        acc.add_bsp_row("S1", bsp_row(D("100"), bal=D("100")), "ISSUES")
        acc.add_bsp_row("S1", bsp_row(D("50"), bal=D("0"), fop="CC"), "ISSUES")
        acc.add_bsp_row("S1", bsp_row(D("-80"), ttype="rfnd", fop="CC", bal=D("-80")), "REFUNDS")
        acc.add_bsp_row("S1", bsp_row(D("0"), ttype="TKTT", fop="MS", bal=None), "ISSUES")
        rows = {(r[1], r[2]): r[3:] for r in block(acc, S.T_SIGNS).rows if r[1] != "Remittance identity"}
        self.assertEqual(rows[("TKTT × CA", "Transaction Amount")], [0, 1, 0, 1])
        self.assertEqual(rows[("TKTT × CC", "Balance Payable")], [0, 0, 1, 1])
        self.assertEqual(rows[("RFND × CC", "Transaction Amount")], [1, 0, 0, 1])
        self.assertEqual(rows[("TKTT × Other", "Transaction Amount")], [0, 0, 1, 1])
        self.assertNotIn(("TKTT × Other", "Balance Payable"), rows)       # blank skipped

    def test_remittance_identity(self):
        acc = SummaryAccumulator()
        # commission stored negative: balance = txn + std + supp + toc
        acc.add_bsp_row("S", bsp_row(D("1000"), std=D("-50"), supp=D("-10"), toc=D("9"), bal=D("949")), None)
        # commission stored positive: balance = txn − std − supp + toc
        acc.add_bsp_row("S", bsp_row(D("1000"), std=D("50"), supp=None, toc=D("9"), bal=D("959.40")), None)
        # no commission: both hold
        acc.add_bsp_row("S", bsp_row(D("500"), bal=D("500")), None)
        acc.add_bsp_row("S", bsp_row(D("500"), std=D("-5"), bal=D("300")), None)
        acc.add_bsp_row("S", bsp_row(None, bal=D("1")), None)
        rows = {r[2]: r[6] for r in block(acc, S.T_SIGNS).rows if r[1] == "Remittance identity"}
        self.assertEqual(rows, {S.IDENTITY_PLUS: 1, S.IDENTITY_MINUS: 1, S.IDENTITY_BOTH: 1,
                                S.IDENTITY_NEITHER: 1, S.IDENTITY_SKIPPED: 1})

    def test_generic_sign_check(self):
        acc = SummaryAccumulator()
        for v in (D("-1"), "-2.50", 3, 0.0, None, "n/a"):
            acc.add_sign_check("ADM", "Stored amount", "Amount", v)
        (r,) = block(acc, S.T_SIGNS).rows
        self.assertEqual(r, ["ADM", "Stored amount", "Amount", 2, 1, 1, 4])


class ExposureTests(unittest.TestCase):
    def setUp(self):
        acc = self.acc = SummaryAccumulator()
        tgq = dict(source_type="TGQ HMPR")
        acc.add_combined(row(**tgq, counts_in_net=C.NET_TGQ_NOT_IN_BSP, not_in_bsp="Yes", gross_amount=D("100")))
        acc.add_combined(row(**tgq, counts_in_net=C.NET_TGQ_NOT_IN_BSP, not_in_bsp="Yes", gross_amount=D("50")))
        acc.add_combined(row(**tgq, counts_in_net=C.NET_TGQ_OTHER_UPLOAD,
                             not_in_bsp="In another BSP upload (JULY.pdf)", gross_amount=D("70")))
        acc.add_combined(row(**tgq, counts_in_net=C.NET_TGQ_UNMATCHABLE, not_in_bsp="Unknown – no ticket no.",
                             gross_amount=D("5"), currency="USD"))
        acc.add_combined(row(**tgq, counts_in_net=C.NET_TGQ_OUTSIDE_PERIOD, gross_amount=D("9")))
        acc.add_combined(row(**tgq, counts_in_net=C.NET_SUPERSEDED, not_in_bsp="Yes", gross_amount=D("100")))
        acc.add_combined(row(source_type="ADM",
                             counts_in_net=C.NET_MEMO_NOT_BILLED + " (status: PENDING, sent to DPC)",
                             gross_amount=D("250"), net_payable=D("250")))
        acc.add_combined(row(source_type="ACM", counts_in_net=C.NET_MEMO_NOT_BILLED, status="OPEN",
                             net_payable=D("-40")))
        acc.add_combined(row(source_type="RA", counts_in_net=C.NET_MEMO_IN_BILLING, net_payable=D("-10")))
        acc.add_combined(row(source_type="NDC", counts_in_net=C.NET_NDC_IN_BSP, gross_amount=D("800"),
                             net_payable=D("780")))
        acc.add_combined(row(source_type="NDC", counts_in_net=C.NET_YES, net_payable=D("1")))
        acc.add_combined(row(category="Third Party", source_type="Third Party GDS",
                             counts_in_net=C.NET_CANCELLED_VERIFY, gross_amount=D("300"), net_payable=D("290")))
        self.rows = block(acc, S.T_EXPOSURE).rows

    def test_tgq_buckets_by_currency(self):
        self.assertEqual(find(self.rows, S.EXP_TGQ, "Not in BSP"), [[S.EXP_TGQ, "Not in BSP", "INR", 2, D("150"), None]])
        self.assertEqual(find(self.rows, S.EXP_TGQ, "In another BSP upload")[0][3:5], [1, D("70")])
        self.assertEqual(find(self.rows, S.EXP_TGQ, "Unmatchable")[0][2:5], ["USD", 1, D("5")])
        self.assertEqual(find(self.rows, S.EXP_TGQ, "BSP row outside period")[0][3:5], [1, D("9")])

    def test_superseded_is_not_exposure(self):
        total = sum(r[3] for r in self.rows if r[0] == S.EXP_TGQ)
        self.assertEqual(total, 5)

    def test_memos_not_yet_billed_by_kind_and_status(self):
        self.assertEqual(find(self.rows, S.EXP_MEMOS),
                         [[S.EXP_MEMOS, "ADM – PENDING, sent to DPC", "INR", 1, D("250"), D("250")],
                          [S.EXP_MEMOS, "ACM – OPEN", "INR", 1, None, D("-40")]])

    def test_ndc_and_tp_cancelled(self):
        self.assertEqual(find(self.rows, S.EXP_NDC), [[S.EXP_NDC, S.NDC_THIS_REPORT, "INR", 1, D("800"), D("780")]])
        self.assertEqual(find(self.rows, S.EXP_TP_CANCELLED),
                         [[S.EXP_TP_CANCELLED, "Third Party GDS", "INR", 1, D("300"), D("290")]])

    def test_section_order(self):
        sections = [r[0] for r in self.rows]
        self.assertEqual(sorted(set(sections), key=sections.index), [S.EXP_TGQ, S.EXP_MEMOS, S.EXP_NDC, S.EXP_TP_CANCELLED])


class LinkingTpApiAndFlagsTests(unittest.TestCase):
    def test_derived_and_explicit_link_stats(self):
        acc = SummaryAccumulator()
        for te in ("Yes (document)", "Yes (document)", "Yes (rtdn)", "No", None):
            acc.add_combined(row(tgq_enriched=te))
        acc.add_combined(row(source_type="TGQ HMPR", counts_in_net=C.NET_SUPERSEDED, not_in_bsp="Yes"))
        acc.add_combined(row(source_type="TGQ HMPR", counts_in_net=C.NET_TGQ_NOT_IN_BSP, not_in_bsp="Yes",
                             data_flags=["LINK_AMBIGUOUS"]))
        acc.add_combined(row(category="LCC", source_type="LCC DI Statement", counts_in_net=C.NET_LEDGER,
                             linked_document="lcc_detailed:5"), source_key="lcc-di")
        acc.add_link_stat("TGQ tickets suppressed as in BSP", 6)
        acc.add_link_stat("Duplicate TGQ suppressed")
        acc.add_link_stat("Duplicate TGQ suppressed")
        acc.add_link_stat("BSP rows TGQ-enriched (rtdn)", 5)          # explicit replaces derived
        b = block(acc, S.T_LINKING)
        rows = {r[0]: r[1:] for r in b.rows}
        self.assertEqual(rows["BSP rows TGQ-enriched (document)"], [2, D("40.0"), "BSP rows"])
        self.assertEqual(rows["BSP rows TGQ-enriched (rtdn)"][0], 5)
        # TGQ tickets = 2 emitted rows + 6 suppressed + 2 duplicates
        self.assertEqual(rows["TGQ tickets suppressed as in BSP"], [6, D("60.0"), "TGQ tickets"])
        self.assertEqual(rows["Duplicate TGQ suppressed"], [2, D("20.0"), "TGQ tickets"])
        self.assertEqual(rows["Superseded rows (TGQ HMPR)"], [1, D("50.0"), "TGQ HMPR rows"])
        self.assertEqual(rows["LINK_AMBIGUOUS"], [1, D("12.5"), "All Combined rows"])
        self.assertEqual(rows["LCC ledger rows linked to LCC Detailed"][:2], [1, D("100.0")])
        self.assertEqual(b.rows[0][0], "BSP rows TGQ-enriched (document)")
        self.assertEqual(b.rows[-1][0], "LINK_AMBIGUOUS")
        js = acc.to_json()["link_stats"]
        self.assertEqual(js["BSP rows TGQ-enriched (rtdn)"], 5)
        self.assertEqual(js["Duplicate TGQ suppressed"], 2)

    def test_unknown_link_stat_has_no_percent(self):
        acc = SummaryAccumulator()
        acc.add_link_stat("Something else", 3)
        self.assertEqual(block(acc, S.T_LINKING).rows, [["Something else", 3, None, None]])

    def test_tp_api_by_product(self):
        acc = SummaryAccumulator()
        tp = dict(category="Third Party", source_type="Third Party API")
        acc.add_combined(row(**tp, product="Air", net_payable=D("100")))
        acc.add_combined(row(**tp, product="Hotel", net_payable=D("40")))
        acc.add_combined(row(**tp, product="Hotel", net_payable=D("60")))
        acc.add_combined(row(**tp, product="Hotel", counts_in_net=C.NET_NEEDS_REVIEW, net_payable=D("999")))
        acc.add_combined(row(category="Third Party", source_type="Third Party GDS", product="Air", net_payable=D("7")))
        rows = block(acc, S.T_TP_API).rows
        self.assertEqual(rows, [["Air", "INR", 1, 1, D("100")], ["Hotel", "INR", 3, 2, D("100")]])

    def test_flags_block_top_sources(self):
        acc = SummaryAccumulator()
        for _ in range(3):
            acc.add_combined(row(source_type="TGQ HMPR", counts_in_net=C.NET_TGQ_NOT_IN_BSP,
                                 data_flags=["CURRENCY_ASSUMED", "TGQ_PRE_SPLIT"]))
        for src, n in (("BSP", 5), ("NDC", 2), ("ADM", 1), ("RA", 1)):
            for _ in range(n):
                acc.add_combined(row(source_type=src, data_flags="CURRENCY_ASSUMED; LINK_AMBIGUOUS"))
        rows = block(acc, S.T_FLAGS).rows
        self.assertEqual(rows[0], ["CURRENCY_ASSUMED", C.FLAG_LEGEND["CURRENCY_ASSUMED"], 12,
                                   "BSP (5), TGQ HMPR (3), NDC (2)"])
        self.assertEqual([r[0] for r in rows], ["CURRENCY_ASSUMED", "LINK_AMBIGUOUS", "TGQ_PRE_SPLIT"])
        self.assertEqual(acc.to_json()["flag_counts"], {"CURRENCY_ASSUMED": 12, "LINK_AMBIGUOUS": 9, "TGQ_PRE_SPLIT": 3})


class JsonAndWorkbookTests(unittest.TestCase):
    def build(self, n=6000):
        acc = SummaryAccumulator()
        sources = [("BSP", "BSP"), ("BSP", "TGQ HMPR"), ("LCC", "LCC Detailed Statement"), ("Third Party", "Third Party API")]
        for i in range(n):
            cat, src = sources[i % len(sources)]
            acc.add_combined(row(
                category=cat, source_type=src, currency="USD" if i % 7 == 0 else "INR",
                counts_in_net=C.NET_YES if i % 3 else C.NET_SUPERSEDED, transaction_type=C.SALE,
                product="Hotel" if src == "Third Party API" else None,
                gross_amount=D("100.10"), net_payable=D(i) / 100, data_flags=["CURRENCY_ASSUMED"],
            ))
        acc.set_bsp_statement("S1", statement(), SimpleNamespace(file_name="sum.pdf", match_status="matched",
                                                                 match_detail={"issues": {"summary": "1", "detail": "1", "variance": "0", "ok": True}}), "full")
        acc.add_bsp_row("S1", bsp_row(D("1000"), bal=D("1000")), "ISSUES")
        acc.add_sign_check("RA", "Stored", "Amount", D("-3"))
        acc.add_link_stat("TGQ tickets suppressed as in BSP", 10)
        return acc

    def test_to_json_serialisable_and_small(self):
        acc = self.build()
        js = acc.to_json()
        text = json.dumps(js)
        self.assertLess(len(text), 50_000)
        self.assertEqual(js["rows"], 6000)
        self.assertEqual(js["counted_rows"], 4000)
        self.assertEqual(list(js["by_category"]), ["BSP", "LCC", "Third Party"])
        self.assertEqual(js["by_category"]["BSP"]["rows"], 3000)
        self.assertEqual(set(js["net_by_currency"]), {"INR", "USD"})
        self.assertIsInstance(js["net_by_currency"]["INR"], str)
        counted = sum(D(i) / 100 for i in range(6000) if i % 3)
        self.assertEqual(sum(D(v) for v in js["net_by_currency"].values()), counted)
        self.assertEqual(js["link_stats"]["TGQ tickets suppressed as in BSP"], 10)

    def test_many_link_names_stay_capped(self):
        acc = SummaryAccumulator()
        for i in range(5000):
            acc.add_link_stat(f"metric {i} " + "x" * 300)
        self.assertLess(len(json.dumps(acc.to_json())), 50_000)

    def test_blocks_render_through_workbook(self):
        acc = self.build(400)
        acc.add_combined(row(source_type="NDC", counts_in_net=C.NET_NDC_IN_BSP, net_payable=D("5")))
        blocks = acc.blocks()
        self.assertEqual([b.title for b in blocks], [
            S.T_NET, S.T_BY_TYPE, S.T_TIE_OUT, S.T_SUMMARY_VS_DETAILED, S.T_EXPOSURE, S.T_LINKING,
            S.T_SIGNS, S.T_TP_API, S.T_FLAGS,
        ])
        for b in blocks:
            self.assertEqual(len(b.kinds), len(b.headers), b.title)
            self.assertTrue(all(len(r) == len(b.headers) for r in b.rows), b.title)
        tmp = tempfile.mkdtemp(prefix="report-summary-test-")
        try:
            path = os.path.join(tmp, "s.xlsx")
            wb = ReportWorkbook(path)
            wb.write_blocks("summary", "Summary", 2, blocks)
            wb.save()
            ws = load_workbook(path)["Summary"]
            col_a = [c.value for c in ws["A"]]
            for b in blocks:
                self.assertIn(b.title, col_a)
            # a money cell is a real number with the money format
            header_row = col_a.index(S.T_NET) + 2
            first = ws.cell(row=header_row + 1, column=10)
            self.assertIsInstance(first.value, float)
            self.assertEqual(first.number_format, "#,##0.00")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
