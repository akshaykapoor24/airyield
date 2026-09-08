"""One reader, deciding the format from the bytes and the header row from the data.

Two failures motivated every test here, both from real supplier files:

* ``GLOBE OUR 08 15 AUG 26.xls`` is a genuine OLE2/BIFF workbook whose real header
  sits on row 4, under a title, a period line and the issuing company. Nothing in
  the product could read it: pandas routed BIFF to xlrd (not installed), and every
  header scanner in the codebase gave up after row 3.
* ``23APR2026_HMPR_305T.xls`` is an XLSX wearing a .xls extension. Any reader that
  branches on the file name gets this one wrong in the opposite direction.

So: sniff on content, and score header rows in memory rather than re-reading the
file once per candidate.

No DB, no network. The workbook fixtures are built in memory with openpyxl; the
two real files are used only where they exist, because they carry live passenger
names and are not in the repo.

Run:  python -m unittest discover -s tests      (from backend/)
"""

import io
import os
import sys
import unittest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import pandas as pd  # noqa: E402

from app.services import spreadsheet as ss  # noqa: E402

SAMPLES = r"C:\Users\manve\Desktop\Airyield Docs"
GLOBE = os.path.join(SAMPLES, "GLOBE OUR 08 15 AUG 26.xls")
MISNAMED_XLSX = os.path.join(SAMPLES, "23APR2026_HMPR_305T.xls")

GLOBE_HEADERS = [
    "Doc Date", "DocNo", "Ticket No.", "PNR No.", "Reference", "Narration",
    "Sector / Description", "Pax Name", "AL", "Travel Date", "Basic Fare",
    "YQ Tax", "YR Tax", "K3 Tax", "OC Tax", "Other Taxes", "%%%%", "Discount",
    "Serv. Chrgs", "RAF", "GST", "TDS", "BillAmount",
]


def workbook(rows, sheets=("Sheet1",)) -> bytes:
    """An .xlsx holding `rows` on the first sheet; the rest are left empty."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = sheets[0]
    for r in rows:
        ws.append(list(r))
    for name in sheets[1:]:
        wb.create_sheet(name)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def globe_shaped() -> bytes:
    """The Globe statement's shape: three preamble rows, then the header, then data."""
    return workbook([
        ["TICKET STATEMENT"],
        ["For The Period 8 August 2026 To 15 August 2026"],
        ["GLOBE ALL INDIA SERVICES LIMITED - (WEST BENGAL)", "INR"],
        GLOBE_HEADERS,
        ["2026-08-08", "IS26/ 1067", "607 5808583279", "HNVTK1", "", "",
         "IST-AUH-DEL-   -", "MR. SAHOTA/VIKAS", "EY", "2026-08-11", "36045",
         "14741", "0", "0", "0", "4276", "0.57", "205.81", "100", "0", "18",
         "4.12", "54978"],
    ], sheets=("Sheet1", "Sheet2", "Sheet3"))


class TestSniff(unittest.TestCase):
    """The format comes from the bytes. The name is evidence, not proof."""

    def test_ole2_magic_is_a_legacy_xls(self):
        self.assertEqual(ss.sniff(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64, "x.xlsx"), ss.XLS)

    def test_zip_with_content_types_is_an_xlsx(self):
        self.assertEqual(ss.sniff(b"PK\x03\x04" + b"junk[Content_Types].xml", "x.xls"), ss.XLSX)

    def test_opendocument_zip_is_recognised_not_guessed_as_xlsx(self):
        self.assertEqual(
            ss.sniff(b"PK\x03\x04" + b"mimetypeapplication/vnd.oasis.opendocument.spreadsheet", "x.ods"),
            ss.ODS,
        )

    def test_spreadsheetml_2003_is_recognised(self):
        xml = b'<?xml version="1.0"?>\n<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet">'
        self.assertEqual(ss.sniff(xml, "report.xls"), ss.XML2003)

    def test_html_table_saved_as_xls(self):
        self.assertEqual(ss.sniff(b"<html><body><table><tr><td>1</td></tr></table>", "report.xls"), ss.HTML)

    def test_plain_text_falls_through_to_delimited(self):
        self.assertEqual(ss.sniff(b"a,b,c\n1,2,3\n", "x.csv"), ss.DELIMITED)

    def test_a_real_xlsx_sniffs_as_xlsx_whatever_it_is_called(self):
        content = workbook([["A", "B"], ["1", "2"]])
        self.assertEqual(ss.sniff(content, "definitely_a.xls"), ss.XLSX)

    def test_unsupported_formats_explain_themselves_rather_than_raising_pandas(self):
        for content, word in (
            (b"PK\x03\x04mimetypeapplication/vnd.oasis.opendocument.spreadsheet", ".xlsx"),
            (b'<?xml version="1.0"?><Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet">', "Excel 2003"),
        ):
            with self.assertRaises(ss.SpreadsheetError) as cm:
                ss.read_grid(content, "x.xls")
            self.assertIn(word, str(cm.exception))


class TestCell(unittest.TestCase):
    def test_whitespace_only_is_empty(self):
        # The Globe statement pads unused text cells to a fixed width. Without this
        # every such column reads as populated and maps as if it held data.
        self.assertEqual(ss.cell("      "), "")

    def test_null_spellings_are_empty(self):
        for v in (None, "nan", "NaN", "NaT", "null"):
            self.assertEqual(ss.cell(v), "", f"{v!r} should read as empty")

    def test_placeholders_a_person_typed_are_left_alone(self):
        # The statement importers store cells verbatim on purpose. Deciding that
        # "N/A" means absent belongs to the field reading it, not to the reader.
        for v in ("-", "--", "N/A", "NA"):
            self.assertEqual(ss.cell(v), v, f"{v!r} should survive the reader")

    def test_a_real_value_is_trimmed_not_altered(self):
        self.assertEqual(ss.cell("  MR. SAHOTA/VIKAS   "), "MR. SAHOTA/VIKAS")


class TestHeaderDetection(unittest.TestCase):

    def _grid(self, rows):
        return pd.DataFrame(rows)

    def test_finds_the_header_under_three_preamble_rows(self):
        # The case the old (0, 1, 2) scanners could not reach at all.
        result = ss.read_table(globe_shaped(), "g.xlsx")
        self.assertEqual(result.header_row, 3)
        self.assertEqual(result.columns, GLOBE_HEADERS)
        self.assertEqual(len(result.df), 1)

    def test_finds_a_header_on_row_one_when_that_is_where_it_is(self):
        content = workbook([["Ticket No", "Airline", "Fare"], ["098 12345", "AI", "100"]])
        self.assertEqual(ss.read_table(content, "t.xlsx").header_row, 0)

    def test_a_single_cell_title_row_never_wins(self):
        # A merged banner fills one cell; the header fills the row.
        content = workbook([["MONTHLY SALES REPORT"], ["Ticket No", "Airline", "Fare"],
                            ["098 1", "AI", "1"]])
        self.assertEqual(ss.read_table(content, "t.xlsx").header_row, 1)

    def test_a_two_cell_preamble_row_never_wins_either(self):
        # A real file opens "OLD | 2293898.82" above 40 headers. dropna(how="all")
        # cannot drop that row, which is why header=0 was wrong for it.
        content = workbook([["OLD", "2293898.82"],
                            ["Week", "Booked Date", "Customer Name", "Airline"],
                            ["8-15 AUG", "2026-08-08", "ACME", "TK"]])
        self.assertEqual(ss.read_table(content, "t.xlsx").header_row, 1)

    def test_a_recognised_vocabulary_outranks_everything_else(self):
        from app.services.ticket_extraction import recognised_count
        content = workbook([["a", "b", "c", "d"],
                            ["Ticket No.", "Basic Fare", "YQ Tax", "AL"],
                            ["098 1", "1", "2", "AI"]])
        self.assertEqual(ss.read_table(content, "t.xlsx", recognise=recognised_count).header_row, 1)

    def test_a_pinned_header_row_is_obeyed_not_re_detected(self):
        # The whole reason the caller sends it back: a mapping names columns, and a
        # re-detection one row out would rename all of them.
        result = ss.read_table(globe_shaped(), "g.xlsx", header_row=4)
        self.assertEqual(result.header_row, 4)
        self.assertNotEqual(result.columns, GLOBE_HEADERS)

    def test_a_pinned_row_outside_the_sheet_is_refused(self):
        with self.assertRaises(ss.SpreadsheetError):
            ss.read_table(globe_shaped(), "g.xlsx", header_row=999)

    def test_blank_and_duplicate_headers_are_kept_distinct(self):
        # Two columns called "Tax" must not collapse into one.
        content = workbook([["Tax", "", "Tax"], ["1", "2", "3"]])
        cols = ss.read_table(content, "t.xlsx").columns
        self.assertEqual(len(set(cols)), 3, f"headers collapsed: {cols}")
        self.assertIn("Tax", cols)


class TestSheetSelection(unittest.TestCase):

    def test_picks_the_sheet_with_data_not_the_first_one(self):
        # Excel leaves Sheet2/Sheet3 attached; defaulting to the first is right
        # only by accident.
        result = ss.read_table(globe_shaped(), "g.xlsx")
        self.assertEqual(result.sheet, "Sheet1")
        self.assertEqual(result.sheets, ["Sheet1", "Sheet2", "Sheet3"])

    def test_a_named_sheet_is_honoured(self):
        self.assertEqual(ss.read_table(globe_shaped(), "g.xlsx", sheet="Sheet2").sheet, "Sheet2")

    def test_an_unknown_sheet_names_the_ones_that_exist(self):
        with self.assertRaises(ss.SpreadsheetError) as cm:
            ss.read_table(globe_shaped(), "g.xlsx", sheet="Nope")
        self.assertIn("Sheet1", str(cm.exception))


class TestPreambleAndPeriod(unittest.TestCase):

    def test_preamble_is_everything_above_the_header(self):
        pre = ss.read_table(globe_shaped(), "g.xlsx").preamble
        self.assertIn("TICKET STATEMENT", pre)
        self.assertIn("GLOBE ALL INDIA SERVICES LIMITED - (WEST BENGAL)", pre)

    def test_reads_the_statement_period_for_the_valid_from_to_prefill(self):
        self.assertEqual(
            ss.parse_period(["For The Period 8 August 2026 To 15 August 2026"]),
            ("2026-08-08", "2026-08-15"),
        )

    def test_other_common_phrasings(self):
        for line, expected in (
            ("Period: 01/08/2026 to 15/08/2026", ("2026-08-01", "2026-08-15")),
            ("STATEMENT FOR THE PERIOD 01-Aug-2026 TO 15-Aug-2026", ("2026-08-01", "2026-08-15")),
            ("From 1 Aug 2026 To 15 Aug 2026", ("2026-08-01", "2026-08-15")),
        ):
            self.assertEqual(ss.parse_period([line]), expected, line)

    def test_a_period_that_does_not_fully_parse_returns_nothing(self):
        # Half a range is worse than none: the user would have to notice a wrong
        # date rather than fill in a blank one.
        for line in ("Some other title", "For The Period 8 August 2026", "GLOBE ALL INDIA"):
            self.assertIsNone(ss.parse_period([line]), line)

    def test_a_backwards_period_is_rejected(self):
        self.assertIsNone(ss.parse_period(["Period: 15/08/2026 to 01/08/2026"]))


class TestDelimited(unittest.TestCase):

    def test_csv_reads_with_its_header_detected(self):
        csv = b"TICKET STATEMENT\nTicket No,Airline,Fare\n098 1,AI,100\n"
        result = ss.read_table(csv, "s.csv")
        self.assertEqual(result.columns, ["Ticket No", "Airline", "Fare"])
        self.assertEqual(len(result.df), 1)

    def test_a_semicolon_export_needs_no_special_case(self):
        result = ss.read_table(b"Ticket No;Airline;Fare\n098 1;AI;100\n", "s.csv")
        self.assertEqual(result.columns, ["Ticket No", "Airline", "Fare"])

    def test_an_empty_file_says_so(self):
        with self.assertRaises(ss.SpreadsheetError):
            ss.read_table(b"", "empty.csv")


@unittest.skipUnless(os.path.exists(GLOBE), "real supplier file not present")
class TestRealLegacyXls(unittest.TestCase):
    """The file that could not be read at all. Skipped where it is absent —
    it carries live passenger names and is deliberately not in the repo."""

    def test_a_genuine_biff_workbook_reads(self):
        with open(GLOBE, "rb") as fh:
            content = fh.read()
        self.assertEqual(ss.sniff(content, "x.xls"), ss.XLS)
        result = ss.read_table(content, os.path.basename(GLOBE))
        self.assertEqual(result.header_row, 3)
        self.assertEqual(result.columns, GLOBE_HEADERS)
        self.assertEqual(result.sheet, "Sheet1")
        self.assertGreater(len(result.df), 30)

    def test_its_period_line_prefills_the_statement_dates(self):
        with open(GLOBE, "rb") as fh:
            result = ss.read_table(fh.read(), "g.xls")
        self.assertEqual(ss.parse_period(result.preamble), ("2026-08-08", "2026-08-15"))


@unittest.skipUnless(os.path.exists(MISNAMED_XLSX), "real supplier file not present")
class TestRealMisnamedFile(unittest.TestCase):

    def test_an_xlsx_called_xls_is_read_as_what_it_is(self):
        with open(MISNAMED_XLSX, "rb") as fh:
            content = fh.read()
        self.assertEqual(ss.sniff(content, os.path.basename(MISNAMED_XLSX)), ss.XLSX)
        result = ss.read_table(content, os.path.basename(MISNAMED_XLSX))
        self.assertIn("Ticket_No", result.columns)


if __name__ == "__main__":
    unittest.main()
