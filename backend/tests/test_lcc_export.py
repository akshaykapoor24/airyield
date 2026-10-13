"""The LCC "entries in system" export — what a user downloads to check an import landed.

The workbook is only useful if it can be laid against the uploaded file and trusted, so
the cases worth guarding are alignment (a tax amount under the wrong code is worse than no
export), the row check's verdict, the Billing column agreeing with the uploads screen, and
uploaded text never turning into a live formula.

No DB, no network — lcc_export is pure.

Run:  python -m unittest discover -s tests      (from backend/)
"""

import os
import sys
import tempfile
import unittest
from datetime import date, datetime
from decimal import Decimal

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from openpyxl import load_workbook  # noqa: E402

from app.services import lcc_detailed_spec as spec  # noqa: E402
from app.services import lcc_export as ex  # noqa: E402


class FakeRow:
    """Every LccDetailed attribute the export reads, blank unless told otherwise."""

    def __init__(self, **kw):
        self.__dict__.update({c["field"]: None for c in spec.CORE_COLUMNS})
        self.__dict__.update({
            "id": 1, "taxes": None, "segments": None, "ssr": None,
            "airline_code": "6E", "airline_name": "INDIGO", "departure_date": None,
            "row_kind": "pax", "movement_kind": None,
            "payment_method_code": "AG", "bill_status": "unresolved", "projected_ticket_id": None,
            **kw,
        })


def facts(**kw):
    return ex.UploadFacts(**{"source_file": "indigo.csv", "stored_rows": 2, "source_rows": 2, **kw})


class EntryLayoutTests(unittest.TestCase):

    def test_tax_and_leg_columns_follow_the_template_and_only_cover_what_is_present(self):
        rows = [
            FakeRow(id=1, taxes=[{"code": "UDF", "amount": 61.0}, {"code": "AE", "amount": 10.0}],
                    segments=[{"leg": 1, "route": "DELBOM", "flight_no": "2571", "dep_date": "01-Sep-2026"}]),
            FakeRow(id=2, taxes=[{"code": "PSF", "amount": 91.0}],
                    segments=[{"leg": 1, "route": "BOMDEL", "flight_no": "6012", "dep_date": None},
                              {"leg": 2, "route": "DELCCU", "flight_no": "0715", "dep_date": None}]),
        ]
        taxes, legs = ex._layout(rows)
        # Template order (AE before PSF before UDF), not first-seen order.
        self.assertEqual(taxes, ["AE", "PSF", "UDF"])
        self.assertEqual(legs, 2)

        headers = [c.header for c in ex.entry_columns(taxes, legs)]
        core = [c["header"] for c in spec.CORE_COLUMNS]
        self.assertEqual(headers[:len(core)], core)
        self.assertEqual(headers[len(core):len(core) + 3], ["AE", "PSF", "UDF"])
        self.assertIn("Second Leg Flight No", headers)
        self.assertNotIn("Third Leg", headers)
        self.assertEqual(tuple(headers[-len(ex.SYSTEM_COLUMNS):]), ex.SYSTEM_COLUMNS)

    def test_values_line_up_with_their_headers(self):
        row = FakeRow(id=42, record_locator="ABC123", total=Decimal("5097.00"),
                      taxes=[{"code": "PSF", "amount": 91.0}],
                      segments=[{"leg": 2, "route": "DELCCU", "flight_no": "0715", "dep_date": "02-Sep-2026"}])
        taxes, legs = ex._layout([row])
        cells = dict(zip([c.header for c in ex.entry_columns(taxes, legs)],
                         ex.entry_values(row, taxes, legs)))
        self.assertEqual(cells["RecordLocator"], "ABC123")
        self.assertEqual(cells["Total"], Decimal("5097.00"))
        self.assertEqual(cells["PSF"], 91.0)
        # Leg 2 lands under Second Leg even though it is the row's only segment.
        self.assertIsNone(cells["First Leg"])
        self.assertEqual(cells["Second Leg"], "DELCCU")
        self.assertEqual(cells["Second Leg Dep Date"], "02-Sep-2026")
        self.assertEqual(cells["Airline"], "6E INDIGO")
        self.assertEqual(cells["System ID"], 42)


class BillingLabelTests(unittest.TestCase):

    def test_labels_match_the_uploads_screen(self):
        self.assertEqual(ex.billing_label(FakeRow(payment_method_code="CC")), ex.BILLING_NOT_AG)
        self.assertEqual(ex.billing_label(FakeRow(payment_method_code=None)), ex.BILLING_NOT_AG)
        self.assertEqual(ex.billing_label(FakeRow(bill_status="excluded")), ex.BILLING_PAYMENT)
        self.assertEqual(ex.billing_label(FakeRow(bill_status="defaulted")), ex.BILLING_READY)
        self.assertEqual(ex.billing_label(FakeRow(bill_status="ambiguous")), ex.BILLING_NO_PARTY)
        self.assertEqual(ex.billing_label(FakeRow(bill_status="resolved", projected_ticket_id=9)),
                         ex.BILLING_SENT)
        # " ag " still counts — is_ag trims and upper-cases, and so must this.
        self.assertEqual(ex.billing_label(FakeRow(payment_method_code=" ag ")), ex.BILLING_NO_PARTY)


class RowCheckTests(unittest.TestCase):

    def test_all_stored(self):
        self.assertIn("All 32 lines", ex.row_check(facts(source_rows=32, stored_rows=32)))

    def test_lines_not_stored_are_counted(self):
        self.assertTrue(ex.row_check(facts(source_rows=205, stored_rows=203)).startswith("2 of the 205"))

    # A real Air India Express upload: an account statement plus a passenger report.
    AIX = {"account_source_rows": 48, "account_rows": 32, "skipped_not_ag": 16,
           "pax_source_rows": 12, "pax_rows": 0, "enriched_account_rows": 6,
           "duplicate_pax_lines": 0}
    AIX_FILES = [("account", "a.xlsx", 48), ("pax", "p.xlsx", 12)]

    def test_a_merge_is_not_reported_as_a_loss(self):
        msg = ex.row_check(facts(source_rows=60, stored_rows=32, merge_stats=self.AIX,
                                 files=self.AIX_FILES))
        self.assertIn("All 32 entries the import made from the 60 lines", msg)
        self.assertNotIn("not stored", msg)

    def test_a_merge_that_lost_rows_says_so(self):
        msg = ex.row_check(facts(source_rows=60, stored_rows=30, merge_stats=self.AIX,
                                 files=self.AIX_FILES))
        self.assertTrue(msg.startswith("2 of the 32 entries"))

    def test_every_merged_line_is_accounted_for(self):
        lines = dict(ex.line_breakdown(self.AIX))
        self.assertEqual(lines["Account statement — lines read"], 48)
        self.assertEqual(lines["   stored as entries (AG transactions)"], 32)
        self.assertEqual(lines["   skipped — not AG payment lines (balance lines)"], 16)
        self.assertEqual(lines["Passenger report — lines read"], 12)

    def test_passenger_report_alone_counts_its_folded_segment_lines(self):
        lines = dict(ex.line_breakdown({"pax_source_rows": 10, "pax_rows": 6, "duplicate_pax_lines": 1}))
        self.assertEqual(lines["   stored as entries (one per passenger per PNR)"], 6)
        self.assertEqual(lines["   folded into a passenger's legs (one line per segment)"], 3)
        self.assertEqual(lines["   dropped — exact duplicate lines"], 1)

    def test_a_one_for_one_merge_reads_as_all_stored(self):
        self.assertIn("All 20 lines read from the file",
                      ex.row_check(facts(source_rows=20, stored_rows=20, merge_stats={"pax_rows": 20})))

    def test_unknown_line_count(self):
        self.assertIn("predates", ex.row_check(facts(source_rows=None)))


class WorkbookTests(unittest.TestCase):

    def setUp(self):
        fd, self.path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)

    def tearDown(self):
        if os.path.exists(self.path):
            os.remove(self.path)

    def test_round_trip(self):
        rows = [
            FakeRow(id=1, name1="=HYPERLINK(\"http://x\")", total=Decimal("100.50"),
                    transaction_date=datetime(2026, 8, 15, 13, 49, 20), departure_date=date(2026, 9, 1)),
            FakeRow(id=2, name1="RAHUL SHARMA", total=Decimal("-25.25"), pax_count=2),
        ]
        ex.write_workbook(self.path, rows, facts(filters=[("Passenger", "RAHUL")]))

        wb = load_workbook(self.path)
        self.assertEqual(wb.sheetnames, [ex.ENTRIES_TITLE, ex.SUMMARY_TITLE])
        ws = wb[ex.ENTRIES_TITLE]
        self.assertEqual(ws.max_row, 3)                          # header + 2 entries
        headers = [c.value for c in ws[1]]
        first = dict(zip(headers, [c.value for c in ws[2]]))
        # Uploaded text is written as text, never as a live formula.
        self.assertEqual(ws.cell(row=2, column=headers.index("Name1") + 1).data_type, "s")
        self.assertEqual(first["Name1"], "=HYPERLINK(\"http://x\")")
        # Seconds survive — they are what a timestamp is matched on.
        self.assertEqual(first["Transaction Date"], datetime(2026, 8, 15, 13, 49, 20))
        self.assertIn("ss", ws.cell(row=2, column=headers.index("Transaction Date") + 1).number_format)
        self.assertEqual(first["Total"], 100.5)

        summary = [[c for c in r if c is not None] for r in wb[ex.SUMMARY_TITLE].iter_rows(values_only=True)]
        flat = {r[0]: r[1] for r in summary if len(r) == 2}
        self.assertEqual(flat["Entries stored in the system"], 2)
        self.assertEqual(flat["Entries in this export (filtered)"], 2)
        self.assertEqual(flat["Passenger"], "RAHUL")
        self.assertEqual(flat["Total"], 75.25)
        self.assertEqual(flat["Pax"], 2)

    def test_no_matching_rows_still_writes_the_header(self):
        ex.write_workbook(self.path, [], facts(stored_rows=0, source_rows=0))
        ws = load_workbook(self.path)[ex.ENTRIES_TITLE]
        self.assertEqual(ws.max_row, 1)
        self.assertEqual(ws.cell(row=1, column=1).value, spec.CORE_COLUMNS[0]["header"])


class FileNameTests(unittest.TestCase):

    def test_extension_dropped_and_unsafe_characters_replaced(self):
        self.assertEqual(ex.file_name("indigolccStat.csv"), "indigolccStat - entries in system.xlsx")
        self.assertEqual(ex.file_name("a/b:c.XLSX"), "a_b_c - entries in system.xlsx")
        self.assertEqual(ex.file_name(None), "upload - entries in system.xlsx")


if __name__ == "__main__":
    unittest.main()
