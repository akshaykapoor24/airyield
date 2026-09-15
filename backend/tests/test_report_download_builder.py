"""Report download — the builder wired end to end on a fake query layer (no DB).

One small workspace exercises every cross-source rule the builder is responsible for:

* one BSP statement (issue-date scope): a TKTT enriched by a TGQ ticket, a TKTT issued before
  the period (indexed for linking, not written), an RFND and an ADMA row;
* a TGQ HMPR upload: the enriched ticket (suppressed — BSP already bills it), a ticket BSP
  does not hold (Not In BSP = Yes after the owner lookup finds nothing) and the ticket of the
  out-of-period BSP row ("BSP row outside period"); plus an UNTICKED older TGQ upload that
  must be kept out of the enrichment index;
* an ADM matching the ADMA row ("Yes – this report", never counted);
* the same Third Party GDS statement uploaded twice (older copy superseded, not counted);
* an LCC Detailed row and an LCC Flown row on the same PNR (the ledger links to Detailed).

Asserted: Combined rows per source, suppression, Counts In Net, TGQ enrichment on the BSP row,
the tie-out block, sheet order after save, Read Me files included / excluded, the arguments
the TGQ loader and owner lookup received, cancellation, the deadline, a not-completed
statement being left out and a re-processed upload being noted.

Run:  ..\\venv\\Scripts\\python.exe -m unittest test_report_download_builder -v   (from backend/tests)
"""
import asyncio
import os
import shutil
import sys
import tempfile
import time
import unittest
from collections import Counter
from datetime import date, datetime
from decimal import Decimal
from types import SimpleNamespace

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import app.models  # noqa: F401,E402
from openpyxl import load_workbook  # noqa: E402

from app.models.airline_adjustment import AirlineADM  # noqa: E402
from app.models.bsp_statement import BspStatementRow  # noqa: E402
from app.models.lcc_detailed import LccDetailed  # noqa: E402
from app.services import sector_split  # noqa: E402
from app.services import statement_spec as spec  # noqa: E402
from app.services.bsp_tgq_enrichment import TgqIndex, TgqTicket  # noqa: E402
from app.services.report_download import builder as B  # noqa: E402
from app.services.report_download import columns as C  # noqa: E402
from app.services.report_download.selection import OwnedSelection, OwnedUpload  # noqa: E402
from app.services.report_download.summary import T_TIE_OUT  # noqa: E402
from app.services.report_download.types import (  # noqa: E402
    AirlineInfo, AirlineMaster, AirlineSnap, Period, ReportMeta, ReportOptions, SupplierSnap, TaxComponent,
)

PERIOD = Period(date(2026, 7, 1), date(2026, 7, 15))
D = Decimal


# ── fixture rows ──────────────────────────────────────────────────────────────

def bsp_row(**kw):
    base = {c.name: None for c in BspStatementRow.__table__.columns}
    base.update(statement_id="stmt-1", airline_accounting_code="098", tenant_id=1, created_by_id=2)
    base.update(kw)
    return SimpleNamespace(**base)


BSP_ROWS = [
    bsp_row(id=1, transaction_type="TKTT", document_number="5805708071", ticket_number="5805708071",
            issue_date=date(2026, 7, 3), transaction_amount=D("11800"), fare_amount=D("10000"),
            balance_payable=D("11700"), standard_commission_amount=D("100"), form_of_payment="CA",
            raw_data={"section": "ISSUES"}),
    bsp_row(id=2, transaction_type="TKTT", document_number="5805708072", ticket_number="5805708072",
            issue_date=date(2026, 6, 20), transaction_amount=D("5000"), fare_amount=D("5000"),
            balance_payable=D("5000"), raw_data={"section": "ISSUES"}),
    bsp_row(id=3, transaction_type="RFND", document_number="6000000001", rtdn="5805708073",
            issue_date=date(2026, 7, 5), transaction_amount=D("-4000"), fare_amount=D("-4000"),
            balance_payable=D("-4000"), raw_data={"section": "REFUNDS"}),
    bsp_row(id=4, transaction_type="ADMA", document_number="8000000001", ticket_number="8000000001",
            issue_date=date(2026, 7, 6), transaction_amount=D("250"), balance_payable=D("250"),
            raw_data={"section": "DEBIT MEMOS"}),
]
TAXES = {1: [TaxComponent("TAX", "YQ", D("500")), TaxComponent("TAX", "K3", D("500")),
             TaxComponent("TAX", "IN", D("800"))]}

BSP_HEADER = SimpleNamespace(
    batch_id="stmt-1", statement_name="BSP 1-15 Jul", file_name="bsp.pdf", period_from=date(2026, 7, 1),
    period_to=date(2026, 7, 15), group_id="g1", status="completed", created_at=datetime(2026, 7, 16, 8, 0),
    gt_issues=D("16800"), gt_refunds=D("-4000"), gt_debit_memos=D("250"), gt_credit_memos=D("0"),
    gt_std_comm=D("100"), gt_sup_comm=None, gt_tax_on_comm=None, gt_balance_payable=D("12950"), gt_doc_count=4,
)
SUMMARY_HEADER = SimpleNamespace(
    batch_id="sum-1", group_id="g1", agent_code="14312345", agent_name="ACME TRAVELS", currency="INR",
    billing_period_code="2026071", period_from=date(2026, 7, 1), period_to=date(2026, 7, 15),
    file_name="summary.pdf", status="completed", match_status="matched", match_detail=None,
)

_TGQ_SPLIT = spec.split_config("tgq-hmpr")
_TGQ_TICKET_NO = spec.ticket_no_config("tgq-hmpr")


def tgq_legs(seq, ticket_no, ticket_date, *, batch="tgq-1", start_id=100, when=date(2026, 7, 3)):
    source = {
        "ticket_no": ticket_no, "transaction_type": "SALE", "ticket_date": ticket_date, "airline": "AI",
        "air_name": "AIR INDIA", "pax_name": f"PAX/{seq} MR", "air_pnr": f"PNR{seq}", "sectors": "BOM/DEL",
        "flightno": "AI-101", "traveldt": "20JUL", "class": "Y", "coupon_status": "OPEN",
        "base_fare": "10000", "total_tax": "1800", "total_fare": "11800", "net_remit": "11700",
        "total_refund_amount": "0", "fop": "CA",
    }
    data, _derived = sector_split.apply_ticket_no(source, _TGQ_TICKET_NO)
    legs = []
    for i, (leg_data, leg_taxes, idx, count, status) in enumerate(sector_split.split_row(data, [], _TGQ_SPLIT)):
        legs.append(SimpleNamespace(
            id=start_id + seq * 10 + i, batch_id=batch, source_file=f"{batch}.xlsx",
            uploaded_at=datetime(2026, 7, 10), data=leg_data, taxes=leg_taxes, row_seq=seq,
            sector_index=idx, sector_count=count, split_status=status, is_total=False,
            orig_data=None, orig_taxes=None, fake_date=when,
        ))
    return legs


TGQ_ROWS = (tgq_legs(1, "098 5805708071", "03JUL26")          # enriches BSP row 1 → suppressed
            + tgq_legs(2, "098 5805708099", "04JUL26")        # nowhere in BSP
            + tgq_legs(3, "098 5805708072", "05JUL26"))       # BSP row 2, issued before the period

ENRICHED = TgqTicket(sector="BOM/DEL", booking_class="Y", travel_raw="20JUL", ticket_date_raw="03JUL26",
                     leg_count=1, batch_id="tgq-1", pax_name="SHARMA/RAHUL MR", air_pnr="ABC123",
                     flight_numbers="AI-101", fare_basis="YIN")


def adm_row():
    base = {c.name: None for c in AirlineADM.__table__.columns}
    base.update(id=7, batch_id="adm-1", source_file="adm.xlsx", uploaded_at=datetime(2026, 7, 11),
                airline_code="098", agent_code="14312345", document_number="8000000001", currency="INR",
                status="Billed", type="ADM", issue_date="06/07/2026", amount="250", period="2026071",
                fake_date=date(2026, 7, 6))
    return SimpleNamespace(**base)


def gds_row(rid, batch):
    data = {"ticket_status": "CONFIRMED", "ticket_prefix": "98", "ticket_number": "5805700001",
            "issue_date": "2026-07-07", "passenger_name": "RAHUL SHARMA", "pnr": "S1PNR",
            "base_fare": "900", "other_taxes": "100", "total_fare": "1000", "net_amount": "1000"}
    return SimpleNamespace(id=rid, batch_id=batch, source_file=f"{batch}.xlsx", uploaded_at=datetime(2026, 7, 8),
                           data=data, taxes=None, segments=None, ssr=None, source_format="third-party-gds-v2",
                           fake_date=date(2026, 7, 7))


def lcc_row():
    base = {c.name: None for c in LccDetailed.__table__.columns}
    base.update(id=501, batch_id="lcc-1", airline_code="6E", airline_name="IndiGo",
                transaction_date=datetime(2026, 7, 4, 11, 0), name1="SHARMA/RAHUL MR", total=D("5000.00"),
                base_fare=D("4000.00"), currency_code="INR", record_locator="ABC123", bill_kind="sale",
                fake_date=date(2026, 7, 4))
    return SimpleNamespace(**base)


def flown_row():
    return SimpleNamespace(id=55, batch_id="flown-1", source_file="flown.xlsx", uploaded_at=datetime(2026, 7, 12),
                           data={"pnr": "abc123", "flown_date": "05-Jul-2026", "passenger_name": "SHARMA/RAHUL",
                                 "total_fare": "5000", "net_fare": "4800", "flight_number": "6E-201"},
                           taxes=None, segments=None, ssr=None, source_format="std", fake_date=date(2026, 7, 5))


# ── fake query layer ──────────────────────────────────────────────────────────

def _up(source, uid, *, total, at, status=None):
    return OwnedUpload(source_key=source, upload_id=uid, file_name=f"{uid}.xlsx", uploaded_at=at,
                       status=status, total_rows=total)


class FakeQueries:
    UNDATED_ATTR = "rd_undated"
    IN_PERIOD_ATTR = "rd_in_period"
    STREAM_CHUNK = 5000

    def __init__(self):
        self.rows = {
            ("bsp", "stmt-1"): BSP_ROWS,
            ("adm", "adm-1"): [adm_row()],
            ("tgq-hmpr", "tgq-1"): TGQ_ROWS,
            ("tp-gds", "gds-new"): [gds_row(900, "gds-new")],
            ("tp-gds", "gds-old"): [gds_row(800, "gds-old")],
            ("lcc-detailed", "lcc-1"): [lcc_row()],
            ("lcc-flown-report", "flown-1"): [flown_row()],
        }
        self.status = {("bsp", "stmt-1"): "completed", ("lcc-detailed", "lcc-1"): "completed"}
        self.fingerprinted = set()
        self.change_on_second_fingerprint = None
        self.owner_calls = []
        self.tgq_kwargs = None

    async def upload_fingerprints(self, conn, t, u, key, ids):
        out = {}
        for uid in ids:
            rows = self.rows.get((key, uid))
            if rows is None:
                continue
            count = len(rows)
            if self.change_on_second_fingerprint == (key, uid) and (key, uid) in self.fingerprinted:
                count += 1
            self.fingerprinted.add((key, uid))
            out[uid] = (self.status.get((key, uid)), count, max(r.id for r in rows))
        return out

    async def airline_master(self, conn):
        info = AirlineInfo("AI", "098", "Air India")
        return AirlineMaster(by_numeric={"098": info}, by_iata={"AI": info})

    async def bsp_statement_headers(self, conn, t, u, ids):
        return {"stmt-1": BSP_HEADER} if "stmt-1" in ids else {}

    async def bsp_summaries_by_group(self, conn, t, u, groups):
        return {"g1": SUMMARY_HEADER} if "g1" in groups else {}

    async def bsp_summary_statement_headers(self, conn, t, u, ids):
        return {}

    async def lcc_batch_headers(self, conn, t, u, ids):
        return {"lcc-1": SimpleNamespace(batch_id="lcc-1", source_file="lcc-1.xlsx", source_format="indigo",
                                         airline_name="IndiGo", airline_code="6E", status="completed")}

    async def lcc_batch_airlines(self, conn, t, u, ids):
        return {"lcc-1": AirlineSnap(5, 50, "IndiGo", "6E", "312", "KT1")}

    async def batch_airlines(self, conn, t, u, slug, ids):
        return {i: AirlineSnap(5, 50, "IndiGo", "6E", "312", "KT1") for i in ids}

    async def supplier_snapshots(self, conn, t, u, slug, ids):
        return {i: SupplierSnap(42, "Riya Travel", "RIYA", "Mumbai") for i in ids}

    async def tgq_file_names(self, conn, t, u):
        return {"tgq-1": "hmpr-july.xlsx", "tgq-old": "hmpr-june.xlsx"}

    async def load_tgq_index(self, conn, t, u, **kwargs):
        self.tgq_kwargs = kwargs
        index = TgqIndex({("098", "5805708071"): ENRICHED}, {}, by_serial={})
        index.rows_loaded, index.batches_used = 1, 1
        return index

    def _bsp_in_period(self, row, options):
        if options.basis != "transaction" or options.bsp_scope != "issue_date":
            return True
        d = row.issue_date or BSP_HEADER.period_from
        return PERIOD.contains(d)

    async def stream_bsp_keys(self, conn, t, u, statement_id, *, header, period, options, chunk=5000):
        rows = [SimpleNamespace(**vars(r), in_period=self._bsp_in_period(r, options))
                for r in self.rows[("bsp", statement_id)]]
        yield rows

    async def stream_rows(self, conn, t, u, key, upload_id, *, period, options, header=None, columns=None,
                          filter_period=True, flag_period=False, chunk=5000):
        out = []
        for r in self.rows.get((key, upload_id), ()):
            if key == "bsp":
                keep, undated = self._bsp_in_period(r, options), False
            else:
                d = getattr(r, "fake_date", None)
                undated = d is None
                keep = options.basis != "transaction" or period.contains(d) or (
                    undated and options.undated_rows == "include")
            if filter_period and not flag_period and not keep:
                continue
            out.append(SimpleNamespace(**vars(r), rd_undated=undated, rd_in_period=keep))
        yield out

    async def bsp_tax_components(self, conn, t, u, ids):
        return {i: TAXES[i] for i in ids if i in TAXES}

    async def bsp_tax_codes(self, conn, t, u, ids):
        return ["YQ", "K3", "IN"]

    async def lcc_tax_codes(self, conn, t, u, ids):
        return []

    async def lcc_extra_keys(self, conn, t, u, ids, include_pii):
        return ()

    async def distinct_data_keys(self, conn, t, u, key, ids, include_pii):
        return ()

    async def bsp_owner_lookup(self, conn, t, u, excluded, forms):
        self.owner_calls.append((list(excluded), list(forms)))
        return []


def selection():
    return OwnedSelection(tenant_id=1, user_id=2, uploads=(
        _up("bsp", "stmt-1", total=4, at=datetime(2026, 7, 16), status="completed"),
        _up("adm", "adm-1", total=1, at=datetime(2026, 7, 11)),
        _up("tgq-hmpr", "tgq-1", total=len(TGQ_ROWS), at=datetime(2026, 7, 10)),
        _up("tp-gds", "gds-old", total=1, at=datetime(2026, 7, 8)),
        _up("tp-gds", "gds-new", total=1, at=datetime(2026, 7, 12)),
        _up("lcc-detailed", "lcc-1", total=1, at=datetime(2026, 7, 9), status="completed"),
        _up("lcc-flown-report", "flown-1", total=1, at=datetime(2026, 7, 12)),
    ), unticked=(_up("tgq-hmpr", "tgq-old", total=5, at=datetime(2026, 6, 1)),))


def meta(options):
    return ReportMeta(export_id=9, title="July", generated_at=datetime(2026, 7, 20, 6, 30),
                      generated_by_name="Nitin", generated_by_email="n@example.com", tenant_name="ACME",
                      period=PERIOD, options=options,
                      source_types=("bsp", "adm", "tgq-hmpr", "tp-gds", "lcc-detailed", "lcc-flown-report"))


PARAMS = {"date_from": "2026-07-01", "date_to": "2026-07-15", "basis": "transaction", "bsp_scope": "issue_date"}


class BuilderTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="rd-builder-")
        self.path = os.path.join(self.tmp, "report.xlsx")

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def build(self, q=None, options=None, state=None, sel=None):
        q = q or FakeQueries()
        options = options or ReportOptions(bsp_scope="issue_date")
        state = state or B.BuildState()
        result = asyncio.run(B.build_report(None, sel or selection(), PARAMS, meta(options), self.path, state, q=q))
        return q, state, result

    def combined(self):
        wb = load_workbook(self.path, read_only=True)
        rows = list(wb["Combined"].iter_rows(values_only=True))
        names = wb.sheetnames
        wb.close()
        keys = C.COMBINED_KEYS
        return names, [dict(zip(keys, tuple(r) + (None,) * (len(keys) - len(r)))) for r in rows[1:]]

    def test_full_build(self):
        q, state, result = self.build()
        result.workbook.save()
        names, rows = self.combined()

        self.assertEqual(names, ["Read Me", "Summary", "Combined", "BSP Detailed", "ADM", "TGQ HMPR",
                                 "LCC Detailed", "LCC Flown Report", "TP GDS"])
        by_source = Counter(r["source_type"] for r in rows)
        self.assertEqual(by_source, Counter({"BSP": 3, "TGQ HMPR": 2, "ADM": 1, "Third Party GDS": 2,
                                             "LCC Detailed Statement": 1, "LCC Flown Report": 1}))
        self.assertEqual(result.combined_rows, 10)
        self.assertEqual(result.sheet_row_counts["BSP Detailed"], 3)       # row 2 is outside the period
        self.assertEqual(result.sheet_row_counts["TGQ HMPR"], len(TGQ_ROWS))

        bsp = {r["document_number"]: r for r in rows if r["source_type"] == "BSP"}
        self.assertEqual(set(bsp), {"5805708071", "6000000001", "8000000001"})
        self.assertTrue(all(r["counts_in_net"] == "Yes" for r in bsp.values()))
        enriched = bsp["5805708071"]
        self.assertEqual(enriched["tgq_enriched"], "Yes (document)")
        self.assertEqual(enriched["passenger_name"], "SHARMA/RAHUL MR")
        self.assertEqual(enriched["linked_via"], "TGQ HMPR document (hmpr-july.xlsx)")
        self.assertEqual(enriched["yq"], 500)
        self.assertIn("MEMO_UPLOADED", bsp["8000000001"]["data_flags"])     # the ADM is uploaded too

        tgq = {r["document_number"]: r for r in rows if r["source_type"] == "TGQ HMPR"}
        self.assertNotIn("5805708071", tgq)                                  # suppressed: BSP bills it
        self.assertEqual(tgq["5805708099"]["not_in_bsp"], "Yes")
        self.assertEqual(tgq["5805708099"]["counts_in_net"], C.NET_TGQ_NOT_IN_BSP)
        self.assertEqual(tgq["5805708072"]["not_in_bsp"], "BSP row outside period (stmt-1.xlsx)")
        self.assertEqual(tgq["5805708072"]["counts_in_net"], C.NET_TGQ_OUTSIDE_PERIOD)
        self.assertEqual(result.summary["link_stats"]["TGQ tickets suppressed (in BSP)"], 1)
        self.assertEqual(result.summary["link_stats"]["BSP rows TGQ-enriched (document)"], 1)

        adm = next(r for r in rows if r["source_type"] == "ADM")
        self.assertEqual(adm["also_in_bsp"], "Yes – this report")
        self.assertEqual(adm["counts_in_net"], C.NET_MEMO_IN_BILLING)

        gds = sorted((r for r in rows if r["source_type"] == "Third Party GDS"), key=lambda r: r["upload_id"])
        self.assertEqual([(g["upload_id"], g["counts_in_net"]) for g in gds],
                         [("gds-new", "Yes"), ("gds-old", C.NET_SUPERSEDED)])
        self.assertIn("DUPLICATE_SUPERSEDED", gds[1]["data_flags"])
        self.assertNotIn("DUPLICATE_SUPERSEDED", gds[0]["data_flags"] or "")

        lcc = next(r for r in rows if r["source_type"] == "LCC Detailed Statement")
        self.assertEqual(lcc["counts_in_net"], "Yes")
        flown = next(r for r in rows if r["source_type"] == "LCC Flown Report")
        self.assertEqual(flown["counts_in_net"], C.NET_LEDGER)
        self.assertEqual((flown["linked_document"], flown["linked_via"]), ("ABC123", "LCC Detailed PNR"))
        self.assertIn("ALSO_IN_LCC_DETAILED", flown["data_flags"])

        # counted net: BSP 11700 − 4000 + 250, GDS 1000 (newest copy), LCC 5000
        self.assertEqual(Decimal(result.summary["net_by_currency"]["INR"]), Decimal("13950"))
        self.assertEqual(result.summary["counted_rows"], 5)

        # TGQ loader: report mode, the unticked TGQ upload kept out; owner lookup excludes the
        # included statement and was asked about the one ticket BSP does not hold.
        self.assertEqual(q.tgq_kwargs["exclude_batch_ids"], frozenset({"tgq-old"}))
        self.assertTrue(q.tgq_kwargs["report_mode"])
        self.assertTrue(q.owner_calls)
        excluded, forms = q.owner_calls[0]
        self.assertEqual(excluded, ["stmt-1"])
        self.assertIn("5805708099", forms)
        self.assertNotIn("5805708072", forms)

        # Read Me: every included upload, the unticked one excluded with its reason
        included = {(f["source_type"], f["upload_id"]): f for f in result.files_included}
        self.assertEqual(len(included), 7)
        bsp_file = included[("bsp", "stmt-1")]
        self.assertEqual((bsp_file["rows_in_file"], bsp_file["rows_included"]), (4, 3))
        self.assertEqual(bsp_file["reference"], "ACME TRAVELS")
        self.assertIn("1 rows issued outside the period", bsp_file["notes"][0])
        self.assertEqual(included[("tp-gds", "gds-old")]["rows_superseded"], 1)
        self.assertIn("3 tickets: 2 listed, 1 in BSP", included[("tgq-hmpr", "tgq-1")]["notes"][0])
        self.assertEqual([(f["upload_id"], f["reason"]) for f in result.files_excluded], [("tgq-old", "unticked")])

        # Summary: the tie-out block is written (partial scope → OK n/a); progress advanced
        wb = load_workbook(self.path, read_only=True)
        summary_titles = [r[0] for r in wb["Summary"].iter_rows(values_only=True) if r and r[0]]
        readme_cells = [c for r in wb["Read Me"].iter_rows(values_only=True) for c in r if c]
        wb.close()
        self.assertIn(T_TIE_OUT, summary_titles)
        self.assertIn("tgq-old.xlsx", readme_cells)
        self.assertIn("Unticked by you", readme_cells)
        self.assertGreater(state.processed, 0)
        self.assertEqual(state.stage, "Writing Summary and Read Me")

    def test_whole_statement_writes_every_bsp_row_and_ties_out(self):
        q, _state, result = self.build(options=ReportOptions(bsp_scope="whole_statement", include_detail_sheets=False))
        result.workbook.save()
        names, rows = self.combined()
        self.assertEqual(names, ["Read Me", "Summary", "Combined"])
        self.assertEqual(Counter(r["source_type"] for r in rows)["BSP"], 4)
        # every BSP row is in period now, so the TGQ ticket of row 2 is suppressed too
        self.assertEqual({r["document_number"] for r in rows if r["source_type"] == "TGQ HMPR"}, {"5805708099"})
        self.assertIs(result.summary["tie_out_ok"], True)

    def test_cancelled_state_raises_and_releases_the_workbook(self):
        state = B.BuildState(cancelled=True)
        with self.assertRaises(B.ReportCancelled):
            self.build(state=state)
        self.assertFalse(os.path.exists(self.path))

    def test_deadline_raises_timeout(self):
        with self.assertRaises(B.ReportTimeout):
            self.build(state=B.BuildState(deadline=time.monotonic() - 1))

    def test_not_completed_statement_is_left_out(self):
        q = FakeQueries()
        q.status[("bsp", "stmt-1")] = "processing"
        _q, _state, result = self.build(q=q)
        self.assertIn(("stmt-1", "not_completed"), [(f["upload_id"], f["reason"]) for f in result.files_excluded])
        self.assertIsNone(q.tgq_kwargs)                   # no BSP left, so no enrichment index
        result.workbook.discard()

    def test_upload_changed_during_build_is_noted(self):
        q = FakeQueries()
        q.change_on_second_fingerprint = ("adm", "adm-1")
        _q, _state, result = self.build(q=q)
        notes = next(f["notes"] for f in result.files_included if f["upload_id"] == "adm-1")
        self.assertIn(B.NOTE_CHANGED, notes)
        result.workbook.discard()

    def test_owner_lookup_batches_do_not_change_the_rows(self):
        def rows_of(result):
            result.workbook.save()
            _names, rows = self.combined()
            return sorted((r["source_type"], r["document_number"] or "", r["counts_in_net"], r["not_in_bsp"] or "",
                           r["also_in_bsp"] or "") for r in rows)

        q_one, _s, first = self.build()
        baseline = rows_of(first)
        old = B.PENDING_PROBES_FLUSH
        B.PENDING_PROBES_FLUSH = 1                       # flush after every pending ticket / memo
        second = None
        try:
            q_many, _s, second = self.build(sel=OwnedSelection(
                tenant_id=1, user_id=2,
                uploads=tuple(u for u in selection().uploads if u.source_key != "bsp"),
                unticked=selection().unticked))
            self.assertEqual(len(q_many.owner_calls), 4)   # the ADM, then each of the three TGQ tickets
            rows = rows_of(second)
        finally:
            B.PENDING_PROBES_FLUSH = old
            if second is not None:
                second.workbook.discard()
        self.assertEqual(len(q_one.owner_calls), 1)
        # Without the statement: the ADM and the TGQ ticket of row 2 are no longer linked to it.
        self.assertEqual([r for r in rows if r[0] == "ADM"], [("ADM", "8000000001", "No – not yet billed in BSP (status: Billed)", "", "No")])
        self.assertEqual(sorted(r[1] for r in rows if r[0] == "TGQ HMPR"), ["5805708071", "5805708072", "5805708099"])
        self.assertTrue(any(r[0] == "BSP" for r in baseline))

    def test_tick_updates_progress(self):
        state = B.BuildState()
        state.tick(5, "Writing BSP")
        state.tick(3)
        self.assertEqual((state.processed, state.stage), (8, "Writing BSP"))


@unittest.skipUnless(os.environ.get("REPORT_DB_TESTS") == "1", "set REPORT_DB_TESTS=1 to build a real workbook from the local DB")
class DatabaseSmokeTest(unittest.TestCase):
    """READ-ONLY: list, select and build for the local user with the most BSP rows (else any
    upload), inside a READ ONLY transaction that is rolled back; prints the sheet counts."""

    def test_build_for_first_local_user(self):
        from sqlalchemy import text
        from sqlalchemy.ext.asyncio import create_async_engine
        from sqlalchemy.pool import NullPool
        from app.config import Settings, settings
        from app.services.report_download import queries, selection as sel_mod
        from app.services.report_download.registry import source_keys

        env_file = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".env"))
        url = Settings(_env_file=env_file).DATABASE_URL if os.path.exists(env_file) else settings.DATABASE_URL
        tmp = tempfile.mkdtemp(prefix="rd-db-smoke-")

        async def go():
            engine = create_async_engine(url, poolclass=NullPool)
            try:
                async with engine.connect() as conn:
                    await conn.execute(text("SET TRANSACTION READ ONLY"))
                    owner = (await conn.execute(text(
                        "SELECT tenant_id, created_by_id FROM ("
                        " SELECT tenant_id, created_by_id FROM bsp_statement_rows UNION ALL"
                        " SELECT tenant_id, created_by_id FROM lcc_detailed UNION ALL"
                        " SELECT tenant_id, created_by_id FROM tgq_hmpr) x"
                        " GROUP BY 1, 2 ORDER BY count(*) DESC LIMIT 1"))).first()
                    if owner is None:
                        self.skipTest("no uploads in the local database")
                    period = Period(date(2020, 1, 1), date(2030, 12, 31))
                    listing = await queries.list_matching_uploads(
                        conn, owner[0], owner[1], types=list(source_keys()), date_from=period.date_from,
                        date_to=period.date_to, basis="transaction", bsp_scope="whole_statement")
                    picked = [(u["source_type"], u["upload_id"]) for u in listing["uploads"]
                              if u["selectable"] and u["default_selected"]]
                    if not picked:
                        self.skipTest("nothing selectable")
                    sel = await sel_mod.resolve_selection(conn, owner[0], owner[1], picked)
                    options = ReportOptions()
                    result = await B.build_report(
                        conn, sel, {"date_from": "2020-01-01", "date_to": "2030-12-31"},
                        ReportMeta(None, "DB smoke", datetime(2026, 1, 1), None, None, None, period, options,
                                   tuple(source_keys())),
                        os.path.join(tmp, "smoke.xlsx"), B.BuildState())
                    await conn.rollback()
                    return result
            finally:
                await engine.dispose()

        try:
            result = asyncio.run(go())
            result.workbook.save()
            print("\nDB smoke sheet counts:", result.sheet_row_counts, "summary:", result.summary.get("rows"),
                  "tie_out_ok:", result.summary.get("tie_out_ok"))
            self.assertEqual(result.combined_rows, result.sheet_row_counts.get("Combined", 0))
            wb = load_workbook(os.path.join(tmp, "smoke.xlsx"), read_only=True)
            self.assertEqual(wb.sheetnames[:2], ["Read Me", "Summary"])
            wb.close()
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
