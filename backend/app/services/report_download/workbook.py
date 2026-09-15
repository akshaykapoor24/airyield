"""Streaming .xlsx writer for the report: lazy sheets, rank ordering, row-limit splits, cell safety.

WHY WRITE-ONLY. A report can carry ~750k Combined rows plus the same again across detail
sheets. A normal openpyxl workbook keeps every cell object in memory; write-only mode
streams each appended row to a per-sheet temp file, so memory stays flat. The price is a
set of ordering rules, all verified against openpyxl 3.1.5 (no lxml in this venv, so
et_xmlfile is the XML writer):

* Interleaving is fine. Every ``WriteOnlyWorksheet`` owns its own ``WorksheetWriter``, temp
  file and row generator (``worksheet/_write_only.py``), so rows may be appended to sheet A,
  then B, then A again, and a sheet may be created after other sheets already hold rows.
* Column widths and freeze panes are written by ``WorksheetWriter.write_top()``, which runs on
  the sheet's FIRST append (``_get_writer``). They are set before the header row and can
  never change afterwards. ``auto_filter`` is written by ``write_tail()`` when the sheet is
  closed, so its ref is set at ``save()`` — or when a full split part is closed early.
* Sheet order is ``wb._sheets`` order at save time: ``ExcelWriter._write_worksheets`` numbers
  ``sheetN.xml`` from it and ``WorkbookWriter.write_worksheets`` lists the same order, so
  sorting the list in place just before saving is enough. That is how Summary can be
  written last yet sit second.
* A string that starts with "=" becomes a formula (``Cell._bind_value`` sets data_type "f")
  and one equal to an error literal becomes an error cell ("e"). Forcing ``data_type = "s"``
  on a fresh cell afterwards writes it as an inline string (``cell/_writer.py``: "s" →
  ``t="inlineStr"``), which Excel shows as text and never evaluates — this blocks formula
  injection from uploaded file content. Other strings ("-285.72") are already plain text.
* ``_values_to_row`` reuses ONE cell for plain values and adopts any ``Cell`` passed in,
  overwriting its row/column and — when it carries no style — reusing it for the next
  value. A cell handed to ``append`` is therefore never reused by us; styled, date and
  forced-string values each get a new ``Cell``. To keep that cheap at this volume, each
  style (header font+fill, money / date / datetime number formats, and openpyxl's own
  default date formats) is resolved to its ``StyleArray`` of workbook style-table indices
  once, and new cells copy that array (``Cell(..., style_array=...)``) instead of
  re-resolving Font/Fill/format per cell. Text and unformatted numbers stay plain Python
  values and never become cell objects here.
* Throughput without lxml is bounded by openpyxl's pure-Python XML serialisation
  (``etree_write_cell`` + et_xmlfile), not by this module: measured ~65-70k non-blank
  cells/s on a dev laptop for a 62-column Combined row.
* ``ILLEGAL_CHARACTERS_RE`` (cell.py) only rejects C0 controls. Lone surrogates are written
  as ``&#55296;``-style references (the temp file is opened with
  ``errors="xmlcharrefreplace"``) and U+FFFE/U+FFFF pass through — both make the package
  unreadable, so ``excel_value`` strips every XML-1.0-illegal code point first.
* Temp files: each sheet writer creates ``openpyxl.*`` in ``tempfile.gettempdir()``
  (``NamedTemporaryFile(delete=False)``, registered for atexit removal). ``save()`` deletes
  each one as it is zipped; ``discard()`` closes the row generator and the XML stream
  (which closes the file handle — required before deletion on Windows) and deletes the rest.
  A worker that wants them in a per-export directory points ``tempfile.tempdir`` there.

The row limit is Excel's 1,048,576 rows per sheet; one is the header, so a sheet takes
``MAX_DATA_ROWS`` data rows and the next row opens "<title> (2)" with the header repeated.
"""
from __future__ import annotations

import json
import math
import numbers
import os
import re
from contextlib import suppress
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal
from typing import Any, Optional, Sequence
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.cell.cell import ERROR_CODES, TIME_FORMATS, Cell, WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.cell_style import StyleArray
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._write_only import WriteOnlyWorksheet
from openpyxl.writer.excel import ExcelWriter

MAX_DATA_ROWS = 1_048_575          # Excel's 1,048,576 rows minus the header row
MAX_COLUMNS = 16_384
MAX_CELL_CHARS = 32_767
MAX_SHEET_TITLE = 31
MAX_COLUMN_WIDTH = 255
NO_DATA_TITLE = "No data"

COLUMN_KINDS = frozenset({"text", "int", "date", "datetime", "money", "flag"})
NUMBER_FORMATS: dict[str, str] = {
    "money": "#,##0.00",
    "date": "DD-MMM-YYYY",
    "datetime": "DD-MMM-YYYY HH:MM",
}
HEADER_FILL_RGB = "1E3A5F"

# Every code point XML 1.0 forbids: C0 controls except TAB/LF/CR, surrogates, U+FFFE/U+FFFF.
_ILLEGAL_XML_RE = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_ERROR_LITERALS = frozenset(ERROR_CODES)            # "#N/A", "#VALUE!", …
_INVALID_TITLE_RE = re.compile(r"[\[\]:*?/\\]")
# Excel stores numbers as doubles and keeps 15 significant digits: a longer integer (an id
# or a document number held as a number) would be silently rounded, so it is written as text.
_EXACT_INT_LIMIT = 10 ** 15

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill(fill_type="solid", fgColor=HEADER_FILL_RGB)
_TITLE_FONT = Font(bold=True, size=12)


@dataclass(frozen=True)
class ColSpec:
    """One data-sheet column. ``kind`` is text | int | date | datetime | money | flag."""
    header: str
    kind: str = "text"
    width: int = 14


@dataclass
class Block:
    """One table on a Read Me / Summary sheet: an optional bold title row, an optional
    styled header row, then the rows. ``kinds`` (per column) picks number formats."""
    title: Optional[str]
    headers: Optional[Sequence[str]]
    rows: Sequence[Sequence[Any]]
    kinds: Optional[Sequence[str]] = None


# ── value safety ─────────────────────────────────────────────────────────────

def _clean_text(s: str) -> Optional[str]:
    if _ILLEGAL_XML_RE.search(s):
        s = _ILLEGAL_XML_RE.sub("", s)
    if len(s) > MAX_CELL_CHARS:
        s = s[:MAX_CELL_CHARS]
    return s or None


def _non_finite_text(v: float | Decimal) -> str:
    # NaN / ±Infinity have no cell representation (openpyxl would write "nan" into <v> and
    # corrupt the sheet); shown as text so the anomaly stays visible.
    if v != v:
        return "NaN"
    return "Infinity" if v > 0 else "-Infinity"


def _json_text(v: Any) -> str:
    return json.dumps(v, ensure_ascii=False, default=str, separators=(", ", ": "))


def _item_text(v: Any) -> str:
    """One element of a list/tuple/set cell, as text ("" for blank)."""
    if v is None:
        return ""
    if isinstance(v, str):
        return v
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if isinstance(v, (dict, list, tuple)):
        return _json_text(v)
    if isinstance(v, (date, time)):
        return v.isoformat()
    return str(v)


def excel_value(v: Any) -> Any:
    """Turn any Python value into something openpyxl can write without corrupting the file.

    str → XML-illegal code points removed, truncated to 32,767 chars ("" → None);
    bool → "Yes"/"No"; Decimal → float; list/tuple/set → ", "-joined text; dict → JSON text;
    date/time/timedelta pass (an aware datetime is converted to naive UTC — Excel has no
    time zones and openpyxl raises on them); NaN/±Infinity → text; integers beyond 15
    digits → text; bytes → UTF-8 text; anything else → ``str(v)``. None stays None, so a
    blank source value stays a blank cell. Formula/error-literal neutralisation needs the
    worksheet and happens in ``ReportWorkbook``.
    """
    if v is None:
        return None
    t = type(v)
    if t is str:
        return _clean_text(v)
    if t is float:
        return v if math.isfinite(v) else _non_finite_text(v)
    if t is int:
        return v if -_EXACT_INT_LIMIT < v < _EXACT_INT_LIMIT else str(v)
    if t is bool:
        return "Yes" if v else "No"
    if isinstance(v, Decimal):
        return float(v) if v.is_finite() else _non_finite_text(v)
    if isinstance(v, datetime):
        return v if v.tzinfo is None else v.astimezone(timezone.utc).replace(tzinfo=None)
    if isinstance(v, time):
        return v.replace(tzinfo=None) if v.tzinfo is not None else v
    if isinstance(v, (date, timedelta)):
        return v
    if isinstance(v, str):                       # str subclasses (StrEnum …)
        return _clean_text(str.__str__(v))
    if isinstance(v, (list, tuple)):
        return _clean_text(", ".join(s for s in map(_item_text, v) if s))
    if isinstance(v, (set, frozenset)):
        return _clean_text(", ".join(sorted(s for s in map(_item_text, v) if s)))
    if isinstance(v, dict):
        return _clean_text(_json_text(v))
    if isinstance(v, (bytes, bytearray, memoryview)):
        return _clean_text(bytes(v).decode("utf-8", errors="replace"))
    if isinstance(v, numbers.Integral):          # IntEnum, numpy integers
        return excel_value(int(v))
    if isinstance(v, numbers.Real):              # numpy floats
        return excel_value(float(v))
    return _clean_text(str(v))


# ── workbook ─────────────────────────────────────────────────────────────────

@dataclass(frozen=True)
class _Styles:
    """Style-table index arrays, resolved once per workbook (indices are per workbook).

    ``plain_date`` / ``plain_datetime`` are the formats openpyxl itself gives a bare date or
    datetime (``cell.TIME_FORMATS``); pre-resolving them lets unstyled date cells skip the
    per-cell number-format lookup openpyxl would otherwise do, with an identical result.
    """
    header: StyleArray
    title: StyleArray
    money: StyleArray
    date: StyleArray
    datetime: StyleArray
    plain_date: StyleArray
    plain_datetime: StyleArray

    @classmethod
    def resolve(cls, ws: WriteOnlyWorksheet) -> "_Styles":
        def array(**attrs: Any) -> StyleArray:
            proto = WriteOnlyCell(ws)
            for name, value in attrs.items():
                setattr(proto, name, value)
            return proto._style

        return cls(
            header=array(font=_HEADER_FONT, fill=_HEADER_FILL),
            title=array(font=_TITLE_FONT),
            money=array(number_format=NUMBER_FORMATS["money"]),
            date=array(number_format=NUMBER_FORMATS["date"]),
            datetime=array(number_format=NUMBER_FORMATS["datetime"]),
            plain_date=array(number_format=TIME_FORMATS[date]),
            plain_datetime=array(number_format=TIME_FORMATS[datetime]),
        )


# Per column: (style for int/float values, style for date values, style for datetime values).
# A plain tuple because it is indexed once per cell. A text value in a money column (an
# unparseable amount kept verbatim) matches none of them and stays unformatted text.
_ColFormat = tuple[Optional[StyleArray], StyleArray, StyleArray]


@dataclass
class _Part:
    ws: WriteOnlyWorksheet
    title: str
    index: int              # 1 = the base sheet, 2 = "<title> (2)", …
    rows: int = 0           # data rows (header / block title rows excluded)


@dataclass
class _Sheet:
    key: str
    title: str
    rank: int
    seq: int                                  # creation order, the tie-break within a rank
    columns: Optional[tuple[ColSpec, ...]]    # None for a block sheet
    formats: tuple[_ColFormat, ...] = ()
    parts: list[_Part] = field(default_factory=list)


def _base_title(title: str) -> str:
    cleaned = _INVALID_TITLE_RE.sub("-", _clean_text(str(title)) or "").strip().strip("'")
    if not cleaned:
        raise ValueError(f"Sheet title {title!r} is empty after removing invalid characters")
    return cleaned


def _part_title(base: str, index: int) -> str:
    suffix = "" if index == 1 else f" ({index})"
    return base[: MAX_SHEET_TITLE - len(suffix)].rstrip() + suffix


class ReportWorkbook:
    """A write-only workbook whose sheets are addressed by stable keys, not titles.

    Data sheets (``ensure_sheet`` + ``append``) have one header row, fixed columns, frozen
    panes, an auto filter and split automatically at ``max_data_rows``. Block sheets
    (``write_blocks``) hold small key/value tables such as Read Me and Summary. Sheets may be
    created in any order; ``save()`` sorts them by rank. Not thread-safe; ``save()`` may run
    in a worker thread once nothing else touches the instance.
    """

    def __init__(self, path: str, *, styled: bool = True, max_data_rows: int = MAX_DATA_ROWS) -> None:
        if not 1 <= max_data_rows <= MAX_DATA_ROWS:
            raise ValueError(f"max_data_rows must be between 1 and {MAX_DATA_ROWS}")
        self._path = path
        self._styled = styled
        self._max_data_rows = max_data_rows
        self._wb = Workbook(write_only=True)
        self._wb.properties.creator = "AirYield"
        self._sheets: dict[str, _Sheet] = {}
        self._used_titles: set[str] = set()      # lower-cased: Excel titles are case-insensitive
        self._style_arrays: Optional[_Styles] = None
        self._finished = False                   # saved or discarded

    # ── data sheets ──────────────────────────────────────────────────────────

    def ensure_sheet(self, key: str, title: str, columns: Sequence[ColSpec], rank: int) -> None:
        """Create the data sheet ``key`` with its header row, once. Later calls are no-ops."""
        self._check_open()
        existing = self._sheets.get(key)
        if existing is not None:
            if existing.columns is None:
                raise ValueError(f"Sheet {key!r} is a block sheet, not a data sheet")
            return
        cols = tuple(columns)
        if not cols:
            raise ValueError(f"Sheet {key!r} needs at least one column")
        if len(cols) > MAX_COLUMNS:
            raise ValueError(f"Sheet {key!r} has {len(cols)} columns; Excel allows {MAX_COLUMNS}")
        bad = sorted({c.kind for c in cols} - COLUMN_KINDS)
        if bad:
            raise ValueError(f"Unknown column kind(s) {bad} on sheet {key!r}")
        sheet = _Sheet(key=key, title=_base_title(title), rank=rank, seq=len(self._sheets), columns=cols)
        self._start_part(sheet)                  # resolves styles, so formats come after
        sheet.formats = tuple(self._format_for(c.kind) for c in cols)
        self._sheets[key] = sheet

    def has_sheet(self, key: str) -> bool:
        return key in self._sheets

    def append(self, key: str, values: Sequence[Any]) -> None:
        """Append one data row. ``values`` must line up with the sheet's columns exactly —
        a misaligned amount in an accounting report is worse than a failed build."""
        self._check_open()
        sheet = self._sheets.get(key)
        if sheet is None or sheet.columns is None:
            raise KeyError(f"No data sheet {key!r}; call ensure_sheet first")
        if len(values) != len(sheet.columns):
            raise ValueError(
                f"Sheet {key!r} expects {len(sheet.columns)} values, got {len(values)}"
            )
        part = sheet.parts[-1]
        if part.rows >= self._max_data_rows:
            part = self._start_part(sheet)
        part.ws.append(self._row_cells(part.ws, values, sheet.formats))
        part.rows += 1

    # ── block sheets ─────────────────────────────────────────────────────────

    def write_blocks(self, key: str, title: str, rank: int, blocks: Sequence[Block]) -> None:
        """Write a whole Read Me / Summary style sheet in one call (write-only sheets cannot
        be revisited). Blocks are separated by one blank row."""
        self._check_open()
        if key in self._sheets:
            raise ValueError(f"Sheet {key!r} already exists")
        blocks = list(blocks)
        lines = max(0, len(blocks) - 1) + sum(
            (b.title is not None) + (b.headers is not None) + len(b.rows) for b in blocks
        )
        if lines > MAX_DATA_ROWS + 1:
            raise ValueError(f"Sheet {key!r} would need {lines} rows; Excel allows {MAX_DATA_ROWS + 1}")
        width_count = max(
            [len(b.headers or ()) for b in blocks] + [len(r) for b in blocks for r in b.rows] + [0]
        )
        if width_count > MAX_COLUMNS:
            raise ValueError(f"Sheet {key!r} has {width_count} columns; Excel allows {MAX_COLUMNS}")
        bad = sorted({k for b in blocks for k in (b.kinds or ())} - COLUMN_KINDS)
        if bad:
            raise ValueError(f"Unknown column kind(s) {bad} in a block on sheet {key!r}")

        sheet = _Sheet(key=key, title=_base_title(title), rank=rank, seq=len(self._sheets), columns=None)
        ws = self._create_ws(_part_title(sheet.title, 1))
        for idx, width in enumerate(_block_widths(blocks, width_count), 1):
            ws.column_dimensions[get_column_letter(idx)].width = width

        part = _Part(ws=ws, title=ws.title, index=1)
        title_style = self._styles().title if self._styled else None
        for n, block in enumerate(blocks):
            if n:
                ws.append([])
            if block.title is not None:
                ws.append([self._text_cell(ws, block.title, title_style)])
            if block.headers is not None:
                ws.append(self._header_cells(ws, block.headers))
            # zip() in _row_cells stops at the shorter sequence, so pad formats to the widest row.
            formats = tuple(self._format_for(k) for k in (block.kinds or ()))
            formats += (self._format_for("text"),) * max(0, width_count - len(formats))
            for row in block.rows:
                ws.append(self._row_cells(ws, row, formats))
                part.rows += 1
        sheet.parts.append(part)
        self._sheets[key] = sheet

    # ── reporting ────────────────────────────────────────────────────────────

    def row_counts(self) -> dict[str, int]:
        """Data rows per sheet key, summed over split parts (block sheets: block rows)."""
        return {key: sum(p.rows for p in s.parts) for key, s in self._sheets.items()}

    def sheet_titles(self) -> dict[str, list[str]]:
        return {key: [p.title for p in s.parts] for key, s in self._sheets.items()}

    # ── finish ───────────────────────────────────────────────────────────────

    def save(self) -> None:
        """Set each open data sheet's filter, order sheets by (rank, part), write the file.

        Callable once. On failure the partial file is removed and the exception propagates;
        call ``discard()`` afterwards to release any remaining temp files.
        """
        self._check_open()
        self._finished = True
        order: dict[int, tuple[float, int, int]] = {}
        for sheet in self._sheets.values():
            for part in sheet.parts:
                order[id(part.ws)] = (sheet.rank, sheet.seq, part.index)
                if sheet.columns is not None and not part.ws.closed:
                    self._set_filter(sheet, part)
        if not self._sheets:
            ws = self._wb.create_sheet(NO_DATA_TITLE)
            ws.append(["No rows matched the selected uploads."])
            order[id(ws)] = (0, 0, 0)
        # A sheet not in ``order`` can only be one whose creation failed half-way; keep it last.
        unregistered = (math.inf, 0, 0)
        self._wb._sheets.sort(key=lambda ws: order.get(id(ws), unregistered))

        # openpyxl's save_workbook() leaves the ZipFile open when writing fails; closing it
        # here releases the handle so the partial file can be deleted (Windows locks it).
        try:
            archive = ZipFile(self._path, "w", ZIP_DEFLATED, allowZip64=True)
            try:
                self._wb.properties.modified = datetime.now(timezone.utc).replace(tzinfo=None)
                ExcelWriter(self._wb, archive).save()
            finally:
                archive.close()
        except BaseException:
            with suppress(OSError):
                os.remove(self._path)
            raise

    def discard(self) -> None:
        """Abandon the workbook: close every sheet's row stream and temp file, delete the temp
        files, write nothing. Safe after a failed save and on repeated calls; never raises."""
        self._finished = True
        for ws in list(self._wb._sheets):
            rows = getattr(ws, "_rows", None)
            if rows is not None:
                with suppress(Exception):
                    rows.close()
            writer = getattr(ws, "_writer", None)
            if writer is None:
                continue
            with suppress(Exception):
                writer.close()
            with suppress(Exception):
                if os.path.exists(writer.out):
                    writer.cleanup()

    # ── internals ────────────────────────────────────────────────────────────

    def _check_open(self) -> None:
        if self._finished:
            raise RuntimeError("The report workbook was already saved or discarded")

    def _create_ws(self, title: str) -> WriteOnlyWorksheet:
        if title.lower() in self._used_titles:
            raise ValueError(f"Duplicate sheet title {title!r}")
        ws = self._wb.create_sheet(title)
        self._used_titles.add(title.lower())
        if self._style_arrays is None:
            self._style_arrays = _Styles.resolve(ws)
        return ws

    def _styles(self) -> _Styles:
        assert self._style_arrays is not None, "styles resolve with the first sheet"
        return self._style_arrays

    def _start_part(self, sheet: _Sheet) -> _Part:
        assert sheet.columns is not None
        if sheet.parts:
            previous = sheet.parts[-1]
            self._set_filter(sheet, previous)
            previous.ws.close()                  # flush + release the file handle now
        index = len(sheet.parts) + 1
        ws = self._create_ws(_part_title(sheet.title, index))
        for idx, col in enumerate(sheet.columns, 1):
            ws.column_dimensions[get_column_letter(idx)].width = max(1, min(MAX_COLUMN_WIDTH, col.width))
        ws.freeze_panes = "A2"
        ws.append(self._header_cells(ws, [c.header for c in sheet.columns]))
        part = _Part(ws=ws, title=ws.title, index=index)
        sheet.parts.append(part)
        return part

    @staticmethod
    def _set_filter(sheet: _Sheet, part: _Part) -> None:
        assert sheet.columns is not None
        last = get_column_letter(len(sheet.columns))
        part.ws.auto_filter.ref = f"A1:{last}{part.rows + 1}"

    def _format_for(self, kind: str) -> _ColFormat:
        s = self._styles()
        if self._styled:
            if kind == "money":
                return s.money, s.plain_date, s.plain_datetime
            if kind == "date":             # a datetime in a date column shows as a date
                return None, s.date, s.date
            if kind == "datetime":
                return None, s.datetime, s.datetime
        return None, s.plain_date, s.plain_datetime

    def _header_cells(self, ws: WriteOnlyWorksheet, headers: Sequence[Any]) -> list[Any]:
        style = self._styles().header if self._styled else None
        return [self._text_cell(ws, h, style) for h in headers]

    @staticmethod
    def _text_cell(ws: WriteOnlyWorksheet, value: Any, style: Optional[StyleArray]) -> Any:
        """A header/title value: a styled new cell, forced to text when it would otherwise be
        read as a formula or error literal."""
        v = excel_value(value)
        forced = type(v) is str and ((v[0] == "=" and len(v) > 1) or (v[0] == "#" and v in _ERROR_LITERALS))
        if style is None and not forced:
            return v
        if v is None:
            return None
        cell = Cell(ws, row=1, column=1, value=v, style_array=style)
        if forced:
            cell.data_type = "s"
        return cell

    @staticmethod
    def _row_cells(ws: WriteOnlyWorksheet, values: Sequence[Any], formats: Sequence[_ColFormat]) -> list[Any]:
        row: list[Any] = []
        push = row.append
        for v, fmt in zip(values, formats):
            if v is None:
                push(None)
                continue
            v = excel_value(v)
            t = type(v)
            if t is str:
                if (v[0] == "=" and len(v) > 1) or (v[0] == "#" and v in _ERROR_LITERALS):
                    cell = Cell(ws, row=1, column=1, value=v)
                    cell.data_type = "s"
                    push(cell)
                else:
                    push(v)
            elif t is float:
                # Money reaches here as float (excel_value converts Decimal), so only floats take
                # the money format — a count in a money-formatted column (Summary "Doc count")
                # must not read "4.00".
                push(v if fmt[0] is None else Cell(ws, row=1, column=1, value=v, style_array=fmt[0]))
            elif t is int:
                push(v)
            elif t is datetime:
                push(Cell(ws, row=1, column=1, value=v, style_array=fmt[2]))
            elif t is date:
                push(Cell(ws, row=1, column=1, value=v, style_array=fmt[1]))
            else:
                push(v)     # None (text that cleaned to blank), time, timedelta: openpyxl's own handling
        return row


def _block_widths(blocks: Sequence[Block], count: int) -> list[int]:
    """Column widths for a block sheet from its header and row text, clamped to 10-60.
    Block titles are excluded: they are meant to spill across the row."""
    widths = [10] * count
    for block in blocks:
        lines: list[Sequence[Any]] = list(block.rows)
        if block.headers is not None:
            lines.append(block.headers)
        for line in lines:
            for idx, value in enumerate(line):
                if value is None:
                    continue
                size = 12 if isinstance(value, (date, time)) else len(str(value))
                if size + 2 > widths[idx]:
                    widths[idx] = min(60, size + 2)
    return widths
