"""Customer and Corporate Billing → "Download Tickets XLS": tickets as one accounting sheet.

The layout is the business's own purchase/sales register — one row per ticket, the sale
built up left to right (fare → service charge → GST → bill amount) and the purchase beside
it (gross → commission → TDS → net).

Two halves. `ticket_row` and `build_workbook` are pure — no session, no I/O — and are what
the tests pin down. `export_rows` is the one async piece: it loads what a sheet needs beyond
the tickets (their saved invoices, who each invoice was addressed to, place of supply) for
the four endpoints that download one — both billing lists and both Sold Tickets tabs.

THE MONEY IS BILLING'S MONEY, NOT A THIRD CALCULATION. A billed ticket reads its saved
invoice line — the snapshot the PDF prints — so the sheet and the invoice cannot disagree
even if the party's markup has changed since. An unbilled ticket is priced exactly as the
Sold Tickets tab previews it (`line_markup` + `split_gst`), which is what Save Billing
would then charge.

BASE FARE IS THE RESIDUAL. `total_amt` is the billing base, and the fee and cancellation
buckets are already inside it (see the projections' "recorded, never summed" notes). So
BASE FARE = base − OTHERS − XXLN, and the first TOTAL always lands on the billed base —
including on a refund, where a negative base and a positive penalty still add up.
"""
from __future__ import annotations

import re
from datetime import date
from io import BytesIO

from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.agency import Agency
from app.models.billing import Billing
from app.models.corporate import Corporate
from app.models.customer import Customer
from app.services.billing_calc import passenger_name, safe_date, split_gst, to_float
from app.services.billing_pdf import _invoice_number, supplier_block
from app.services.party_markup import line_markup, pax_of
from app.services.place_of_supply import place_of_supply

__all__ = [
    "HEADERS", "corporate_name", "customer_name", "ticket_row", "build_workbook",
    "booking_order", "export_rows", "xlsx_download",
]

# The column names as the business's sheet spells them — "PASSANGER" included, since a
# sheet that is pasted into an existing register has to match its header.
HEADERS = [
    "Type", "BILL NO", "BILL TO", "SUPPLIER", "BOOKING DATE", "NAME OF PASSANGER",
    "FROM", "TO", "TO", "PNR", "FLIGHT NO / HOTEL NAME", "DATE FROM", "DATE TO",
    "BASE FARE", "OTHERS", "XXLN CHARGES", "ADDITIONAL MARKUP", "TOTAL",
    "SERVICE CHARGE", "TOTAL", "GST CHARGED", "TOTAL AMOUNT",
    "GROSS PURCHASE", "COM", "TDS", "NET PURCHASE", "BILL AMOUNT",
]

# 1-based column numbers, for the workbook's formats.
_DATE_COLUMNS = (5, 12, 13)
_FIRST_MONEY_COLUMN = HEADERS.index("BASE FARE") + 1

# The fee / ancillary buckets the projections write INSIDE `total_amt`.
_OTHER_FIELDS = ("booking_fee_sell", "seat_selection", "excess_baggage", "meals")


def _money(value: float) -> float:
    return round(value, 2)


def _optional(value: float) -> float | None:
    """A component the sheet leaves blank when there is none — OTHERS, COM, TDS…"""
    return _money(value) if abs(value) >= 0.005 else None


def _person_and_company(party) -> tuple[str, str]:
    person = f"{getattr(party, 'first_name', None) or ''} {getattr(party, 'last_name', None) or ''}".strip()
    return person, (getattr(party, "company", None) or "").strip()


def corporate_name(corporate) -> str:
    """BILL TO for a corporate — its company; pre-split rows may carry only a contact's name.
    The same headline billing_pdf._bill_to_lines prints on a corporate invoice."""
    person, company = _person_and_company(corporate)
    return company or person


def customer_name(customer) -> str:
    """BILL TO for a customer — the PERSON, not their employer. A customer bill is a direct
    sale to the traveller, and billing_pdf._bill_to_lines addresses it to them by name."""
    person, company = _person_and_company(customer)
    return person or company


def _excel_date(raw) -> date | str | None:
    """A real date when it parses (so the column sorts), else the text as it came."""
    if raw is None or not str(raw).strip():
        return None
    return safe_date(raw) or str(raw).strip()


def _route_points(sector: str | None) -> list[str]:
    """"COK-BLR / BLR-CCU" → [COK, BLR, CCU]; "DEL-STV / STV-DEL" → [DEL, STV, DEL].

    Legs are joined by " / " (LCC, NDC). Inside a leg the points are split on "/" when it
    has one — Third Party API writes "MOPA Airport/Hyderabad", and a place name there can
    carry a hyphen — else on "-". A point repeated where one leg ends and the next begins
    is the connection, and is kept once.
    """
    if not sector or not sector.strip():
        return []
    points: list[str] = []
    for leg in re.split(r"\s+/\s+", sector.strip()):
        parts = leg.split("/") if "/" in leg else leg.split("-")
        for part in (p.strip() for p in parts):
            if part and (not points or points[-1].upper() != part.upper()):
                points.append(part)
    return points


def _flight_label(flight_no: str | None, airline_code: str | None) -> str | None:
    """"5071" + 6E → "6E5071". A leg that already names its carrier ("AI423") is kept."""
    if not flight_no or not flight_no.strip():
        return None
    code = (airline_code or "").strip().upper()
    legs = [leg.strip() for leg in re.split(r"\s*/\s*", flight_no.strip()) if leg.strip()]
    return " / ".join(f"{code}{leg}" if code and leg.isdigit() else leg for leg in legs)


def _last_leg_date(segments) -> date | None:
    if not isinstance(segments, list):
        return None
    dates = [safe_date(s.get("dep_date")) for s in segments if isinstance(s, dict) and s.get("dep_date")]
    dates = [d for d in dates if d]
    return dates[-1] if len(dates) > 1 else None


def _itinerary(t) -> dict:
    """FROM / TO / TO / PNR / FLIGHT NO-HOTEL NAME / DATE FROM / DATE TO for one line."""
    category = (getattr(t, "product_category", None) or "air").lower()
    travel = _excel_date(getattr(t, "travel_dt", None) or getattr(t, "departure_datetime", None))

    if category == "air":
        points = _route_points(getattr(t, "sector", None))
        date_to = _last_leg_date(getattr(t, "segments", None))
        if isinstance(travel, date) and date_to == travel:
            date_to = None          # a same-day connection has no separate "to" date
        return {
            "from": points[0] if points else None,
            "to": points[1] if len(points) > 1 else None,
            "to2": "-".join(points[2:]) or None,
            "pnr": getattr(t, "air_pnr", None) or getattr(t, "gds_pnr", None) or getattr(t, "booking_ref", None),
            "service": _flight_label(getattr(t, "flight_no", None), getattr(t, "airlines_code", None)),
            "date_from": travel,
            "date_to": date_to,
        }

    # Hotel / train / bus / car: what the line IS lives in service_details
    # (services/tp_api_itinerary), whose `route` is "A → B", or just a hotel's city.
    details = getattr(t, "service_details", None)
    details = details if isinstance(details, dict) else {}
    ends = [p.strip() for p in (details.get("route") or "").split("→") if p.strip()]
    title = details.get("title")
    if title and title == details.get("label"):
        title = None                # "Hotel" is the placeholder for an unnamed property
    return {
        "from": ends[0] if ends else None,
        "to": ends[1] if len(ends) > 1 else None,
        "to2": None,
        "pnr": details.get("pnr") or getattr(t, "booking_ref", None),
        "service": title,
        "date_from": _excel_date(details.get("check_in") or details.get("journey_date")
                                 or details.get("start_date")) or travel,
        "date_to": _excel_date(details.get("check_out") or details.get("end_date")),
    }


def ticket_row(t, party, *, line: dict | None, bill_no: str | None,
               interstate: bool | None, bill_to: str | None) -> list:
    """One sheet row, in HEADERS order.

    `party` is the customer or corporate the ticket is billed under — its markup and billing
    type price an unbilled line. `line` is the ticket's saved invoice line when it is billed;
    its figures are used as stored. `interstate` is the party's place-of-supply decision
    (only the GST total is shown, so it moves nothing but paise).

    `bill_to` is required rather than read off `party`, because a party's name depends on
    what KIND it is (a customer is billed by person, a corporate by company) and on whether
    the invoice went to someone else: a ticket naming both an employee and their employer
    can already sit on the employee's own bill, raised before one ticket had one payer.
    """
    others = sum(to_float(getattr(t, f, None)) for f in _OTHER_FIELDS)
    xxln = to_float(getattr(t, "can_charge", None))

    if line is not None:
        base = to_float(line.get("base_amount"))
        markup = to_float(line.get("markup_amount"))
        additional = to_float(line.get("additional_markup"))
        gst = to_float(line.get("gst_amount"))
        bill_amount = to_float(line.get("total"))
        pax = line.get("pax_count") or pax_of(t)
    else:
        base = to_float(t.total_amt) if t.total_amt is not None else to_float(t.sell_fare)
        markup, _note = line_markup(base, party, t)
        additional = 0.0
        gst = split_gst(base, markup, party.billing_type, interstate=interstate)["gst_amount"]
        bill_amount = base + markup + gst
        pax = pax_of(t)

    fare_total = base + additional
    with_service = fare_total + markup
    total_amount = with_service + gst

    # What the supplier said we earned, else what the deal engine worked out.
    comm_sell = getattr(t, "comm_sell", None)
    commission = to_float(comm_sell) if comm_sell is not None else to_float(getattr(t, "calculated_incentive", None))
    tds = to_float(getattr(t, "tds_sell", None))

    name = passenger_name(t)
    if isinstance(pax, int) and pax > 1:
        name = f"{name}^{pax}"

    it = _itinerary(t)
    return [
        (getattr(t, "product_category", None) or "air").upper(),
        bill_no,
        bill_to,
        getattr(t, "booking_agency_name", None) or getattr(t, "airline_name", None),
        _excel_date(getattr(t, "ticket_date", None)),
        name,
        it["from"], it["to"], it["to2"],
        it["pnr"],
        it["service"],
        it["date_from"], it["date_to"],
        _money(base - others - xxln),
        _optional(others),
        _optional(xxln),
        _optional(additional),
        _money(fare_total),
        _money(markup),
        _money(with_service),
        _optional(gst),
        _money(total_amount),
        _money(base),
        _optional(commission),
        _optional(tds),
        _money(base - commission + tds),
        _money(bill_amount),
    ]


def build_workbook(rows: list[list], title: str = "Tickets") -> BytesIO:
    """The rows under a styled header — navy, frozen, filterable — as .xlsx bytes."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]           # Excel's limit on a sheet name
    ws.append(HEADERS)

    head_fill = PatternFill("solid", fgColor="1E3A5F")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in rows:
        ws.append(row)

    for r in range(2, ws.max_row + 1):
        for c in _DATE_COLUMNS:
            ws.cell(row=r, column=c).number_format = "DD-MMM-YY"
        for c in range(_FIRST_MONEY_COLUMN, len(HEADERS) + 1):
            ws.cell(row=r, column=c).number_format = "#,##0.00"

    widths = [7, 16, 28, 18, 12, 30, 8, 8, 8, 12, 22, 12, 12] + [14] * (len(HEADERS) - 13)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def booking_order(t) -> tuple:
    """Sort key: booking date, undated last, then upload order."""
    return (safe_date(getattr(t, "ticket_date", None)) or date.max, t.id)


def _party_label(party) -> str:
    return customer_name(party) if isinstance(party, Customer) else corporate_name(party)


async def export_rows(db: AsyncSession, current_user, pairs: list[tuple]) -> list[list]:
    """Sheet rows for `(ticket, party)` pairs, in the order given.

    `party` is the Customer or Corporate each ticket is billed under — a pair rather than
    read off the ticket, because an untagged ticket reached by passenger name names nobody,
    and an employee's ticket names two parties of which only one pays.

    BILL TO follows the invoice when there is one: whoever the billing was raised to — a
    person, a company or an agency — which is not always `party`. Unbilled, it is `party`.
    """
    tickets = [t for t, _party in pairs]
    billing_ids = {t.billing_id for t in tickets if t.billing_id}
    billings: dict[int, Billing] = {}
    if billing_ids:
        billings = {b.id: b for b in (await db.execute(
            select(Billing).where(
                Billing.id.in_(billing_ids),
                Billing.tenant_id == current_user.tenant_id,
                Billing.created_by_id == current_user.id,
            )
        )).scalars().all()}

    async def addressees(model, ids: set[int], label) -> dict[int, str]:
        if not ids:
            return {}
        q = select(model).where(model.id.in_(ids))
        if model is not Agency:     # an agency has no creator; the billing naming it is scoped
            q = q.where(model.tenant_id == current_user.tenant_id,
                        model.created_by_id == current_user.id)
        return {row.id: label(row) for row in (await db.execute(q)).scalars().all()}

    billed = billings.values()
    names = {
        "customer": await addressees(Customer, {b.customer_id for b in billed if b.customer_id}, customer_name),
        "corporate": await addressees(Corporate, {b.corporate_id for b in billed if b.corporate_id}, corporate_name),
        "agency": await addressees(Agency, {b.agency_id for b in billed if b.agency_id}, lambda a: a.name),
    }

    def invoice_addressee(billing: Billing) -> str | None:
        if billing.corporate_id:
            return names["corporate"].get(billing.corporate_id)
        if billing.customer_id:
            return names["customer"].get(billing.customer_id)
        if billing.agency_id:
            return names["agency"].get(billing.agency_id)
        return None

    # The invoice number is derived from the billing and the workspace name, exactly as
    # the PDF derives it — so BILL NO is the number printed on the bill.
    agency = supplier_block(current_user.tenant, current_user)
    interstate: dict[tuple, bool | None] = {}

    rows = []
    for t, party in pairs:
        key = (type(party).__name__, party.id)
        if key not in interstate:
            interstate[key] = place_of_supply(current_user.tenant, party).interstate
        billing = billings.get(t.billing_id) if t.billing_id else None
        line = next(
            (it for it in (billing.line_items or []) if isinstance(it, dict) and it.get("ticket_id") == t.id),
            None,
        ) if billing else None
        rows.append(ticket_row(
            t, party,
            line=line,
            bill_no=_invoice_number(billing, agency) if billing else None,
            interstate=interstate[key],
            bill_to=(invoice_addressee(billing) if billing else None) or _party_label(party),
        ))
    return rows


def xlsx_download(rows: list[list], filename: str, title: str = "Tickets") -> StreamingResponse:
    return StreamingResponse(
        build_workbook(rows, title),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
