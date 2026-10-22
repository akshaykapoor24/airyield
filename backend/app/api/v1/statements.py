"""Generic spec-driven vendor statements (TGQ HMPR, later NDC / LCC / GDS).

Upload an XLS/CSV export; the spec's fixed columns are stored verbatim in a JSONB
``data`` blob and the repeating ``Tax_TypeN`` / ``TaxN`` pairs are folded into a JSONB
``taxes`` array (so any number of taxes is supported — no fixed Tax1..Tax20 columns).
Each statement type has its OWN dedicated table, resolved by ``{slug}`` via
``STATEMENT_MODELS`` (e.g. ``tgq-hmpr`` → ``tgq_hmpr``). Same REST shape as the
airline-adjustments router, so the frontend batch view is reused unchanged. The
original spreadsheet is stored in GCS for preview/download. Scoped per user/tenant.
"""
from __future__ import annotations

import asyncio
import hashlib
import io
import json
import math
import os
import re
import tempfile
import uuid
from contextlib import suppress
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal

import pandas as pd
from fastapi import (
    APIRouter, Depends, File, Form, HTTPException, Query, Request, UploadFile, status,
)
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import Numeric, and_, case, cast, delete, func, literal, or_, select, update
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import aliased, defer
from starlette.background import BackgroundTask

from app.config import settings
from app.database import get_db
from app.dependencies import get_current_user
from app.models.commission_calculation import CommissionCalculation
from app.models.commission_run import CommissionRun
from app.models.statement_batch_supplier import StatementBatchSupplier
from app.models.statement_batch_billing import StatementBatchBilling
from app.models.statement_batch_airline_id import StatementBatchAirlineId
from app.models.statement_row import STATEMENT_MODELS
from app.models.third_party_statement import THIRD_PARTY_SLUGS, ThirdPartyStatement
from app.models.typed_statement import UNPARSED_KEY, is_typed
from app.models.tenant_airline import TenantAirline
from app.models.ticket_statement import TicketStatement
from app.models.uploaded_ticket import UploadedTicket
from app.models.user import User
from app.services import flat_statement as _flat
from app.services import income_board
from app.services import sector_split, spec_mapping, spreadsheet, statement_spec as spec
from app.services import statement_balance, statement_export, statement_period, tp_gds_reupload
from app.services import date_format_detect, excel_dates
from app.models.statement_batch_control import StatementBatchControl
from app.services import statement_supplier_selection as supplier_selection
from app.services import tp_airline_resolution as tp_airline
# Kept under the old private name: tests (test_tp_api_spec) and this router call it so.
from app.services.statement_display import display_columns as _display_columns, flatten_record
from app.services.lcc_airline_selection import resolve_for_upload

router = APIRouter()

_MIN_MATCHED_COLUMNS = 3
_TAX_TYPE_RE = re.compile(r"^tax_type(\d+)$")
_TAX_AMT_RE = re.compile(r"^tax(\d+)$")
# Drill-in filters arrive as `f.<field>=value` so adding one is a spec edit, not a
# signature change. Unknown fields are ignored, never rejected (a stale bookmark should
# degrade, not 400).
_FILTER_PREFIX = "f."
# Facet dropdowns past this are unusable anyway; the client falls back to typing.
_MAX_FACET_VALUES = 200


def _bucket() -> str:
    return settings.GCS_BSP_BUCKET_NAME or settings.GCS_TICKETS_BUCKET_NAME


def _resolve(slug: str):
    """(normalized slug, dedicated model) or 404."""
    s = (slug or "").lower()
    model = STATEMENT_MODELS.get(s)
    if model is None or spec.spec_for(s) is None:
        raise HTTPException(status_code=404, detail=f"Unknown statement type '{slug}'.")
    return s, model


def _scope(model, user: User):
    return (model.tenant_id == user.tenant_id, model.created_by_id == user.id)


def _not_balance_row(model):
    """The statement's own OLD/BALANCE/total lines are summaries, not entries.

    Stamped at ingest (`data.row_kind`, services/statement_balance.py) for the types whose
    spec declares `balance_lines`. Kept out of every count, list and sum here for the same
    reason a TGQ export's `is_total` line is.
    """
    return model.data[statement_balance.ROW_KIND_KEY].astext.is_(None)


def _splits(model) -> bool:
    """Does this table carry the per-sector leg columns (models/statement_row._SplitMixin)?

    Only the ticket-shaped tables (tgq_hmpr, ndc) do; the ledger-shaped ones don't, so
    every leg/total reference below is gated on this.
    """
    return hasattr(model, "sector_count")


def _active_filters(slug: str, request: Request | None) -> dict[str, tuple[dict, str]]:
    """Declared `f.<field>` params present on the request → {field: (filter_spec, value)}."""
    if request is None:
        return {}
    declared = {f["field"]: f for f in spec.filter_specs(slug)}
    out: dict[str, tuple[dict, str]] = {}
    for key, raw in request.query_params.multi_items():
        if not key.startswith(_FILTER_PREFIX):
            continue
        f = declared.get(key[len(_FILTER_PREFIX):])
        value = (raw or "").strip()
        if f and value:
            out[f["field"]] = (f, value)
    return out


def _record_conds(model, slug: str, user: User, batch_id: str | None, request: Request | None) -> list:
    """Scope + batch + declared filters, with the export's own grand-total line excluded.

    The total line is a summary, not a ticket — it must never be a data row, get counted,
    or be swept into a SUM.
    """
    conds = [*_scope(model, user)]
    if batch_id:
        conds.append(model.batch_id == batch_id)
    if _splits(model):
        conds.append(model.is_total.is_(False))
    if spec.balance_lines(slug):
        conds.append(_not_balance_row(model))
    for field, (f, value) in _active_filters(slug, request).items():
        col = model.data[field].astext
        conds.append(col == value if f.get("type") == "select" else col.ilike(f"%{value}%"))
    return conds


def _num(model, field: str):
    """`data->>field` as NUMERIC, treating anything non-numeric as 0.

    Values are stored verbatim as strings (see `_clean`), so a stray "N/A" or a header
    fragment must not blow up the aggregate. The CASE guard is safe — Postgres won't
    constant-fold a cast of a column reference, so it only evaluates on matching rows.
    """
    txt = func.replace(func.coalesce(model.data[field].astext, ""), ",", "")
    return case(
        (txt.op("~")(literal(r"^-?\d+(\.\d+)?$")), cast(txt, Numeric(20, 4))),
        else_=literal(0, Numeric(20, 4)),
    )


def _money_str(value) -> str | None:
    """Serialize a Decimal as a string — a float round-trip would lose paise."""
    return None if value is None else format(value.normalize(), "f")


def _clean(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    s = str(value).strip()
    if s == "" or s.lower() == "nan":
        return None
    return s


_MIDNIGHT = re.compile(r"^\d{4}-\d{2}-\d{2} 00:00:00$")


def _sample_cell(value) -> str | None:
    """A cell as the review grid shows it — a stored Excel date as its day, "2026-08-01",
    not "2026-08-01 00:00:00". Display only: `/confirm` reads the file, not these."""
    s = _clean(value)
    return s[:10] if s and _MIDNIGHT.match(s) else s


# Delimited text. `sep=None` with the python engine sniffs the separator, so tab- and
# semicolon-separated exports need no extra branch — only their extensions do. Before
# these were listed, a .tsv or .txt fell through to `read_excel` and simply threw, and a
# consolidator's "spreadsheet" is routinely a tab-separated .txt.
_DELIMITED_EXT = (".csv", ".tsv", ".tab", ".txt")


def _read_df(content: bytes, filename: str, header_row: int) -> pd.DataFrame:
    """One line, because services/spreadsheet.py now owns the format question.

    This used to branch on the extension, which was wrong twice over: a legacy .xls
    reached pandas with no engine and raised ImportError, and a consolidator's
    "xls" is as often an xlsx or an HTML table as it is a real BIFF workbook. The
    reader decides from the bytes.
    """
    return spreadsheet.read_df(content, filename, header_row)


def _parse(content: bytes, filename: str, slug: str):
    """(df, {col->field}, {col->normalized}) for the header row that recognises the most columns.

    Aliases count as recognition where the spec declares them, so a direct `/upload` of an
    airline's own export sees the same columns the mapping wizard would propose — the two
    entry points disagreeing about what a file contains would be worse than either rule.
    `normmap` stays the raw normalized header regardless: tax folding matches on
    `Tax_TypeN`/`TaxN`, which are not fields and have no aliases.
    """
    field_set = set(spec.fields(slug))
    alias_map: dict[str, str] = {}
    for field, names in spec.aliases(slug).items():
        for a in names:
            alias_map.setdefault(spec.norm(a), field)
    best = None
    last_err: Exception | None = None
    for header_row in (0, 1, 2):
        try:
            df = _read_df(content, filename, header_row)
        except Exception as e:  # noqa: BLE001 — unreadable at this offset, try the next
            last_err = e
            continue
        df.dropna(how="all", inplace=True)
        colmap: dict = {}
        normmap: dict = {}
        for col in df.columns:
            key = spec.norm(col)
            normmap[col] = key
            field = key if key in field_set else alias_map.get(key)
            if field and field not in colmap.values():
                colmap[col] = field
        if best is None or len(colmap) > best[0]:
            best = (len(colmap), df, colmap, normmap)
    if best is None:
        raise HTTPException(status_code=400, detail=f"Could not read the file: {last_err}. Upload a valid .xlsx, .xls or .csv.")
    matched, df, colmap, normmap = best
    if matched < _MIN_MATCHED_COLUMNS:
        sample = ", ".join(c["header"] for c in spec.columns(slug)[:5])
        raise HTTPException(
            status_code=400,
            detail=f"This doesn't look like a {spec.spec_for(slug)['label']} file — only {matched} known "
                   f"column(s) were recognised. Expected headers such as: {sample}…",
        )
    return df, colmap, normmap


def _fold_taxes(row, normmap: dict) -> list[dict]:
    """Fold every Tax_TypeN / TaxN pair (any N) into [{type, amount}, ...]; drop empty types."""
    types: dict[int, str | None] = {}
    amounts: dict[int, str | None] = {}
    for col, nz in normmap.items():
        m = _TAX_TYPE_RE.match(nz)
        if m:
            types[int(m.group(1))] = _clean(row[col])
            continue
        m = _TAX_AMT_RE.match(nz)
        if m:
            amounts[int(m.group(1))] = _clean(row[col])
    out: list[dict] = []
    for n in sorted(types):
        t = types[n]
        if t:
            out.append({"type": t, "amount": amounts.get(n)})
    return out


def _detect_df(content: bytes, filename: str, score_fn):
    """Pick the header row (0/1/2) whose columns score highest via `score_fn(cols)`."""
    best = None
    last_err: Exception | None = None
    for hr in (0, 1, 2):
        try:
            df = _read_df(content, filename, hr)
        except Exception as e:  # noqa: BLE001
            last_err = e
            continue
        df.dropna(how="all", inplace=True)
        score = score_fn([str(c) for c in df.columns])
        if best is None or score > best[0]:
            # Carried on the frame rather than returned beside it, so the existing callers
            # are untouched; the control-figure capture needs it to read the preamble.
            df.attrs["header_row"] = hr
            best = (score, df)
    if best is None:
        raise HTTPException(status_code=400, detail=f"Could not read the file: {last_err}. Upload a valid .xlsx, .xls or .csv.")
    if best[0] < _MIN_MATCHED_COLUMNS:
        raise HTTPException(status_code=400, detail="This doesn't look like a valid statement — too few recognised columns.")
    return best[1]


def _preamble_opening(content: bytes, filename: str, header_row: int | None) -> Decimal | None:
    """The statement's OLD/opening line, read off the rows ABOVE the header.

    Everything above the header is discarded on a normal read, and that is where a running
    account prints its opening balance ("OLD" | 627257). A marker cell followed by a number
    — or one cell holding both — is the opening; a title row ("TICKET STATEMENT") has no
    number after it and is ignored. Best effort: any failure means "not printed".
    """
    if not header_row:
        return None
    try:
        cells = spreadsheet.read_table(content, filename, header_row=header_row).preamble
    except Exception:  # noqa: BLE001 — a preamble we cannot read is a preamble we skip
        return None
    for i, c in enumerate(cells):
        text = " ".join(str(c).upper().split())
        if not any(text == m or text.startswith(m + " ") for m in statement_balance.OPENING_MARKERS):
            continue
        same_cell = statement_balance.to_decimal(text.split()[-1]) if " " in text else None
        if same_cell is not None:
            return same_cell
        if i + 1 < len(cells):
            nxt = statement_balance.to_decimal(cells[i + 1])
            if nxt is not None:
                return nxt
    return None


def _control_fields(control, objs: list | None = None) -> dict | None:
    """What the upload wizard's done screen compares: what was declared vs what landed."""
    if control is None:
        return None
    tickets = [o for o in (objs or [])
               if not (o.data or {}).get(statement_balance.ROW_KIND_KEY)]
    net = sum((statement_balance.to_decimal((o.data or {}).get("net_amount")) or Decimal("0")
               for o in tickets), Decimal("0"))
    return {
        "file_rows": control.file_rows, "loaded_rows": control.loaded_rows,
        "ticket_rows": len(tickets), "ticket_net": _money_str(net),
        "expected_count": control.expected_count,
        "expected_amount": _money_str(control.expected_amount)
        if control.expected_amount is not None else None,
        "opening_balance": _money_str(control.opening_balance)
        if control.opening_balance is not None else None,
        "closing_balance": _money_str(control.closing_balance)
        if control.closing_balance is not None else None,
        # Entries of the file left out as already uploaded — still part of what the file declares.
        "duplicate_rows": control.duplicate_rows or 0,
        "duplicate_net": _money_str(control.duplicate_net)
        if control.duplicate_net is not None else None,
    }


def _to_decimal_form(value: str | None, label: str) -> Decimal | None:
    """An optional numeric form field ('1,23,456.78' allowed), or 400 when it is not one."""
    if value is None or not str(value).strip():
        return None
    d = statement_balance.to_decimal(value)
    if d is None:
        raise HTTPException(status_code=400, detail=f"{label} must be a number.")
    return d


def _control_row(slug: str, user: User, batch_id: str, *, content: bytes, filename: str,
                 header_row: int | None, file_rows: int, objs: list,
                 expected_count: int | None, expected_amount: Decimal | None,
                 decision: tp_gds_reupload.Decision | None = None):
    """The upload's control figures (models/statement_batch_control.py), or None for types
    that do not capture them. Written in the same transaction as the rows.

    A Third Party GDS upload may save only part of its file; the lines it left out as
    duplicates — their count, net and the upload most of them are in — are recorded here,
    because the file's own arithmetic still counts them."""
    if not spec.captures_controls(slug):
        return None
    skipped = len(decision.left_out) if decision is not None else 0
    closing = None
    for o in objs:
        data = o.data or {}
        if data.get(statement_balance.ROW_KIND_KEY) == "closing":
            closing = statement_balance.to_decimal(data.get("net_amount")) or closing
    opening = _preamble_opening(content, filename, header_row)
    return StatementBatchControl(
        tenant_id=user.tenant_id, created_by_id=user.id, slug=slug, batch_id=batch_id,
        file_rows=file_rows, loaded_rows=len(objs), header_row=header_row,
        opening_balance=opening, opening_source="file" if opening is not None else None,
        closing_balance=closing,
        expected_count=expected_count if (expected_count or 0) > 0 else None,
        expected_amount=expected_amount,
        duplicate_rows=skipped,
        duplicate_net=decision.left_net if skipped else None,
        duplicate_of=decision.duplicate_of if skipped else None,
    )


def _build_rows(model, slug: str, prov: dict, data: dict, taxes: list[dict], seq: int,
                extra: dict | None = None) -> list:
    """One parsed source line → the model row(s) it becomes.

    Shared by upload and re-process so both always produce identical output. For a type
    that declares no `split_sectors` this is a single unchanged row, i.e. exactly the old
    behaviour.

    `extra` is the provenance a spec-driven type on a `_NormalizedBase` table carries and a
    plain `_StatementBase` one does not — `source_format` and `raw_data`. It is empty for
    every type that does not declare `detect_format`, and it is only ever passed on the
    non-splitting branch: no splitting type has those columns.
    """
    if not _splits(model):
        return [model(**prov, **(extra or {}), data=data, taxes=taxes)]

    # The export's own grand-total line: kept (it's the vendor's declared figure, shown in
    # the summary slab) but flagged, so it never behaves like a ticket. Checked before any
    # derivation — it carries no ticket number to split.
    if sector_split.is_total_row(data, spec.total_row_config(slug)):
        return [model(**prov, data=data, taxes=taxes, row_seq=seq,
                      sector_index=1, sector_count=1,
                      split_status=sector_split.TOTAL, is_total=True)]

    source_data = data   # verbatim, before anything is derived from it
    data, derived = sector_split.apply_ticket_no(data, spec.ticket_no_config(slug))

    cfg = spec.split_config(slug)
    if not cfg:
        return [model(**prov, data=data, taxes=taxes, row_seq=seq,
                      sector_index=1, sector_count=1, split_status=sector_split.SINGLE,
                      orig_data=source_data if derived else None)]

    out = []
    for leg_data, leg_taxes, idx, count, status in sector_split.split_row(data, taxes, cfg):
        # Stored whenever the row was actually altered — it's the audit trail and the
        # source for re-deriving legs under a different allocation rule.
        changed = count > 1 or derived
        out.append(model(
            **prov, data=leg_data, taxes=leg_taxes, row_seq=seq,
            sector_index=idx, sector_count=count, split_status=status,
            orig_data=source_data if changed else None,
            orig_taxes=taxes if count > 1 else None,
        ))
    return out


# The JSONB companions a normalized line carries beside `data`. A typed table
# (models/typed_statement.py) keeps none of them — its fields are columns, and it keeps no
# verbatim copy of the vendor's row — so they are left off its lines.
_JSON_EXTRAS = ("taxes", "segments", "ssr", "raw_data")


def _line(model, **values):
    """One statement line of `model`, minus the JSONB companions a typed table does not keep."""
    if is_typed(model):
        for k in _JSON_EXTRAS:
            values.pop(k, None)
    return model(**values)


def _fmt_extra(model, detect, src_fmt: str | None, row, columns: list[str]) -> dict:
    """`source_format` + `raw_data` for a spec-driven type whose table carries them.

    `{}` for every type that does not declare `detect_format`, so the eight existing types
    write exactly the rows they wrote before. The `hasattr` is not belt-and-braces: a
    spec-driven type could be put on a plain `_StatementBase` table, which has neither
    column, and passing them would be a TypeError at insert rather than at import.

    `raw_data` is the whole source line keyed by its ORIGINAL header — not the mapped row.
    That is what makes the mapping recoverable: a column the uploader mapped wrongly, or did
    not map at all, is still on the row and can be re-read without asking for the file again.
    """
    if detect is None or not hasattr(model, "source_format"):
        return {}
    return {
        "source_format": src_fmt,
        "raw_data": {c: v for c in columns if (v := _clean(row.get(c))) is not None},
    }


async def _clear_commission(db: AsyncSession, slug: str, user: User,
                            batch_id: str | None = None, row_ids: list[int] | None = None) -> int:
    """Drop commission figures derived from rows that are about to change or disappear.

    `commission_calculations.source_row_id` points at a statement row by id with NO foreign
    key — it names a different table per source, so there is nothing for a constraint to
    reference. Nothing removes these for us, and a stale one is worse than a missing one:
    a reprocess re-inserts its rows with NEW ids, so the old figures would sit there
    describing rows that no longer exist while the grid showed the new rows as unpriced.

    Returns how many were cleared, so the caller can tell the user to re-run.
    """
    conds = [CommissionCalculation.tenant_id == user.tenant_id,
             CommissionCalculation.created_by_id == user.id,
             CommissionCalculation.source == slug]
    if batch_id:
        conds.append(CommissionCalculation.batch_id == batch_id)
    if row_ids is not None:
        conds.append(CommissionCalculation.source_row_id.in_(row_ids))
    res = await db.execute(delete(CommissionCalculation).where(*conds))
    if batch_id and row_ids is None:
        await db.execute(delete(CommissionRun).where(
            CommissionRun.tenant_id == user.tenant_id,
            CommissionRun.created_by_id == user.id,
            CommissionRun.source == slug,
            CommissionRun.batch_id == batch_id,
        ))
    return res.rowcount or 0


async def _project(db: AsyncSession, slug: str, user: User, batch_id: str) -> int:
    """Refresh this batch on the income/revenue board. Best effort; never raises.

    Called BEFORE the caller's commit and inside its transaction, so an upload that
    rolls back leaves no projection of rows that never landed. The flush is what makes
    the pending `db.add_all(objs)` visible to a set-based INSERT ... SELECT that reads
    the table rather than the session.

    A no-op for the types the board has no arm for. Only NDC and the three third-party
    types carry sale it can project; TGQ HMPR and the LCC ledgers restate another
    statement and are read live instead — see models/income_board.SALE_SOURCES.
    """
    await db.flush()
    return await income_board.refresh(
        db, tenant_id=user.tenant_id, user_id=user.id, source=slug, batch_id=batch_id,
    )


async def _airline_stamper(db: AsyncSession, slug: str):
    """`(data) -> bool` that resolves and stamps a row's carrier, or None for types that
    don't need it.

    Built ONCE per upload/reprocess — the airline master is one query, not one per row.
    Returns None (rather than a no-op) so the caller can skip the loop entirely and so a
    type that never opted in cannot accidentally acquire an `airline_code`.
    """
    if not spec.resolves_airline(slug):
        return None
    index = await tp_airline.build_index(db)

    def stamp(data: dict) -> bool:
        match = tp_airline.resolve_tp_airline(
            data.get("ticket_prefix"), data.get("airline_code"), data.get("airline_name"), index,
        )
        return tp_airline.stamp(data, match)

    return stamp


def _get_builder(parser_name: str):
    """Resolve a custom-parser module exposing build_col_map(cols) + build_row(row, cols)."""
    if parser_name == "lcc":
        from app.services import lcc_statement as m
        return m
    if parser_name == "di":
        from app.services import di_statement as m
        return m
    if parser_name == "divided-pnr":
        from app.services import divided_pnr as m
        return m
    if parser_name == "flown-report":
        from app.services import flown_report as m
        return m
    if parser_name == "cta-bta":
        from app.services import cta_bta_report as m
        return m
    # Generic flat parsers (Third Party GDS/LCC, …) registered in flat_statement.
    from app.services import flat_statement
    b = flat_statement.get(parser_name)
    if b is not None:
        return b
    raise HTTPException(status_code=500, detail=f"No parser registered for '{parser_name}'.")


# ── routes ───────────────────────────────────────────────────────────────────
@router.get("/{slug}/columns")
async def get_columns(slug: str, current_user: User = Depends(get_current_user)):
    slug, _ = _resolve(slug)
    return _display_columns(slug)


async def _airline_selection(
    db: AsyncSession, slug: str, user: User, tenant_airline_ids: list[int]
) -> list[TenantAirline]:
    """The Airline Master ids declared for this upload.

    Required only for the types whose spec says so — the LCC ones, whose exports name
    no carrier. TGQ HMPR, NDC and the third-party types share this endpoint and carry
    their own airline, so a selection sent for them is ignored rather than stored.
    """
    if not spec.requires_airline_id(slug):
        return []
    try:
        return await resolve_for_upload(db, user.tenant_id, tenant_airline_ids)
    except ValueError as exc:
        # MixedAirlineSelection and UnknownAirlineId are both ValueErrors and both
        # carry a message written for the user.
        raise HTTPException(status_code=400, detail=str(exc))


def _supplier_link(tenant_id: int, slug: str, batch_id: str, supplier) -> StatementBatchSupplier:
    """The batch → consolidator row, with its name/code/branch SNAPSHOTTED.

    Snapshotted rather than joined so a renamed vendor cannot rewrite what a past statement
    says about itself, and so the uploads list renders without a join. `code` is the unique
    one — 141 names in the master repeat across branches.
    """
    return StatementBatchSupplier(
        tenant_id=tenant_id, slug=slug, batch_id=batch_id,
        supplier_id=supplier.id, supplier_name=supplier.name,
        supplier_code=supplier.code,
        supplier_branch=supplier.branch or supplier.city,
        agency_id=getattr(supplier, "agency_id", None),
    )


def _supplier_fields(supplier) -> dict:
    """The consolidator keys every upload/extract/confirm response carries."""
    return {
        "supplier_id": supplier.id if supplier else None,
        "supplier_name": supplier.name if supplier else None,
        "supplier_code": supplier.code if supplier else None,
        "supplier_branch": (supplier.branch or supplier.city) if supplier else None,
        "agency_id": getattr(supplier, "agency_id", None) if supplier else None,
    }


# The channel a third-party statement trades on — its agency must be onboarded on it.
_SLUG_CHANNEL = {"tp-gds": "GDS", "tp-lcc": "LCC"}


async def _agency_counterparty(db: AsyncSession, slug: str, agency_id: int, user: User):
    """The picked Agency Master row, as the counterparty a batch is linked to.

    Checked the way services/statement_agency_selection.py lays down: this user's own
    agency, active, and trading on the statement's channel (a GDS statement against an
    agency's LCC row would be priced against the wrong deal). An agency AWAITING approval
    is allowed: it has no supplier-master row yet, so the batch carries supplier_id None
    and matches deals by name — exactly as that agency's deals do — until approval
    back-fills it (api/v1/suppliers.py::approve_supplier).
    """
    from types import SimpleNamespace
    from app.models.supplier import Supplier
    from app.services import statement_agency_selection as agency_sel

    channel = _SLUG_CHANNEL.get(slug)
    try:
        agency = await agency_sel.resolve_for_upload(db, user, agency_id)
        if channel:
            agency_sel.resolve_agency_choice(agency, channel)
        elif agency is None:
            agency_sel.resolve_agency_choice(None, "")
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))

    supplier = None
    if agency.supplier_id is not None:
        supplier = (await db.execute(select(Supplier).where(Supplier.id == agency.supplier_id))).scalar_one_or_none()
    return SimpleNamespace(
        id=supplier.id if supplier else None,
        name=agency.name,
        code=(supplier.code if supplier else None) or agency.agency_code or agency.branch_code,
        branch=agency.branch_name or (supplier.branch if supplier else None),
        city=agency.city,
        agency_id=agency.id,
    )


async def _supplier_selection(db: AsyncSession, slug: str, supplier_id: int | None,
                              agency_id: int | None = None, user: User | None = None):
    """The Supplier master row this upload is attributed to, or None for types that don't
    need one.

    Required only for the third-party types, whose statements are issued BY a consolidator
    the file never names — see `requires_supplier` in services/statement_spec.py, which is
    what confines the requirement. An id sent for any other type is ignored rather than
    rejected, mirroring `_airline_selection`.

    No user or tenant scope: `suppliers` is global platform-admin master data, unlike
    `agencies`, which every user keeps their own copy of.
    """
    if not spec.requires_supplier(slug):
        return None
    if agency_id is not None and user is not None:
        return await _agency_counterparty(db, slug, agency_id, user)
    try:
        return await supplier_selection.resolve(db, supplier_id)
    except ValueError as exc:
        # UnknownSupplier / InactiveSupplier are both ValueErrors and both carry a message
        # written for the user.
        raise HTTPException(status_code=400, detail=str(exc))


def _period_form(slug: str, period_from: str | None, period_to: str | None) -> tuple[date, date] | None:
    """The statement period an upload declares, checked — for the types that capture one
    (statement_spec.captures_period). None for every other type, whatever was sent.

    Required, not optional: the tickets' Booked and Issue Dates are checked against it on
    the review step, and the wizard asks before reading a file whose period is already
    uploaded for the agency (services/tp_gds_reupload.py)."""
    if not spec.captures_period(slug):
        return None
    if not (period_from or "").strip() or not (period_to or "").strip():
        raise HTTPException(
            status_code=400,
            detail="Enter the statement period (From and To) — the tickets' dates are checked "
                   "against it.")
    try:
        start = date.fromisoformat(period_from.strip()[:10])
        end = date.fromisoformat(period_to.strip()[:10])
    except ValueError:
        raise HTTPException(status_code=400, detail="The statement period must be two dates (YYYY-MM-DD).")
    if start > end:
        raise HTTPException(status_code=400, detail="The statement period's From date is after its To date.")
    if (end - start).days > 366:
        raise HTTPException(status_code=400, detail="A statement period cannot be longer than a year.")
    return start, end


def _date_formats_form(slug: str, raw: str | None, *, required: bool) -> dict[str, str] | None:
    """The order the uploader declared each date column is written in — {field: "DMY" |
    "MDY" | "YMD"} for the types that ask (statement_spec.date_format_fields), checked. None
    for every other type, whatever was sent.

    `required` wherever lines are read to be saved or previewed (`/save-preview`, `/confirm`,
    `/upload`): it decides how every date in those columns is read — text in that order, and
    Excel-stored dates put in it (`_read_stored_dates`) — and it is stored with the upload.
    There is no day-first fallback for these types. Only `/extract`, which saves nothing,
    takes it as optional."""
    fields = spec.date_format_fields(slug)
    if not fields:
        return None
    if not (raw or "").strip():
        if not required:
            return None
        raise HTTPException(
            status_code=400,
            detail=f"Choose the date format of {' and '.join(f['header'] for f in fields)} — "
                   f"it decides how the file's dates are read.")
    try:
        sent = json.loads(raw)
    except ValueError:
        sent = None
    if not isinstance(sent, dict):
        raise HTTPException(status_code=400, detail="The date formats were not readable.")
    out: dict[str, str] = {}
    for f in fields:
        order = sent.get(f["field"])
        if order not in _flat.DATE_ORDERS:
            raise HTTPException(
                status_code=400,
                detail=f"Choose the date format of {f['header']} (DD-MM-YYYY, MM-DD-YYYY or YYYY-MM-DD).")
        out[f["field"]] = order
    return out


def _read_stored_dates(slug: str, content: bytes, filename: str, df: pd.DataFrame,
                       period: tuple[date, date] | None, colmap: dict | None,
                       formats: dict | None) -> excel_dates.DateFix | None:
    """Third Party GDS: the Excel-stored dates of the declared date columns, put in the order
    the sheet is written in (services/excel_dates.py) — "01-08-2026" typed day-first into a
    month-first Excel and stored as 8 Jan is put back to 1 Aug, row by row against the period.
    `colmap` ({field: column}) says which columns hold the dates, `formats` their declared
    order — None only while the file is still being read for it (`/detect-date-formats`).

    Applied by every reader of the file — `/extract`, `/save-preview`, `/confirm`, `/upload`,
    `/detect-date-formats` — so all of them see the same dates. None, and `df` untouched, for
    every other type, any non-Excel file, and a sheet entered the way its Excel reads."""
    if period is None or not spec.date_format_fields(slug):
        return None
    if spreadsheet.sniff(content, filename) not in (spreadsheet.XLS, spreadsheet.XLSX):
        return None
    columns = {f: c for f, c in (colmap or {}).items() if f in _flat.DATE_FIELDS}
    orders = _flat.date_orders(formats) if formats else None
    return excel_dates.read_stored_dates(df, period, columns, orders)


def _period_fields(period: tuple[date, date] | None) -> dict:
    return {
        "period_from": period[0].isoformat() if period else None,
        "period_to": period[1].isoformat() if period else None,
    }


def _agency_lock(user: User, supplier) -> int:
    """A Postgres advisory-lock key for one agency's statements in one tenant (a signed
    bigint, stable across processes — Python's `hash()` is salted per process)."""
    who = getattr(supplier, "agency_id", None) or f"s{getattr(supplier, 'id', None)}"
    raw = hashlib.blake2b(f"tp-gds:{user.tenant_id}:{who}".encode(), digest_size=8).digest()
    return int.from_bytes(raw, "big", signed=True)


def _duplicate_choices_form(slug: str, raw: str | None) -> dict | None:
    """What Save does with each bucket of duplicates, and row by row (Third Party GDS —
    services/tp_gds_reupload.normalize_choices), checked; the defaults when none is sent.
    None for every other type, whatever was sent."""
    if slug != tp_gds_reupload.SLUG:
        return None
    try:
        sent = json.loads(raw) if (raw or "").strip() else {}
        if not isinstance(sent, dict):
            raise ValueError("The duplicate choices were not readable.")
        return tp_gds_reupload.normalize_choices(sent)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=f"Duplicate choices: {exc}")


async def _screen_duplicates(db: AsyncSession, slug: str, user: User, supplier,
                             period: tuple[date, date] | None, objs: list, *,
                             key_data: list | None = None, indexes: list[int] | None = None,
                             choices: dict | None = None):
    """Third Party GDS: screen the file against every line already saved for this agency, and
    keep what the user chose to save (services/tp_gds_reupload.py) — `(kept, screen,
    decision)`. `(objs, None, None)`, unchanged, for every other type.

    Refused outright when no entry is left to save: an upload of nothing but duplicates
    would put an empty statement in the list. Checked before anything is written or stored,
    so nothing is."""
    if slug != tp_gds_reupload.SLUG:
        return objs, None, None
    if supplier is not None:
        # Two saves of one agency's statements at the same moment would each find the other
        # not yet there and both save in full. One at a time, per agency, until the commit.
        await db.execute(select(func.pg_advisory_xact_lock(_agency_lock(user, supplier))))
    screen = await tp_gds_reupload.compare(db, user, supplier, period, objs,
                                           key_data=key_data, indexes=indexes)
    decision = screen.decide(choices)
    reason = screen.refusal(decision)
    if reason:
        raise HTTPException(status_code=409, detail=f"{reason} Nothing was saved.")
    return decision.kept, screen, decision


async def _write_lines(db: AsyncSession, slug: str, user: User, objs: list, *, batch_id: str,
                       source_file: str | None, file_url: str | None, uploaded_at: datetime,
                       period: tuple[date, date] | None = None,
                       decision: tp_gds_reupload.Decision | None = None,
                       date_formats: dict[str, str] | None = None) -> None:
    """Add an upload's lines — after its header, for the types that have one
    (models/third_party_statement.py), with its statement period on the header and on every
    line when it declares one, its declared date formats on the header, and its tally
    counting the lines it left out as duplicates (Third Party GDS). Shared by `/upload` and
    `/confirm`."""
    if slug in THIRD_PARTY_SLUGS:
        new = sum(1 for o in objs
                  if not (getattr(o, "data", None) or {}).get(statement_balance.ROW_KIND_KEY))
        repeats = len(decision.left_out) if decision is not None else 0
        db.add(ThirdPartyStatement(
            batch_id=batch_id, tenant_id=user.tenant_id, created_by_id=user.id, slug=slug,
            source_file=source_file, file_url=file_url, uploaded_at=uploaded_at,
            period_from=period[0] if period else None, period_to=period[1] if period else None,
            date_formats=date_formats,
            total_entries=new + repeats, duplicate_entries=repeats, new_entries=new,
        ))
        # Flushed on its own, ahead of the lines whose batch_id is a foreign key to it, so
        # the INSERT order never rests on how the unit of work sorts two mappers that have
        # no relationship() between them.
        await db.flush()
    if period is not None:
        # Stamped here rather than in the shared provenance: only the types that capture a
        # period have the columns (MO GDS shares the parser and has none).
        for o in objs:
            o.period_from, o.period_to = period
    db.add_all(objs)


def _upload_facts(slug: str, model, batch_id: str, user: User):
    """`(entity, conds)` to read one upload's own facts — file name, stored copy, upload
    time — from: its header for the third-party types, any one of its lines for the rest,
    which still carry them on every line."""
    if slug in THIRD_PARTY_SLUGS:
        H = ThirdPartyStatement
        return H, [H.batch_id == batch_id, H.slug == slug, *_scope(H, user)]
    return model, [model.batch_id == batch_id, *_scope(model, user)]


async def _store_original(content: bytes, file, slug: str, batch_id: str, user: User) -> str | None:
    """Keep the uploaded file in GCS. Best-effort on purpose: losing the ability to
    re-download the original is worth far less than losing a successful import, so a
    storage failure returns None instead of raising."""
    try:
        from app.services import gcs
        blob_name = f"statements/{user.tenant_id}/{slug}/{batch_id}/{file.filename}"
        await gcs.upload_bytes(content, blob_name,
                               file.content_type or "application/octet-stream", _bucket())
        return blob_name
    except Exception:  # noqa: BLE001
        return None


# ── Map · review · confirm ──────────────────────────────────────────────────
# A consolidator writes whatever spreadsheet it likes. `/upload` handles the shapes the
# alias map already knows — including our own template — and these three handle the ones it
# does not, by letting the uploader say which of their columns is which.
#
# THE FILE IS SENT TWICE, deliberately, and there is no staging table. `/extract` reads it
# to propose a mapping; `/confirm` reads it again to write the rows. The alternative — hold
# the parsed file server-side between the two — needs somewhere to put it, a way to expire
# it, and a header row that exists before `/confirm`. The third-party header
# (models/third_party_statement.py) is written at `/confirm`, with its lines, and NDC has
# none. A consolidator statement is a weekly file of a few hundred lines, so the
# second read costs nothing, and `file_digest` is what stops the second file being a
# different one from the first.
_PREVIEW_ROWS = 1000


def _mapping_capable(slug: str):
    """The mapper for a type that supports the wizard, or 400.

    Two kinds answer to the same three methods. A `parser` type (third-party) gets its
    normalizing builder; a spec-driven type (NDC) gets a mapper built from its own ordered
    columns and aliases — see services/spec_mapping.py for why they are not the same
    object. Everything between here and the row-writing branch in `/confirm` is shared.
    """
    if not spec.supports_mapping(slug):
        raise HTTPException(
            status_code=400,
            detail=f"{spec.spec_for(slug)['label']} statements come from a system with a fixed "
                   f"export format, so there is nothing to map. Upload the file directly.",
        )
    parser_name = spec.parser(slug)
    if not parser_name:
        return spec_mapping.mapper_for(slug)
    builder = _get_builder(parser_name)
    if not hasattr(builder, "build_row_mapped"):
        raise HTTPException(status_code=500, detail="This type's parser cannot take a column map.")
    return builder


def _digest(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def _read_for_mapping(content: bytes, filename: str, builder, header_row: int | None):
    """(df, columns, header_row).

    The header row is DETECTED at extract and PINNED at confirm. Re-detecting it on the
    second read is the trap: a mapping is a list of column NAMES, so a different header row
    would silently rename every column and drop the whole mapping on the floor.

    Detection is by recognised-column count, like `_detect_df`, but it returns the INDEX as
    well as the frame — which is the only reason this exists rather than calling that.
    """
    if header_row is None:
        best = None
        last_err: Exception | None = None
        for hr in (0, 1, 2):
            try:
                df = _read_df(content, filename, hr)
            except Exception as e:  # noqa: BLE001 — unreadable at this offset, try the next
                last_err = e
                continue
            df.dropna(how="all", inplace=True)
            score = len(builder.build_col_map([str(c) for c in df.columns]))
            if best is None or score > best[0]:
                best = (score, df, hr)
        if best is None:
            raise HTTPException(
                status_code=400,
                detail=f"Could not read the file: {last_err}. Upload a valid .xlsx, .xls or .csv.")
        _score, df, header_row = best
    else:
        try:
            df = _read_df(content, filename, header_row)
        except Exception as exc:  # noqa: BLE001
            raise HTTPException(status_code=400, detail=f"Could not re-read the file: {exc}")
        df.dropna(how="all", inplace=True)
    return df, [str(c) for c in df.columns], header_row


@router.get("/{slug}/standard-columns")
async def standard_columns(slug: str, current_user: User = Depends(get_current_user)):
    """The fields a file can be mapped ONTO, grouped for the mapping screen.

    `required` is what the API will refuse without; `advisory` is what it will accept but
    warn about — see flat_statement.REQUIRED_GROUPS for why those two lists differ.
    """
    slug, _model = _resolve(slug)
    _mapping_capable(slug)
    parser_name = spec.parser(slug)
    return {
        "groups": _flat.column_groups(parser_name) if parser_name else spec_mapping.column_groups(slug),
        "required": spec.required_groups(slug),
        "advisory": spec.advisory_groups(slug),
        "total": len(spec.columns(slug)),
        # Null for every type but NDC. Present so the review step can grey out the rows the
        # import is going to skip and say why, instead of the user meeting the shortfall as
        # an unexplained entry count. The rule is enforced at /confirm either way.
        "row_filter": spec.row_filter(slug),
        # Third Party GDS: the upload step asks for the statement period, says when it is
        # already uploaded, and the review step asks `/save-preview` what Save would leave out.
        "captures_period": spec.captures_period(slug),
        # Third Party GDS: the date columns whose written order the upload step asks for
        # ([{field, header}]); [] for every other type.
        "date_formats": spec.date_format_fields(slug),
    }


@router.get("/{slug}/last-date-formats")
async def last_date_formats(
    slug: str,
    supplier_id: int | None = Query(default=None),
    agency_id: int | None = Query(default=None),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """The date formats this agency's latest upload declared, to pre-fill the upload step —
    a consolidator writes its dates the same way statement after statement. Nulls when the
    agency has none. Writes nothing."""
    slug, _model = _resolve(slug)
    none = {"date_formats": None, "source_file": None, "uploaded_at": None}
    if not spec.date_format_fields(slug) or slug not in THIRD_PARTY_SLUGS:
        return none
    supplier = await _supplier_selection(db, slug, supplier_id, agency_id, current_user)
    vendor = tp_gds_reupload.vendor_condition(supplier)
    if vendor is None:
        return none
    H, L = ThirdPartyStatement, StatementBatchSupplier
    row = (await db.execute(
        select(H.date_formats, H.source_file, H.uploaded_at)
        .join(L, (L.batch_id == H.batch_id) & (L.slug == H.slug) & (L.tenant_id == H.tenant_id))
        .where(*_scope(H, current_user), H.slug == slug, vendor, H.date_formats.isnot(None))
        .order_by(H.uploaded_at.desc(), H.batch_id.desc())
        .limit(1)
    )).first()
    if row is None:
        return none
    return {"date_formats": row.date_formats, "source_file": row.source_file,
            "uploaded_at": row.uploaded_at.isoformat() if row.uploaded_at else None}


@router.post("/{slug}/detect-date-formats")
async def detect_date_formats(
    slug: str,
    file: UploadFile = File(...),
    # The statement period, when the upload step has one: it is what settles a date like
    # 05/08/2026, which is a real date both ways round. Without it only cells that can be
    # read one way alone ("13/08/2026") count.
    period_from: str | None = Form(default=None),
    period_to: str | None = Form(default=None),
    current_user: User = Depends(get_current_user),
):
    """The order each declared date column of the file is written in, as far as the file
    shows it (services/date_format_detect.py) — to fill in the upload step's Date format.
    Read through the columns `/extract` would map them to. Writes nothing."""
    slug, _model = _resolve(slug)
    fields = spec.date_format_fields(slug)
    if not fields:
        return {"period_label": None, "fields": []}
    builder = _mapping_capable(slug)
    try:
        period = _period_form(slug, period_from, period_to) if period_from and period_to else None
    except HTTPException:
        period = None       # a period still being typed: detect without it

    content = await file.read()
    max_bytes = settings.MAX_UPLOAD_SIZE_MB * 1024 * 1024
    if len(content) > max_bytes:
        raise HTTPException(status_code=413, detail=f"File exceeds {settings.MAX_UPLOAD_SIZE_MB} MB limit.")
    if not content:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")
    df, columns, _header_row = _read_for_mapping(content, file.filename or "", builder, None)
    suggested = builder.suggest_mapping(columns)
    # No formats yet — they are what this finds out — so every date column may be put back.
    fix = _read_stored_dates(slug, content, file.filename or "", df, period, suggested, None)
    found = date_format_detect.detect(df, {f["field"]: suggested.get(f["field"]) for f in fields}, period, fix)
    headers = {f["field"]: f["header"] for f in fields}
    return {
        "period_label": statement_period.label(*period) if period else None,
        "fields": [{**d.as_dict(), "header": headers[d.field]} for d in found],
    }


@router.get("/{slug}/period-guess")
async def guess_period(
    slug: str,
    file_name: str = Query(..., max_length=500),
    current_user: User = Depends(get_current_user),
):
    """The statement period a file's name gives — "…01-07TH AUG 26.xlsx" is 01–07 Aug 2026 —
    to pre-fill the upload's Statement period; nulls when the name gives none, and the user
    types it (services/statement_period.py). Writes nothing."""
    slug, _model = _resolve(slug)
    found = statement_period.from_file_name(file_name, date.today())
    return {**_period_fields(found), "period_label": statement_period.label(*found) if found else None}


@router.get("/{slug}/period-uploads")
async def period_uploads(
    slug: str,
    period_from: str = Query(...),
    period_to: str = Query(...),
    supplier_id: int | None = Query(default=None),
    agency_id: int | None = Query(default=None),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """This agency's uploads whose period shares a day with the one given — what the wizard
    names in "already uploaded — Continue / Cancel" before the file is read. Only the
    question: every upload is screened against all of the agency's saved lines whatever the
    answer (services/tp_gds_reupload.py). Empty for every type that does not capture a
    period. Writes nothing."""
    slug, _model = _resolve(slug)
    period = _period_form(slug, period_from, period_to)
    if period is None or slug != tp_gds_reupload.SLUG:
        return {"uploads": []}
    supplier = await _supplier_selection(db, slug, supplier_id, agency_id, current_user)
    earlier = await tp_gds_reupload.earlier_uploads(db, current_user, supplier, *period)
    return {**_period_fields(period), "period_label": statement_period.label(*period),
            "uploads": [e.as_dict() for e in earlier]}


@router.post("/{slug}/extract")
async def extract_statement(
    slug: str,
    file: UploadFile = File(...),
    supplier_id: int | None = Form(default=None),
    # The Agency Master row picked on upload. Preferred over supplier_id: it is
    # channel-checked, works for agencies awaiting approval, and carries the vendor
    # service charge the SF columns are checked against.
    agency_id: int | None = Form(default=None),
    # The statement period and the declared date formats (Third Party GDS), when the wizard
    # has them: the Excel-stored dates are put in the sheet's order against them
    # (`_read_stored_dates`), as `/confirm` will read them.
    period_from: str | None = Form(default=None),
    period_to: str | None = Form(default=None),
    date_formats: str | None = Form(default=None),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Read the file, propose a mapping, and hand back rows to review. Writes nothing.

    The supplier is validated HERE as well as at confirm, so a wrong consolidator is caught
    before the user spends time on a mapping they cannot save.
    """
    slug, _model = _resolve(slug)
    builder = _mapping_capable(slug)
    period = _period_form(slug, period_from, period_to) if period_from and period_to else None
    formats = _date_formats_form(slug, date_formats, required=False)
    supplier = await _supplier_selection(db, slug, supplier_id, agency_id, current_user)

    content = await file.read()
    max_bytes = settings.MAX_UPLOAD_SIZE_MB * 1024 * 1024
    if len(content) > max_bytes:
        raise HTTPException(status_code=413, detail=f"File exceeds {settings.MAX_UPLOAD_SIZE_MB} MB limit.")
    if not content:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")

    df, columns, header_row = _read_for_mapping(content, file.filename or "", builder, None)
    if not columns:
        raise HTTPException(status_code=400, detail="No columns were found in the file.")
    suggested = builder.suggest_mapping(columns)
    date_fix = _read_stored_dates(slug, content, file.filename or "", df, period, suggested, formats)
    # Null for every type but Third Party API. Shown on the mapping step as reassurance, and
    # deliberately NOT offered as a dropdown: for a spec-driven type the mapping IS the
    # parser, so overriding this would write a provenance claim that changes nothing and
    # contradicts what the user is looking at. Null here means "vendor not recognised",
    # which the wizard says plainly — the file still imports.
    detect = spec.detect_format(slug)
    src_fmt = detect(columns) if detect else None
    # Rows are returned by POSITION, and confirm enumerates identically, so an edit made
    # against row 7 here lands on row 7 there.
    sample = []
    for i, (_, row) in enumerate(df.iterrows()):
        if i >= _PREVIEW_ROWS:
            break
        sample.append({"__index__": i,
                       **{c: _sample_cell(row[c]) for c in columns}})

    return {
        "file_name": file.filename,
        "file_digest": _digest(content),
        "header_row": header_row,
        "total_rows": int(len(df)),
        "preview_rows": len(sample),
        "preview_limit": _PREVIEW_ROWS,
        "columns": columns,
        "suggested_mapping": suggested,
        "matched_columns": len(suggested),
        "standard_total": len(spec.columns(slug)),
        # A file built from our own Template comes back fully mapped; the wizard says so
        # rather than making the user check 56 dropdowns to find out.
        "is_template_match": len(suggested) >= len(spec.columns(slug)) - 2,
        "source_format": src_fmt,
        "source_format_label": spec.format_label(slug, src_fmt),
        "detects_format": detect is not None,
        "ai_narration": spec.ai_narration(slug),
        "sample_rows": sample,
        # Third Party GDS: the dates of a sheet typed day-first that were read the other way
        # round, cell by cell, so the review step can mark them. Null when none were.
        "date_fix": date_fix.as_dict(rows_shown=len(sample)) if date_fix else None,
        **_supplier_fields(supplier),
    }


class AiParseRow(BaseModel):
    index: int                       # the wizard's own `__index__`
    text: str
    product_type: str | None = None  # the mapped hint; far better than making the model guess


class AiParseRequest(BaseModel):
    rows: list[AiParseRow] = Field(default_factory=list)
    source_format: str | None = None


@router.post("/{slug}/ai-parse")
async def ai_parse_narration(
    slug: str,
    body: AiParseRequest,
    current_user: User = Depends(get_current_user),
):
    """Read the free-text column and return the fields it contains. Writes NOTHING.

    THE ROWS ARE POSTED, NOT THE FILE, and that is not an optimisation. The parse has to run
    against the mapping the user is looking at — which column is the narration is client
    state that exists only between `/extract` and `/confirm` — and the result has to key
    back onto `__index__` so it can be merged into the review step's per-cell edits. Sending
    the file instead would mean also sending `column_map`, `header_row` and `file_digest`,
    i.e. re-implementing `/confirm` to read one column, and would make this a third transfer
    of a file this router deliberately moves exactly twice (see `_store_original`).

    Because the result lands in that edit map, `/confirm` needs no knowledge of any of this:
    an AI-proposed value is reviewed, editable and overridable exactly like a typed one.
    """
    slug, _model = _resolve(slug)
    cfg = spec.ai_narration(slug)
    if not cfg:
        raise HTTPException(
            status_code=400,
            detail=f"{spec.spec_for(slug)['label']} statements have no free-text column to "
                   f"read — every field on them comes from a column of its own.",
        )
    if not settings.OPENAI_API_KEY:
        raise HTTPException(
            status_code=503,
            detail="AI parsing is not configured on this server (OPENAI_API_KEY is unset). "
                   "The narration column still imports as text — only the fields inside it "
                   "would have to be filled in by hand.",
        )
    if not body.rows:
        raise HTTPException(status_code=400, detail="No rows were sent to parse.")

    from app.services import ai_client, ai_statement_narration

    try:
        return await ai_statement_narration.parse_rows(
            [{"index": r.index, "text": r.text, "product_type": r.product_type}
             for r in body.rows]
        )
    except ai_client.FatalAIError as exc:
        # Auth/permission: every retry reaches the same answer, so say so once rather than
        # spending four attempts per chunk to get there.
        raise HTTPException(
            status_code=502,
            detail=f"The AI provider rejected the request: {exc}",
        ) from exc


@dataclass
class _MappingRequest:
    colmap: dict
    row_edits: dict
    content: bytes
    df: pd.DataFrame
    columns: list
    header_row: int
    # {date field: declared order} — Third Party GDS (`_date_formats_form`); None elsewhere.
    date_formats: dict | None = None


async def _read_mapping_request(slug: str, builder, file: UploadFile, column_map: str,
                                header_row: int, file_digest: str, edits: str,
                                period: tuple[date, date] | None = None,
                                date_formats: dict | None = None) -> _MappingRequest:
    """A `/confirm`-shaped request's mapping, edits and file — checked, and the file re-read
    under the header row `/extract` pinned, its Excel-stored dates put in the declared order
    against `period` (`_read_stored_dates`). Shared with `/save-preview`."""
    try:
        colmap = json.loads(column_map or "{}")
        row_edits = {int(k): v for k, v in json.loads(edits or "{}").items()}
    except (ValueError, TypeError, AttributeError):
        raise HTTPException(status_code=400, detail="The column map or edits were not readable.")
    if not isinstance(colmap, dict) or not colmap:
        raise HTTPException(status_code=400, detail="Map at least one column before saving.")

    missing = [g for g in spec.required_groups(slug) if not any(colmap.get(f) for f in g["fields"])]
    if missing:
        detail = " and ".join(f"{g['label']} ({g['headers']})" for g in missing)
        raise HTTPException(
            status_code=400,
            detail=f"Map {detail} before saving — without it these rows cannot be found or "
                   f"reconciled, so importing them would be a mis-mapping rather than a choice.",
        )

    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")
    if _digest(content) != file_digest:
        raise HTTPException(
            status_code=400,
            detail="This is not the file the mapping was built against. Start the upload again "
                   "so the columns are read from the file you are saving.",
        )
    df, columns, hr = _read_for_mapping(content, file.filename or "", builder, header_row)
    _read_stored_dates(slug, content, file.filename or "", df, period, colmap, date_formats)
    return _MappingRequest(colmap, row_edits, content, df, columns, hr, date_formats)


@dataclass
class _MappedLines:
    objs: list
    source_rows: int
    edited_rows: int
    excluded_rows: int
    airline_rows: int | None    # None for the types that resolve no carrier
    # Parallel to `objs`: the file row (`__index__`) each came from, and — Third Party GDS,
    # rows the review step corrected — the line as the file had it, which is what its key
    # is built from (services/tp_gds_reupload.py). None where the line is its own source.
    indexes: list[int] = field(default_factory=list)
    key_data: list[dict | None] = field(default_factory=list)


async def _mapped_lines(db: AsyncSession, slug: str, model, builder, req: _MappingRequest,
                        prov: dict) -> _MappedLines:
    """The statement lines a mapped file makes, the reviewer's edits applied — the rows
    `/confirm` writes and `/save-preview` compares. Refuses a file that makes none."""
    df, columns, colmap, row_edits = req.df, req.columns, req.colmap, req.row_edits
    stamp_airline = await _airline_stamper(db, slug)
    # The two row shapes this router writes. A `parser` type is normalized and carries
    # segments/ssr/raw_data/source_format; a spec-driven type is a verbatim copy of the
    # vendor's columns and its table has none of those. Everything before this point —
    # the mapping, the digest guard, the edits — is identical for both.
    normalized = spec.parser(slug) is not None
    fold = spec.fold_taxes(slug)
    normmap = {} if normalized else {c: spec.norm(c) for c in columns}
    drop, rowfilter = spec.drop_row(slug), spec.row_filter(slug)
    derive = spec.derive_row(slug)
    normalize, stamp_fmt = spec.normalize_row(slug), spec.stamp_format(slug)
    # The vendor is a property of the FILE, not of a line, so it is detected once from the
    # header row. None for every type that does not declare `detect_format`, and None for a
    # declaring type whose file matches no known vendor — which still imports, just without
    # a vendor label.
    detect = spec.detect_format(slug)
    src_fmt = detect(columns) if detect else None

    # A period uploaded again is matched line by line on the line as the VENDOR sent it, so a
    # row corrected on the review step is keyed on its uncorrected reading.
    keyed = slug == tp_gds_reupload.SLUG
    # The declared date formats, for the builder that reads them (only the flat one is
    # asked to — spec-driven mappers have no such argument and no type of theirs declares).
    dated = {"date_formats": req.date_formats} if req.date_formats else {}

    objs, airline_rows, edited_rows, source_rows, excluded_rows = [], 0, 0, 0, 0
    indexes: list[int] = []
    key_data: list[dict | None] = []
    for i, (_, row) in enumerate(df.iterrows()):
        overrides = row_edits.get(i)
        b = builder.build_row_mapped(row, columns, colmap, overrides, **dated)
        data = b["data"]
        sent = None
        if keyed and overrides:
            sent = builder.build_row_mapped(row, columns, colmap, None, **dated)["data"]
            if derive is not None:
                extra = derive({n: _clean(row.get(c)) for c, n in normmap.items()})
                if extra:
                    sent = {**extra, **sent}
            if stamp_airline is not None:
                stamp_airline(sent)
        # Derived fields read the WHOLE line, not the mapping — same reason as the tax fold
        # below: NDC's `Other Taxes` is eleven of the airline's columns added up, and no
        # one-field-to-one-column mapping can express that. Spread FIRST so a column the
        # user did map, or corrected on the review step, wins over the computed value.
        if derive is not None:
            extra = derive({n: _clean(row.get(c)) for c, n in normmap.items()})
            if extra:
                data = {**extra, **data}
        # Taxes are folded from the WHOLE row, not from the mapping: a Tax_TypeN/TaxN pair
        # is not a mappable field, and folding here keeps the mapped path and the verbatim
        # `/upload` path producing the same `taxes` array.
        # Rewrite the FINAL row — after the mapping, after `derive`, and after the
        # reviewer's overrides. That order is what makes the review step's promise true:
        # a cell corrected to "13-08-2026" or "1,23,456.78" is normalised exactly as the
        # file's own value was. None for every type but Third Party API.
        if normalize is not None:
            if stamp_fmt is not None:
                stamp_fmt(data, src_fmt)
            normalize(data)
        taxes = b["taxes"] if normalized else (_fold_taxes(row, normmap) if fold else [])
        if not data and not taxes:
            continue    # every mapped column empty on this line
        # A row this type does not keep — NDC's unpaid holds and free seats. Checked AFTER
        # the edits are applied, so correcting a mis-read TXN Type in the review step
        # decides the row's fate, and counted separately from `source_rows` so the response
        # can say the file had 106 lines and 78 of them were transactions.
        if drop is not None and drop(data):
            excluded_rows += 1
            continue
        source_rows += 1
        if overrides:
            edited_rows += 1
        if stamp_airline is not None and stamp_airline(data):
            airline_rows += 1
        if normalized:
            objs.append(_line(
                model, **prov, data=data, taxes=taxes, segments=b["segments"], ssr=b["ssr"],
                raw_data=b["raw_data"], source_format=b["source_format"],
            ))
            indexes.append(i)
            # Read back through the model, as the corrected line is, so both keys see the
            # values exactly as a stored line holds them.
            key_data.append(model(data=sent).data if sent is not None and is_typed(model) else sent)
        else:
            # Shared with `/upload`, so a mapped import and a direct one produce identical
            # rows — including the leg/total flags for the types that split.
            made = _build_rows(
                model, slug, prov, data, taxes, source_rows,
                extra=_fmt_extra(model, detect, src_fmt, row, columns),
            )
            objs.extend(made)
            indexes.extend([i] * len(made))
            key_data.extend([None] * len(made))

    if not objs:
        # Two different failures with the same symptom, and telling them apart is the whole
        # point: a bad mapping is the user's to fix, a file of nothing but unpaid holds is
        # not — re-mapping it would not produce a single row.
        if excluded_rows and rowfilter:
            raise HTTPException(
                status_code=400,
                detail=f"Every row in this file is {rowfilter['label']} "
                       f"({', '.join(rowfilter['exclude'])}), so there is nothing to import. "
                       f"{rowfilter['note']}")
        raise HTTPException(
            status_code=400,
            detail="Every row came out empty under this mapping. Check that the mapped columns "
                   "are the ones holding the data.")
    return _MappedLines(objs, source_rows, edited_rows, excluded_rows,
                        airline_rows if stamp_airline is not None else None,
                        indexes=indexes, key_data=key_data)


_ISO_DAY = re.compile(r"^\d{4}-\d{2}-\d{2}$")
# Unreadable cells listed back per column. The counts always cover every line.
_MAX_DATE_CELLS = 1000


def _source_cell(req: _MappingRequest, index: int, field: str) -> str | None:
    """What a line's `field` was read from: the reviewer's correction, else the mapped cell."""
    edited = (req.row_edits.get(index) or {}).get(field)
    if edited is not None:
        return _clean(edited)
    col = req.colmap.get(field)
    if not col or col not in req.df.columns:
        return None
    return _clean(req.df.iloc[index][col])


def _dates_read(slug: str, req: _MappingRequest, lines: _MappedLines) -> list[dict]:
    """How each declared date column read in the format chosen for it — for the review step,
    so a wrong choice shows before Save rather than as tickets filed under the wrong month.

    Per column: how many dates were read and the span they cover (a statement's Booked Dates
    sit within days of each other; read the wrong way round they scatter across the year),
    and the cells that could not be read in that format at all, by row with their value —
    those save without the date. Balance lines carry no dates and are passed over. [] for
    the types that declare no date formats."""
    out = []
    for f in spec.date_format_fields(slug):
        field = f["field"]
        read, first, last, unreadable, cells = 0, None, None, 0, []
        for obj, index in zip(lines.objs, lines.indexes):
            data = obj.data or {}
            if data.get(statement_balance.ROW_KIND_KEY):
                continue
            value = data.get(field)
            iso = str(value)[:10] if value not in (None, "") else ""
            if _ISO_DAY.match(iso):
                read += 1
                first = iso if first is None or iso < first else first
                last = iso if last is None or iso > last else last
                continue
            # Not read: a JSONB line keeps the cell as it was; a typed line stores no value
            # and names the field in `unparsed_fields`.
            if value in (None, "") and field not in str(data.get(UNPARSED_KEY) or "").split(","):
                continue
            unreadable += 1
            if len(cells) < _MAX_DATE_CELLS:
                cells.append({"index": index,
                              "value": str(value) if value not in (None, "") else _source_cell(req, index, field)})
        out.append({
            "field": field, "header": f["header"],
            "format": (req.date_formats or {}).get(field),
            "mapped": bool(req.colmap.get(field)),
            "read": read, "first": first, "last": last,
            "unreadable": unreadable, "cells": cells,
        })
    return out


def _row_marks(lines: _MappedLines) -> dict:
    """Which file rows are the statement's own balance lines, and which hold a value that could
    not be read into its field (stored without it) — so the review step can count them and
    filter its grid to them. By file row (`__index__`)."""
    balance, unparsed = [], []
    for obj, index in zip(lines.objs, lines.indexes):
        data = obj.data or {}
        if data.get(statement_balance.ROW_KIND_KEY):
            balance.append(index)
            continue
        # A typed line names them in `unparsed_fields`; a JSONB line keeps its first bad date.
        fields = [f for f in str(data.get(UNPARSED_KEY) or "").split(",") if f]
        if not fields and data.get("date_parse_failed"):
            fields = [data["date_parse_failed"]]
        if fields:
            unparsed.append({"index": index, "fields": fields})
    return {"balance_rows": balance, "unparsed_rows": unparsed}


@router.post("/{slug}/confirm")
async def confirm_statement(
    slug: str,
    file: UploadFile = File(...),
    column_map: str = Form(...),
    supplier_id: int | None = Form(default=None),
    # The Agency Master row picked on upload. Preferred over supplier_id: it is
    # channel-checked, works for agencies awaiting approval, and carries the vendor
    # service charge the SF columns are checked against.
    agency_id: int | None = Form(default=None),
    header_row: int = Form(...),
    file_digest: str = Form(...),
    edits: str = Form(default="{}"),
    # Optional control totals the uploader declares (the vendor's covering note) — kept for
    # the completeness check on the types that capture controls, ignored by the rest.
    expected_count: int | None = Form(default=None),
    expected_amount: str | None = Form(default=None),
    # The statement period (YYYY-MM-DD) — required for the types that capture one
    # (statement_spec.captures_period), ignored by the rest.
    period_from: str | None = Form(default=None),
    period_to: str | None = Form(default=None),
    # JSON {date field: "DMY" | "MDY" | "YMD"} — required for the types that declare their
    # date formats (statement_spec.date_format_fields), ignored by the rest.
    date_formats: str | None = Form(default=None),
    # JSON — what to do with the duplicates the review step showed (Third Party GDS):
    # {"exact" | "partial" | "in_file": "save" | "skip", "rows": {file row: "save" | "skip"}}.
    # The defaults (services/tp_gds_reupload.DEFAULT_CHOICES) when absent; ignored elsewhere.
    duplicate_choices: str | None = Form(default=None),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Apply the reviewed mapping and edits, and write the batch.

    `file_digest` must match the file `/extract` read. Without it a mapping built against
    one spreadsheet could be applied to another — same column count, different columns —
    and every value would land in the wrong field with nothing to show for it.

    Third Party GDS: the file is screened against every line already saved for the agency,
    and the duplicates are saved or left out as the user chose (services/tp_gds_reupload.py).
    """
    slug, model = _resolve(slug)
    builder = _mapping_capable(slug)
    period = _period_form(slug, period_from, period_to)
    formats = _date_formats_form(slug, date_formats, required=True)
    choices = _duplicate_choices_form(slug, duplicate_choices)
    supplier = await _supplier_selection(db, slug, supplier_id, agency_id, current_user)
    req = await _read_mapping_request(slug, builder, file, column_map, header_row, file_digest, edits,
                                      period, formats)
    expected_amt = _to_decimal_form(expected_amount, "Expected net amount")

    batch_id = str(uuid.uuid4())
    uploaded_at = datetime.utcnow()
    prov = dict(
        tenant_id=current_user.tenant_id, created_by_id=current_user.id,
        batch_id=batch_id, source_file=file.filename, file_url=None,
        uploaded_at=uploaded_at,
    )
    lines = await _mapped_lines(db, slug, model, builder, req, prov)

    objs, screen, decision = await _screen_duplicates(
        db, slug, current_user, supplier, period, lines.objs,
        key_data=lines.key_data, indexes=lines.indexes, choices=choices)
    # Stored only once the upload is going ahead: a file refused as nothing but duplicates
    # must not leave an orphaned copy of itself behind.
    file_url = await _store_original(req.content, file, slug, batch_id, current_user)
    for o in objs:
        o.file_url = file_url
    await _write_lines(db, slug, current_user, objs, batch_id=batch_id,
                       source_file=file.filename, file_url=file_url, uploaded_at=uploaded_at,
                       period=period, decision=decision, date_formats=formats)
    if supplier is not None:
        db.add(_supplier_link(current_user.tenant_id, slug, batch_id, supplier))
    control = _control_row(
        slug, current_user, batch_id, content=req.content, filename=file.filename or "",
        header_row=req.header_row, file_rows=int(len(req.df)), objs=objs,
        expected_count=expected_count, expected_amount=expected_amt, decision=decision)
    if control is not None:
        db.add(control)
    await _project(db, slug, current_user, batch_id)
    await db.commit()
    rowfilter = spec.row_filter(slug)
    return {
        "batch_id": batch_id,
        "type": slug,
        "file_name": file.filename,
        "inserted": len(objs),
        # The statement's own balance lines among `inserted` — summaries, not entries.
        "balance_rows": sum(1 for o in objs
                            if (o.data or {}).get(statement_balance.ROW_KIND_KEY)),
        "controls": _control_fields(control, objs),
        "matched_columns": len(req.colmap),
        "source_rows": lines.source_rows,
        "edited_rows": lines.edited_rows,
        # Rows the type discards on content (NDC's unpaid holds and free seats). Always
        # present so the wizard can account for the shortfall between the file's line count
        # and the entry count on the uploads list; 0 for every other type.
        "excluded_rows": lines.excluded_rows,
        # Null unless the rule could not run — TXN Type was left unmapped, so nothing was
        # classified and nothing was skipped. The wizard says so rather than letting the
        # user believe the filter applied.
        "filter_column_mapped": (
            bool(req.colmap.get(rowfilter["field"])) if rowfilter else None
        ),
        "leg_rows": sum(1 for o in objs if not getattr(o, "is_total", False)),
        "airline_resolved_rows": lines.airline_rows,
        **_period_fields(period),
        # The date formats the dates were read in, as saved with the upload. Null for the
        # types that declare none.
        "date_formats": formats,
        # Third Party GDS: each bucket of duplicates — found, saved, left out — and the
        # uploads the repeated lines are in. Null for every other type.
        "duplicates": screen.summary(decision) if screen is not None else None,
        # Third Party GDS: the saved tickets dated outside the period — reviewed and kept on
        # the review step, said once more on the done screen. Null for every other type.
        "outside_period": (len(tp_gds_reupload.dates_outside(objs, period).lines)
                           if screen is not None and period is not None else None),
        **_supplier_fields(supplier),
    }


@router.post("/{slug}/save-preview")
async def preview_save(
    slug: str,
    file: UploadFile = File(...),
    column_map: str = Form(...),
    supplier_id: int | None = Form(default=None),
    agency_id: int | None = Form(default=None),
    header_row: int = Form(...),
    file_digest: str = Form(...),
    edits: str = Form(default="{}"),
    period_from: str | None = Form(default=None),
    period_to: str | None = Form(default=None),
    date_formats: str | None = Form(default=None),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """What `/confirm` would do with this file — nothing is written.

    The review step shows it before Save is pressed: each entry's bucket against the lines
    already saved for this agency (`rows` — already uploaded, partial duplicate with what
    changed, repeated in this file; the rest is new), the uploads the repeats are saved in,
    and the defaults Save applies — the user's choices are applied by the wizard to these
    marks, and by `/confirm` on Save. Also any Week value in the file that falls outside the
    declared period, and the tickets whose Booked Date or Issue Date does (`outside_period` —
    every entry, whatever its bucket: the wizard counts the ones that will be saved). And how
    each declared date column read in the format chosen for it (`dates` — see
    `_dates_read`). Read through the very code `/confirm` uses, edits applied, so the two
    cannot disagree; `/confirm` still decides again at save time.

    `{"applies": false}` for every type that does not capture a period.
    """
    slug, model = _resolve(slug)
    if slug != tp_gds_reupload.SLUG:
        return {"applies": False}
    builder = _mapping_capable(slug)
    period = _period_form(slug, period_from, period_to)
    formats = _date_formats_form(slug, date_formats, required=True)
    supplier = await _supplier_selection(db, slug, supplier_id, agency_id, current_user)
    req = await _read_mapping_request(slug, builder, file, column_map, header_row, file_digest, edits,
                                      period, formats)
    # Never added to the session: these lines exist only to be compared.
    prov = dict(tenant_id=current_user.tenant_id, created_by_id=current_user.id,
                batch_id="preview", source_file=file.filename, file_url=None,
                uploaded_at=datetime.utcnow())
    lines = await _mapped_lines(db, slug, model, builder, req, prov)
    screen = await tp_gds_reupload.compare(db, current_user, supplier, period, lines.objs,
                                           key_data=lines.key_data, indexes=lines.indexes)
    outside = tp_gds_reupload.dates_outside(lines.objs, period, indexes=lines.indexes)
    return {"applies": True, **screen.preview(),
            "week_outside": tp_gds_reupload.weeks_outside(lines.objs, period, date.today()),
            "outside_period": outside.as_dict(),
            "dates": _dates_read(slug, req, lines),
            **_row_marks(lines)}


@router.post("/{slug}/upload")
async def upload_statement(
    slug: str,
    file: UploadFile = File(...),
    tenant_airline_ids: list[int] = Form(default=[]),
    supplier_id: int | None = Form(default=None),
    # The Agency Master row picked on upload. Preferred over supplier_id: it is
    # channel-checked, works for agencies awaiting approval, and carries the vendor
    # service charge the SF columns are checked against.
    agency_id: int | None = Form(default=None),
    expected_count: int | None = Form(default=None),
    expected_amount: str | None = Form(default=None),
    # The statement period — required for the types that capture one, as at `/confirm`.
    period_from: str | None = Form(default=None),
    period_to: str | None = Form(default=None),
    # The declared date formats — required for the types that declare them, as at `/confirm`.
    date_formats: str | None = Form(default=None),
    # What to do with the duplicates found (Third Party GDS), as at `/confirm`; the defaults
    # when absent — exact repeats left out, partial and in-file ones saved and flagged.
    duplicate_choices: str | None = Form(default=None),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Parse a statement export → verbatim fixed columns + folded taxes, under one batch.

    For the LCC types the uploader must also declare which of their Airline Master ids
    the file covers — the export names no carrier, so nothing else identifies it. For the
    third-party types they must declare the consolidator from the Supplier master, for the
    same reason: the file does not name its sender, and the B2B deal is matched against it.

    Both are resolved BEFORE the file is read, stored or parsed, so a rejected selection
    cannot leave a half-imported batch or an orphaned file behind.
    """
    slug, model = _resolve(slug)
    period = _period_form(slug, period_from, period_to)
    formats = _date_formats_form(slug, date_formats, required=True)
    choices = _duplicate_choices_form(slug, duplicate_choices)
    airlines = await _airline_selection(db, slug, current_user, tenant_airline_ids)
    supplier = await _supplier_selection(db, slug, supplier_id, agency_id, current_user)

    content = await file.read()
    max_bytes = settings.MAX_UPLOAD_SIZE_MB * 1024 * 1024
    if len(content) > max_bytes:
        raise HTTPException(status_code=413, detail=f"File exceeds {settings.MAX_UPLOAD_SIZE_MB} MB limit.")
    if not content:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")

    parser_name = spec.parser(slug)
    batch_id = str(uuid.uuid4())
    uploaded_at = datetime.utcnow()
    # Stored once the lines are known to go ahead (below) — a file refused as nothing but
    # duplicates must not leave an orphaned copy of itself behind.
    file_url = None

    objs = []
    source_rows = 0
    airline_rows = 0
    excluded_rows = 0   # rows the type discards on content; only the verbatim path sets it
    stamp_airline = await _airline_stamper(db, slug)
    if parser_name:
        # Multi-format normalizer (LCC / DI): alias-map → canonical fields (+ fold taxes/segments/ssr for LCC).
        builder = _get_builder(parser_name)
        df = _detect_df(content, file.filename or "", lambda cols: len(builder.build_col_map(cols)))
        cols = [str(c) for c in df.columns]
        colmap = builder.build_col_map(cols)
        matched = len(colmap)
        _read_stored_dates(slug, content, file.filename or "", df, period,
                           {field: col for col, field in colmap.items()}, formats)
        dated = {"date_formats": formats} if formats else {}
        for _, row in df.iterrows():
            b = builder.build_row(row, cols, **dated)
            if not (b["data"] or b["taxes"] or b["segments"] or b["ssr"]):
                continue  # blank row
            source_rows += 1
            if stamp_airline is not None and stamp_airline(b["data"]):
                airline_rows += 1
            objs.append(_line(
                model, tenant_id=current_user.tenant_id, created_by_id=current_user.id,
                batch_id=batch_id, source_file=file.filename, file_url=file_url, uploaded_at=uploaded_at,
                data=b["data"], taxes=b["taxes"], segments=b["segments"],
                ssr=b["ssr"], raw_data=b["raw_data"], source_format=b["source_format"],
            ))
    else:
        # Verbatim spec-driven path (TGQ HMPR / NDC): fixed columns + Tax_TypeN/TaxN folding.
        df, colmap, normmap = _parse(content, file.filename or "", slug)
        matched = len(colmap)
        fold = spec.fold_taxes(slug)
        prov = dict(
            tenant_id=current_user.tenant_id, created_by_id=current_user.id,
            batch_id=batch_id, source_file=file.filename, file_url=file_url,
            uploaded_at=uploaded_at,
        )
        drop, derive = spec.drop_row(slug), spec.derive_row(slug)
        normalize, stamp_fmt = spec.normalize_row(slug), spec.stamp_format(slug)
        detect = spec.detect_format(slug)
        cols = [str(c) for c in df.columns]
        src_fmt = detect(cols) if detect else None
        for seq, (_, row) in enumerate(df.iterrows(), start=1):
            data = {field: _clean(row[col]) for col, field in colmap.items()}
            if derive is not None:
                # Same precedence as `/confirm`: derived first, the file's own column wins.
                extra = derive({n: _clean(row[c]) for c, n in normmap.items()})
                if extra:
                    data = {**extra, **data}
            # Same rule and same order as `/confirm`, so a direct upload and a mapped one
            # store identical values rather than agreeing only on which rows exist.
            if normalize is not None:
                if stamp_fmt is not None:
                    stamp_fmt(data, src_fmt)
                normalize(data)
            taxes = _fold_taxes(row, normmap) if fold else []
            if not any(v is not None for v in data.values()) and not taxes:
                continue  # blank row
            # Same rule as `/confirm`, so a direct upload and a mapped one agree on what
            # belongs in the table. NDC only reaches this path if something bypasses the
            # wizard, but a row that is not a transaction is not a transaction either way.
            if drop is not None and drop(data):
                excluded_rows += 1
                continue
            source_rows += 1
            objs.extend(_build_rows(
                model, slug, prov, data, taxes, seq,
                extra=_fmt_extra(model, detect, src_fmt, row, cols),
            ))

    if not objs:
        rowfilter = spec.row_filter(slug)
        if excluded_rows and rowfilter:
            raise HTTPException(
                status_code=400,
                detail=f"Every row in this file is {rowfilter['label']} "
                       f"({', '.join(rowfilter['exclude'])}), so there is nothing to import.")
        raise HTTPException(status_code=400, detail="No data rows were found in the file.")

    objs, screen, decision = await _screen_duplicates(db, slug, current_user, supplier, period, objs,
                                                      choices=choices)
    file_url = await _store_original(content, file, slug, batch_id, current_user)
    for o in objs:
        o.file_url = file_url
    await _write_lines(db, slug, current_user, objs, batch_id=batch_id,
                       source_file=file.filename, file_url=file_url, uploaded_at=uploaded_at,
                       period=period, decision=decision, date_formats=formats)
    for ta in airlines:
        db.add(StatementBatchAirlineId(
            tenant_id=current_user.tenant_id, slug=slug,
            batch_id=batch_id, tenant_airline_id=ta.id,
        ))
    if supplier is not None:
        db.add(_supplier_link(current_user.tenant_id, slug, batch_id, supplier))
    control = _control_row(
        slug, current_user, batch_id, content=content, filename=file.filename or "",
        header_row=df.attrs.get("header_row") if parser_name else None,
        file_rows=int(len(df)), objs=objs, expected_count=expected_count,
        expected_amount=_to_decimal_form(expected_amount, "Expected net amount"),
        decision=decision)
    if control is not None:
        db.add(control)
    await _project(db, slug, current_user, batch_id)
    await db.commit()
    return {
        "batch_id": batch_id,
        "type": slug,
        "file_name": file.filename,
        "inserted": len(objs),
        "balance_rows": sum(1 for o in objs
                            if (o.data or {}).get(statement_balance.ROW_KIND_KEY)),
        "controls": _control_fields(control, objs),
        "matched_columns": matched,
        "source_rows": source_rows,
        "excluded_rows": excluded_rows,
        "leg_rows": sum(1 for o in objs if not getattr(o, "is_total", False)),
        "airline_name": airlines[0].airline_name if airlines else None,
        "airline_code": airlines[0].airline_code if airlines else None,
        "airline_ref_ids": [ta.ref_id for ta in airlines],
        # How many rows found their carrier in the airline master. Only meaningful for the
        # types that resolve one; None elsewhere so the shared view shows nothing. A low
        # count is worth saying out loud — an unresolved row matches no deal and is absent
        # from PLB accrual, and neither failure is visible in the rows themselves.
        "airline_resolved_rows": airline_rows if stamp_airline is not None else None,
        **_period_fields(period),
        "duplicates": screen.summary(decision) if screen is not None else None,
        **_supplier_fields(supplier),
    }


async def _annotate_duplicates(db: AsyncSession, slug: str, user: User, rows, out_rows: list[dict]) -> None:
    """Third Party GDS: `__dup__` on each served line saved although it repeats one — its kind,
    what differs, and the upload the line it repeats is in — and `__outside__` on each line
    dated outside its statement period ({field: "before" | "after"}). Nothing elsewhere."""
    if slug != tp_gds_reupload.SLUG or not rows:
        return
    _slug, model = _resolve(slug)
    H = ThirdPartyStatement
    of_ids = {r.duplicate_of_id for r in rows if r.duplicate_of_id}
    where: dict[int, tuple] = {}
    if of_ids:
        for line_id, source_file, p_from, p_to in (await db.execute(
            select(model.id, H.source_file, H.period_from, H.period_to)
            .join(H, H.batch_id == model.batch_id)
            .where(model.id.in_(of_ids), *_scope(model, user))
        )).all():
            where[line_id] = (source_file, statement_period.label(p_from, p_to) if p_from else None)
    labels = dict(tp_gds_reupload.DIFF_FIELDS)
    for row, out in zip(rows, out_rows):
        if row.duplicate_kind:
            of_file, of_period = where.get(row.duplicate_of_id, (None, None))
            out["__dup__"] = {
                "kind": row.duplicate_kind,
                "fields": [labels.get(f, f) for f in (row.duplicate_fields or "").split(",") if f],
                "of_file": of_file, "of_period": of_period,
            }
        if row.period_from and row.period_to:
            sides = tp_gds_reupload.outside_sides(row.data or {}, (row.period_from, row.period_to))
            sides = {k: v for k, v in sides.items() if v != "unreadable"}
            if sides:
                out["__outside__"] = sides


async def _annotate_service_fees(db: AsyncSession, slug: str, user: User, rows, out_rows: list[dict],
                                 batch_id: str | None) -> dict:
    """Adds `__kind__` / `__checks__` to each served row (services/statement_sf_check.py)
    and returns the batch-level summary — over the WHOLE batch when one is open, not just
    this page, so "3 rows overcharged" means the statement, not the screen."""
    from app.models.agency import Agency
    from app.services import statement_sf_check as sfc

    batch_ids = {r.batch_id for r in rows}
    links = {l.batch_id: l for l in (await db.execute(
        select(StatementBatchSupplier).where(
            StatementBatchSupplier.slug == slug,
            StatementBatchSupplier.tenant_id == user.tenant_id,
            StatementBatchSupplier.batch_id.in_(batch_ids),
        )
    )).scalars().all()}
    agency_ids = {l.agency_id for l in links.values() if l.agency_id}
    agencies = {a.id: a for a in (await db.execute(
        select(Agency).where(Agency.id.in_(agency_ids), Agency.user_id == user.id)
    )).scalars().all()} if agency_ids else {}
    agreements = {bid: sfc.agreement_for(agencies.get(l.agency_id)) for bid, l in links.items()}

    for row, out in zip(rows, out_rows):
        out.update(sfc.check_row(row.data or {}, agreements.get(row.batch_id)))

    summary_rows = out_rows
    if batch_id:
        _slug, model = _resolve(slug)
        all_rows = (await db.execute(select(model.data).where(
            model.batch_id == batch_id, *_scope(model, user)))).scalars().all()
        summary_rows = [sfc.check_row(d or {}, agreements.get(batch_id)) for d in all_rows]

    ag = agreements.get(batch_id) if batch_id else None
    link = links.get(batch_id) if batch_id else None
    return {
        **sfc.summarize(summary_rows),
        "agreement": ag.label if ag else None,
        "agency_linked": bool(link and link.agency_id),
    }


@router.get("/{slug}/records")
async def list_records(
    slug: str,
    request: Request,
    limit: int = Query(100, ge=1, le=500),
    offset: int = Query(0, ge=0),
    batch_id: str | None = Query(None),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Paginated rows — fixed columns flattened from `data`, plus a folded `Taxes` column.

    Types that declare them additionally get `f.<field>` filters, a `summary` block totalled
    over the whole filtered set, and a `Leg` column. Every one of those keys is absent for a
    type that declares none, so the shared frontend view is unchanged for it.
    """
    slug, model = _resolve(slug)
    conds = _record_conds(model, slug, current_user, batch_id, request)
    active = _active_filters(slug, request)

    total = await db.scalar(select(func.count()).select_from(select(model).where(*conds).subquery()))

    # File order, not insertion order reversed: `id DESC` used to surface the export's last
    # line (its grand total) first. Sort on row_seq so re-processed legs stay under their
    # parent rather than landing at the end of the batch.
    order = [model.uploaded_at.desc()]
    if _splits(model):
        order += [func.coalesce(model.row_seq, model.id).asc(),
                  func.coalesce(model.sector_index, 1).asc()]
    order.append(model.id.asc())
    q = select(model).where(*conds).order_by(*order).limit(limit).offset(offset)
    rows = (await db.execute(q)).scalars().all()

    split_cfg = spec.split_config(slug)
    cols = _display_columns(slug)
    out_rows = [flatten_record(slug, r) for r in rows]

    payload = {"total": total or 0, "limit": limit, "offset": offset, "columns": cols, "rows": out_rows}

    # Third-party statements: each row's kind (issue / refund) and, where the agency has a
    # vendor service charge on record, whether its SF and GST On SF agree with it.
    if slug in _SLUG_CHANNEL and rows:
        payload["sf_check"] = await _annotate_service_fees(db, slug, current_user, rows, out_rows, batch_id)
    await _annotate_duplicates(db, slug, current_user, rows, out_rows)

    filters = spec.filter_specs(slug)
    if filters:
        payload["filters"] = filters

    sum_specs = spec.summary_fields(slug)
    if sum_specs:
        payload["summary"] = await _summary(
            db, model, slug, current_user, batch_id, conds, sum_specs, bool(active)
        )

    # Gated on the spec, not just the table: `ndc` carries the leg columns but has
    # splitting switched off, and offering to re-process it would be a prompt to nowhere.
    if _splits(model) and split_cfg and batch_id:
        # Batch-scoped on purpose: whether this upload predates splitting is a property of
        # the upload, not of whatever filter happens to be applied.
        payload["needs_reprocess"] = bool(await db.scalar(
            select(func.count()).select_from(model).where(
                model.batch_id == batch_id, *_scope(model, current_user),
                model.sector_count.is_(None),
            )
        ))
    # Was this batch parsed by an OLDER alias map than the one running now? `source_format`
    # is the schema version, so a mismatch means fields the current parser extracts are
    # simply absent from `data` — which renders as blank cells rather than as an error.
    # Batch-scoped for the same reason as above.
    current_fmt = _flat.CURRENT_SOURCE_FORMATS.get(spec.parser(slug) or slug)
    # Only where a Reprocess could act on it: a typed table keeps no raw row to re-read.
    if current_fmt and batch_id and hasattr(model, "raw_data"):
        stale = await db.scalar(
            select(func.count()).select_from(model).where(
                model.batch_id == batch_id, *_scope(model, current_user),
                func.coalesce(model.source_format, "") != current_fmt,
            )
        )
        if stale:
            payload["needs_reprocess"] = True
            payload["reprocess_reason"] = "stale_format"
            payload["stale_rows"] = stale
    return payload


async def _summary(db, model, slug, user, batch_id, conds, sum_specs, has_filters) -> dict:
    """Column totals over the ENTIRE filtered set, plus the file's own declared total.

    Both are shown side by side rather than picking one: real exports are not always
    internally consistent, and quietly presenting a single number would hide that.
    """
    fields = [f["field"] for f in sum_specs]
    agg = (await db.execute(
        select(func.count().label("n"), *[func.sum(_num(model, f)).label(f"s_{f}") for f in fields])
        .where(*conds)
    )).one()

    computed = {f: _money_str(getattr(agg, f"s_{f}")) or "0" for f in fields}

    declared: dict | None = None
    note: str | None = None
    if _splits(model) and batch_id:
        totals = (await db.execute(
            select(model).where(model.is_total.is_(True), model.batch_id == batch_id, *_scope(model, user))
            .limit(2)
        )).scalars().all()
        if len(totals) == 1:
            declared = {f: (totals[0].data or {}).get(f) for f in fields}
        elif len(totals) > 1:
            note = "This upload has more than one Total line — showing computed figures only."

    # A running-account statement declares its period net through its own balances:
    # opening + Σ Net Amount = closing. Show the file's figure against the computed one,
    # the same ✓/≠ the TGQ total line gets.
    balance = None
    if spec.balance_lines(slug) and batch_id:
        balance = await _statement_balance(db, model, slug, user, batch_id)
        net = statement_balance.derived_ticket_net(
            _opt_dec(balance.get("opening")), _opt_dec(balance.get("closing")),
            _opt_dec(balance.get("payments")))
        if net is not None and "net_amount" in fields:
            declared = {**(declared or {}), "net_amount": _money_str(net)}

    leg_count = agg.n or 0
    source_rows = leg_count
    if _splits(model):
        source_rows = await db.scalar(
            select(func.count(func.distinct(func.coalesce(model.row_seq, model.id)))).where(*conds)
        ) or 0

    return {
        "fields": sum_specs,
        "computed": computed,
        "declared": declared,
        # A filtered subset is not comparable to a whole-file total; saying so beats
        # flashing a delta that only means "you filtered something out".
        "declared_comparable": bool(declared) and not has_filters,
        "declared_note": note,
        "row_count": source_rows,
        "leg_count": leg_count,
        # {opening, closing, payments, opening_source} for statements that carry their
        # own balances.
        "balance": balance,
    }


def _opt_dec(v) -> Decimal | None:
    return statement_balance.to_decimal(v)


async def _statement_balance(db, model, slug, user, batch_id: str) -> dict:
    """The upload's opening, closing and in-period payments, as strings at file precision.

    The closing comes from the stamped BALANCE row and the payments from its LESS PAYMENT
    rows; the opening from the OLD row when the file kept one as a row, else from the
    control figures captured off the preamble at upload (or typed in later). Any may be None.
    """
    rows = (await db.execute(
        select(model.data).where(
            model.batch_id == batch_id, *_scope(model, user),
            model.data[statement_balance.ROW_KIND_KEY].astext.isnot(None))
        .order_by(model.id.asc())
    )).scalars().all()
    figures = statement_balance.account_figures(
        ((data or {}).get(statement_balance.ROW_KIND_KEY), (data or {}).get("net_amount"))
        for data in rows)
    out: dict = {
        "opening": str(figures["opening"]) if figures["opening"] is not None else None,
        "closing": str(figures["closing"]) if figures["closing"] is not None else None,
        "payments": str(figures["payments"]) if figures["payments"] is not None else None,
        "opening_source": "file" if figures["opening"] is not None else None,
    }
    ctl = (await db.execute(
        select(StatementBatchControl).where(
            StatementBatchControl.slug == slug, StatementBatchControl.batch_id == batch_id,
            StatementBatchControl.tenant_id == user.tenant_id)
    )).scalar_one_or_none()
    if ctl is not None:
        if out["opening"] is None and ctl.opening_balance is not None:
            out["opening"], out["opening_source"] = str(ctl.opening_balance), ctl.opening_source
        if out["closing"] is None and ctl.closing_balance is not None:
            out["closing"] = str(ctl.closing_balance)
    return out


@router.get("/{slug}/records/facets")
async def record_facets(
    slug: str,
    batch_id: str | None = Query(None),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Distinct values for every `select` filter this type declares — `{}` when it has none.

    Post-split these are genuinely useful: `traveldt` holds one date per row instead of
    "28APR 28APR 14MAY 15MAY", so it collapses to a short, pickable list.
    """
    slug, model = _resolve(slug)
    selects = [f for f in spec.filter_specs(slug) if f.get("type") == "select"]
    if not selects:
        return {}

    conds = [*_scope(model, current_user)]
    if batch_id:
        conds.append(model.batch_id == batch_id)
    if _splits(model):
        conds.append(model.is_total.is_(False))

    out: dict[str, list[str]] = {}
    for f in selects:
        col = model.data[f["field"]].astext
        values = (await db.execute(
            select(func.distinct(col)).where(*conds, col.isnot(None), col != "")
            .order_by(col).limit(_MAX_FACET_VALUES)
        )).scalars().all()
        out[f["field"]] = [v for v in values if v]
    return out


@router.post("/{slug}/batches/{batch_id}/resplit")
async def resplit_batch(
    slug: str,
    batch_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Re-derive this upload's per-sector rows in place — no re-upload needed.

    Covers batches imported before splitting existed, and re-deriving legs if the
    allocation rule changes. Idempotent: `orig_data` holds the pre-split row, so running
    it on already-split rows reproduces exactly the same set. All-or-nothing — the delete
    and the re-insert share one transaction, so a failure can't leave a half-split batch.
    """
    slug, model = _resolve(slug)
    if not _splits(model):
        raise HTTPException(status_code=400, detail=f"{spec.spec_for(slug)['label']} rows are not split by sector.")

    rows = (await db.execute(
        select(model).where(model.batch_id == batch_id, *_scope(model, current_user))
        .order_by(func.coalesce(model.row_seq, model.id).asc(),
                  func.coalesce(model.sector_index, 1).asc(), model.id.asc())
    )).scalars().all()
    if not rows:
        raise HTTPException(status_code=404, detail="Upload not found.")

    # Collapse back to source lines first: sibling legs share a row_seq and carry the same
    # orig_data, so only the first of each group is a distinct source row.
    sources: list[tuple[dict, list, dict]] = []
    seen: set = set()
    for i, r in enumerate(rows, start=1):
        key = r.row_seq if r.row_seq is not None else f"id:{r.id}"
        if key in seen:
            continue
        seen.add(key)
        prov = dict(
            tenant_id=r.tenant_id, created_by_id=r.created_by_id, batch_id=r.batch_id,
            source_file=r.source_file, file_url=r.file_url, uploaded_at=r.uploaded_at,
        )
        sources.append((dict(r.orig_data or r.data or {}), list(r.orig_taxes or r.taxes or []), prov))

    objs = []
    for seq, (data, taxes, prov) in enumerate(sources, start=1):
        objs.extend(_build_rows(model, slug, prov, data, taxes, seq))

    await db.execute(delete(model).where(model.batch_id == batch_id, *_scope(model, current_user)))
    db.add_all(objs)
    # A re-split mints new row ids for the same documents, so the projection has to be
    # rewritten: its orphan sweep is what removes the pre-split rows.
    await _project(db, slug, current_user, batch_id)
    await db.commit()
    return {
        "source_rows": len(sources),
        "leg_rows": sum(1 for o in objs if not o.is_total),
        "total_rows": len(objs),
    }


@router.post("/{slug}/batches/{batch_id}/reprocess")
async def reprocess_batch(
    slug: str,
    batch_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Re-run the current parser over this upload's stored source rows — no re-upload needed.

    `raw_data` holds every original cell keyed by its original header (models/statement_row
    ._NormalizedBase), so a batch ingested under an older alias map can be brought up to the
    current one without the file. That is the whole reason this is possible: the third-party
    GDS schema was written before any real export existed and missed every money column, and
    those batches are recoverable rather than lost.

    Idempotent — re-running reproduces the same set. All-or-nothing: the delete and the
    re-insert share one transaction, so a failure cannot leave a half-parsed batch.

    Deliberately NOT the same endpoint as `resplit`: that one re-derives per-sector legs from
    `orig_data` for the ticket-shaped tables, and re-running it is about an allocation rule,
    not about a schema version. Both exist; neither applies to the other's types.
    """
    slug, model = _resolve(slug)
    parser_name = spec.parser(slug)
    builder = _get_builder(parser_name) if parser_name else None
    if builder is None or not hasattr(builder, "build_row"):
        raise HTTPException(
            status_code=400,
            detail=f"{spec.spec_for(slug)['label']} statements are stored verbatim and have nothing to re-parse.",
        )
    if not hasattr(model, "raw_data"):
        # A typed table keeps each field in its own column and no copy of the vendor's row,
        # so there is nothing here to read again.
        raise HTTPException(
            status_code=400,
            detail=f"{spec.spec_for(slug)['label']} statement lines are stored field by field, "
                   f"not as the original spreadsheet row. Re-upload the file to read it again.",
        )

    rows = (await db.execute(
        select(model).where(model.batch_id == batch_id, *_scope(model, current_user))
        .order_by(model.id.asc())
    )).scalars().all()
    if not rows:
        raise HTTPException(status_code=404, detail="Upload not found.")

    stamp_airline = await _airline_stamper(db, slug)
    objs, skipped, airline_rows = [], 0, 0
    for r in rows:
        raw = dict(r.raw_data or {})
        if not raw:
            # Ingested before raw_data was kept, or a genuinely empty line. Carried through
            # UNCHANGED rather than dropped — re-parsing nothing would delete the row.
            skipped += 1
            objs.append(model(
                tenant_id=r.tenant_id, created_by_id=r.created_by_id, batch_id=r.batch_id,
                source_file=r.source_file, file_url=r.file_url, uploaded_at=r.uploaded_at,
                data=r.data, taxes=r.taxes, segments=r.segments, ssr=r.ssr,
                raw_data=r.raw_data, source_format=r.source_format,
            ))
            continue
        cols = list(raw.keys())
        b = builder.build_row(raw, cols)
        if stamp_airline is not None and stamp_airline(b["data"]):
            airline_rows += 1
        objs.append(model(
            tenant_id=r.tenant_id, created_by_id=r.created_by_id, batch_id=r.batch_id,
            source_file=r.source_file, file_url=r.file_url, uploaded_at=r.uploaded_at,
            data=b["data"], taxes=b["taxes"], segments=b["segments"], ssr=b["ssr"],
            raw_data=b["raw_data"], source_format=b["source_format"],
        ))

    # Before the delete: the new rows get new ids, so every figure keyed on an old one
    # would be orphaned. Cleared, not migrated — the inputs may have changed too.
    cleared = await _clear_commission(db, slug, current_user, batch_id=batch_id)
    await db.execute(delete(model).where(model.batch_id == batch_id, *_scope(model, current_user)))
    db.add_all(objs)
    # Same reason as the re-split: new ids for the same documents, and the commission
    # figures keyed on the old ones were just cleared.
    await _project(db, slug, current_user, batch_id)
    await db.commit()
    return {
        "rows": len(objs),
        "reparsed": len(objs) - skipped,
        "skipped_no_raw": skipped,
        "commission_cleared": cleared,
        "airline_resolved_rows": airline_rows if stamp_airline is not None else None,
        "source_format": _flat.CURRENT_SOURCE_FORMATS.get(spec.parser(slug) or slug),
    }


_FLAGS = ("exact", "partial", "in_file", "outside")


def _flag_counts(model) -> list:
    """Third Party GDS, per upload: its saved lines flagged as repeating a line already saved
    (by kind — services/tp_gds_reupload.py), and its tickets whose Booked Date or Issue Date
    (unless copied from the Booked Date) falls outside its statement period."""
    def outside(col):
        return or_(col < model.period_from, col > model.period_to)
    dated_outside = and_(model.row_kind.is_(None), model.period_from.isnot(None),
                         or_(outside(model.booking_date),
                             and_(model.issue_date_source.is_(None), outside(model.issue_date))))
    return [
        func.count().filter(model.duplicate_kind == "exact").label("flag_exact"),
        func.count().filter(model.duplicate_kind == "partial").label("flag_partial"),
        func.count().filter(model.duplicate_kind == "in_file").label("flag_in_file"),
        func.count().filter(dated_outside).label("flag_outside"),
    ]


@router.get("/{slug}/batches")
async def list_batches(slug: str, db: AsyncSession = Depends(get_db), current_user: User = Depends(get_current_user)):
    """One row per upload — file name, when, how many rows, whether the file is stored."""
    slug, model = _resolve(slug)
    # Match what the drill-in actually shows: the declared grand-total line is a summary,
    # not an entry — and so are a statement's own OLD/BALANCE lines.
    if _splits(model):
        row_count = func.count().filter(model.is_total.is_(False))
    elif spec.balance_lines(slug):
        row_count = func.count().filter(_not_balance_row(model))
    else:
        row_count = func.count()
    if slug in THIRD_PARTY_SLUGS:
        # One header per upload, its entries counted off its lines, its tally and statement
        # period off the header.
        H = ThirdPartyStatement
        flagged = slug == tp_gds_reupload.SLUG
        counts = (
            select(model.batch_id, row_count.label("row_count"), *(_flag_counts(model) if flagged else ()))
            .where(*_scope(model, current_user))
            .group_by(model.batch_id)
            .subquery()
        )
        q = (
            select(
                H.batch_id, H.source_file, H.file_url, H.uploaded_at,
                func.coalesce(counts.c.row_count, 0).label("row_count"),
                H.total_entries, H.duplicate_entries, H.new_entries,
                H.period_from, H.period_to,
                *((func.coalesce(counts.c[f"flag_{f}"], 0).label(f"flag_{f}") for f in _FLAGS)
                  if flagged else ()),
            )
            .outerjoin(counts, counts.c.batch_id == H.batch_id)
            .where(H.slug == slug, *_scope(H, current_user))
            .order_by(H.uploaded_at.desc())
        )
    else:
        q = (
            select(
                model.batch_id,
                model.source_file,
                func.max(model.file_url).label("file_url"),
                func.max(model.uploaded_at).label("uploaded_at"),
                row_count.label("row_count"),
            )
            .where(*_scope(model, current_user))
            .group_by(model.batch_id, model.source_file)
            .order_by(func.max(model.uploaded_at).desc())
        )
    rows = (await db.execute(q)).all()

    # The declared airline ids, for the types that have them. ONE query for the whole
    # list rather than a lookup per batch. Empty for every other type, and for uploads
    # made before the airline was captured — those show "—" rather than a guess.
    links: dict[str, list[TenantAirline]] = {}
    if spec.requires_airline_id(slug) and rows:
        for batch_id, ta in (await db.execute(
            select(StatementBatchAirlineId.batch_id, TenantAirline)
            .join(TenantAirline, TenantAirline.id == StatementBatchAirlineId.tenant_airline_id)
            .where(
                StatementBatchAirlineId.slug == slug,
                StatementBatchAirlineId.tenant_id == current_user.tenant_id,
                StatementBatchAirlineId.batch_id.in_([r.batch_id for r in rows]),
            )
            .order_by(TenantAirline.ref_id)
        )).all():
            links.setdefault(batch_id, []).append(ta)

    # The declared consolidator, for the types that have one. ONE query for the whole list,
    # and read off the snapshot columns so a renamed vendor doesn't rewrite history. Absent
    # for every other type, and for third-party uploads made before the supplier was
    # captured — those show "—" rather than a guess.
    suppliers: dict[str, StatementBatchSupplier] = {}
    if spec.requires_supplier(slug) and rows:
        for link in (await db.execute(
            select(StatementBatchSupplier).where(
                StatementBatchSupplier.slug == slug,
                StatementBatchSupplier.tenant_id == current_user.tenant_id,
                StatementBatchSupplier.batch_id.in_([r.batch_id for r in rows]),
            )
        )).scalars().all():
            suppliers[link.batch_id] = link

    # For an upload that left lines out as duplicates: the earlier upload most of them are
    # saved in (statement_batch_controls.duplicate_of), named by file and period so the list
    # can say "already saved in test1.xlsx (01–07 Aug 2026)". ONE query for the whole list.
    already_in: dict[str, tuple] = {}
    if slug in THIRD_PARTY_SLUGS and spec.captures_controls(slug) and rows:
        C, E = StatementBatchControl, aliased(ThirdPartyStatement)
        for batch_id, earlier_file, earlier_from, earlier_to in (await db.execute(
            select(C.batch_id, E.source_file, E.period_from, E.period_to)
            .join(E, E.batch_id == C.duplicate_of)
            .where(C.slug == slug, C.tenant_id == current_user.tenant_id,
                   C.batch_id.in_([r.batch_id for r in rows]))
        )).all():
            already_in[batch_id] = (earlier_file, statement_period.label(earlier_from, earlier_to)
                                    if earlier_from else None)

    # Where each upload stands on its way into billing, for the types that have a billing
    # flow — NDC only today. ONE query for the whole list, exactly like the two blocks
    # above, and absent for every other type: a batch with no header row has never had its
    # Billing screen opened, and the frontend reads that as "Set up billing".
    billing: dict[str, StatementBatchBilling] = {}
    if spec.supports_billing(slug) and rows:
        for link in (await db.execute(
            select(StatementBatchBilling).where(
                StatementBatchBilling.slug == slug,
                StatementBatchBilling.tenant_id == current_user.tenant_id,
                StatementBatchBilling.created_by_id == current_user.id,
                StatementBatchBilling.batch_id.in_([r.batch_id for r in rows]),
            )
        )).scalars().all():
            billing[link.batch_id] = link

    def _billing_fields(link: StatementBatchBilling | None) -> dict | None:
        """A NESTED object, not seven flat keys: eight statement types share this response
        and this frontend view, so the shared `Batch` type grows one optional field rather
        than seven that are meaningless for the other seven."""
        if link is None:
            return {"resolution_status": "none", "billable_rows": 0, "resolved_rows": 0,
                    "unresolved_rows": 0, "projected_rows": 0, "projected_tickets": 0,
                    "unlatched_rows": 0}
        return {
            "resolution_status": link.resolution_status,
            "billable_rows": link.billable_rows,
            "resolved_rows": link.resolved_rows,
            "unresolved_rows": link.unresolved_rows,
            "projected_rows": link.projected_rows,
            "projected_tickets": link.projected_tickets,
            "unlatched_rows": link.unlatched_rows,
        }

    return [
        {
            "batch_id": r.batch_id,
            "source_file": r.source_file,
            "uploaded_at": r.uploaded_at,
            "row_count": r.row_count,
            "has_file": bool(r.file_url),
            "created_by_name": current_user.full_name,
            # Every id in a batch shares one carrier, so the first row's airline is
            # the batch's airline — see services/lcc_airline_selection.py.
            "airline_name": links[r.batch_id][0].airline_name if links.get(r.batch_id) else None,
            "airline_code": links[r.batch_id][0].airline_code if links.get(r.batch_id) else None,
            "airline_ref_ids": [ta.ref_id for ta in links.get(r.batch_id, [])],
            "tenant_airline_ids": [ta.id for ta in links.get(r.batch_id, [])],
            "supplier_id": suppliers[r.batch_id].supplier_id if r.batch_id in suppliers else None,
            "supplier_name": suppliers[r.batch_id].supplier_name if r.batch_id in suppliers else None,
            "supplier_branch": suppliers[r.batch_id].supplier_branch if r.batch_id in suppliers else None,
            "supplier_code": suppliers[r.batch_id].supplier_code if r.batch_id in suppliers else None,
            "agency_id": suppliers[r.batch_id].agency_id if r.batch_id in suppliers else None,
            # None for the types with no billing flow, so the shared view renders no column.
            "billing": _billing_fields(billing.get(r.batch_id)) if spec.supports_billing(slug) else None,
            # The third-party tally and statement period (models/third_party_statement.py):
            # entries in the file = new (saved) + duplicate (already saved for an overlapping
            # period of the agency, left out). Absent for every other type.
            **({"total_entries": r.total_entries, "duplicate_entries": r.duplicate_entries,
                "new_entries": r.new_entries,
                "period_from": r.period_from.isoformat() if r.period_from else None,
                "period_to": r.period_to.isoformat() if r.period_to else None,
                "duplicate_of_file": already_in.get(r.batch_id, (None, None))[0],
                "duplicate_of_period": already_in.get(r.batch_id, (None, None))[1],
                # Third Party GDS: its saved lines flagged as duplicates, by kind, and those
                # dated outside its period. Null for LCC and API.
                "flags": ({f: getattr(r, f"flag_{f}") for f in _FLAGS}
                          if slug == tp_gds_reupload.SLUG else None)}
               if slug in THIRD_PARTY_SLUGS else {}),
        }
        for r in rows
    ]


@router.get("/{slug}/batches/{batch_id}/file-url")
async def get_batch_file_url(
    slug: str,
    batch_id: str,
    inline: bool = Query(True, description="inline (preview) vs attachment (download)"),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Signed GCS URL for the original uploaded file — inline for preview, attachment for download."""
    slug, model = _resolve(slug)
    src, conds = _upload_facts(slug, model, batch_id, current_user)
    row = (await db.execute(
        select(src.file_url, src.source_file).where(*conds).limit(1)
    )).first()
    if not row or not row.file_url:
        raise HTTPException(status_code=404, detail="No file stored for this upload.")
    from app.services import gcs
    url = await gcs.generate_signed_url(row.file_url, _bucket(), expiry_minutes=60, inline=inline)
    return {"url": url, "file_name": row.source_file}


def _billing_export(slug: str):
    """`(billing_state, extra_headers, extras)` for a type with a billing flow, else None.

    The type's OWN `billing_state(row, ticket)` and the per-row columns its worklist shows
    beside it, so the export can never describe a row differently from that screen: NDC's
    Line (which ticket an ancillary's money is on), Third Party API's Category and Kind
    (what it bills under, or why it does not bill).
    """
    if not spec.supports_billing(slug):
        return None
    if slug == "ndc":
        from app.services.ndc_billing_projection import billing_state

        def extras(r) -> dict:
            s = r.bill_latch_status
            return {"Line": statement_export.LINE_LABEL.get(s or "", s)}
        return billing_state, ("Line",), extras
    if slug == "tp-api":
        from app.services.markup_categories import CATEGORY_LABELS
        from app.services.tp_api_billing_projection import billing_state, row_category

        def extras(r) -> dict:
            cat = row_category(r.data)
            return {"Category": CATEGORY_LABELS.get(cat, cat) if cat else "No category",
                    "Kind": statement_export.KIND_LABEL.get(r.bill_kind or "", r.bill_kind)}
        return billing_state, ("Category", "Kind"), extras
    return None


@router.get("/{slug}/batches/{batch_id}/export")
async def export_entries(
    slug: str,
    batch_id: str,
    request: Request,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """The upload's entries AS STORED — one row per entry — as an .xlsx.

    The other half of the check a user runs after an import: `file-url` hands back the file
    they uploaded, this hands back what the database made of it, so the two can be laid side
    by side. Same rows, order and columns as `/records`, and the same `f.<field>` filters, so
    "export what I'm looking at" works; the Summary sheet says when it was filtered. See
    services/statement_export.py.
    """
    slug, model = _resolve(slug)
    src, src_conds = _upload_facts(slug, model, batch_id, current_user)
    first = (await db.execute(
        select(src.source_file, src.uploaded_at).where(*src_conds).limit(1)
    )).first()
    if first is None:
        raise HTTPException(status_code=404, detail="Upload not found.")

    conds = _record_conds(model, slug, current_user, batch_id, request)
    order = []
    if _splits(model):
        order += [func.coalesce(model.row_seq, model.id).asc(),
                  func.coalesce(model.sector_index, 1).asc()]
    order.append(model.id.asc())
    # The verbatim/pre-split copies are the heaviest columns and flatten_record reads
    # neither, so they are never loaded.
    heavy = [defer(getattr(model, c)) for c in ("raw_data", "orig_data", "orig_taxes")
             if hasattr(model, c)]

    billing_export = _billing_export(slug)
    billing = None
    if billing_export is not None:
        link = (await db.execute(
            select(StatementBatchBilling).where(
                StatementBatchBilling.slug == slug, StatementBatchBilling.batch_id == batch_id,
                StatementBatchBilling.tenant_id == current_user.tenant_id,
                StatementBatchBilling.created_by_id == current_user.id)
        )).scalar_one_or_none()
        billing = {} if link is None else {
            "resolution_status": link.resolution_status, "billable_rows": link.billable_rows,
            "resolved_rows": link.resolved_rows, "projected_tickets": link.projected_tickets,
            "unlatched_rows": link.unlatched_rows,
        }
    # Not before billing is set up: billing_state reads an unmatched row as "Needs a party",
    # which on an upload nobody has opened Billing for would put that on every line. The
    # Summary says billing is not set up instead — what the uploads list says too.
    with_billing = (billing_export is not None and bool(billing)
                    and billing.get("resolution_status") != "none")
    extra_headers: tuple[str, ...] = billing_export[1] if with_billing else ()

    if with_billing:
        state_fn, _, extras_fn = billing_export
        # LEFT JOIN, as the worklist does: billing_state needs the ticket a row was
        # projected into to tell "in billing" from "on an invoice" from "re-send".
        T = aliased(UploadedTicket)
        pairs = (await db.execute(
            select(model, T).options(*heavy)
            .outerjoin(T, T.id == model.projected_ticket_id)
            .where(*conds).order_by(*order)
        )).all()
        entries = [statement_export.Entry(
            record=flatten_record(slug, r), billing=state_fn(r, t), extras=extras_fn(r),
        ) for r, t in pairs]
    else:
        rows = (await db.execute(
            select(model).options(*heavy).where(*conds).order_by(*order)
        )).scalars().all()
        entries = [statement_export.Entry(record=flatten_record(slug, r)) for r in rows]

    # A live COUNT with no filters — exactly how the uploads list counts this upload.
    stored = await db.scalar(
        select(func.count()).select_from(model)
        .where(*_record_conds(model, slug, current_user, batch_id, None))
    ) or 0
    control = (await db.execute(
        select(StatementBatchControl).where(
            StatementBatchControl.slug == slug, StatementBatchControl.batch_id == batch_id,
            StatementBatchControl.tenant_id == current_user.tenant_id)
    )).scalar_one_or_none()

    facts = statement_export.UploadFacts(
        type_label=spec.spec_for(slug)["label"],
        source_file=first.source_file,
        stored_rows=stored,
        uploaded_by=current_user.full_name,
        uploaded_at=first.uploaded_at,
        file_rows=control.file_rows if control is not None else None,
        row_filter=spec.row_filter(slug),
        billing=billing,
        filters=[(f["label"], value) for f, value in _active_filters(slug, request).values()],
    )

    fd, path = tempfile.mkstemp(prefix=f"{slug}-entries-", suffix=".xlsx")
    os.close(fd)
    try:
        await asyncio.get_running_loop().run_in_executor(
            None, lambda: statement_export.write_workbook(
                path, _display_columns(slug), entries, facts, spec.summary_fields(slug),
                extra_headers=extra_headers, with_billing=with_billing))
    except BaseException:
        with suppress(OSError):
            os.remove(path)
        raise
    return FileResponse(
        path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=statement_export.file_name(first.source_file),
        background=BackgroundTask(os.remove, path),
    )


class RecordEdit(BaseModel):
    """{field: new value} for the cells being changed. Only display columns are editable."""
    fields: dict[str, str | None]


# Row editing is offered where a person reconciles the file line by line against a
# consolidator: the third-party statements. Other types are system exports (BSP, airline
# LCC) whose rows should be fixed by re-uploading, not hand-edited.
_EDITABLE_SLUGS = {"tp-gds", "tp-lcc"}


@router.patch("/{slug}/records/{record_id}")
async def edit_record(
    slug: str,
    record_id: int,
    payload: RecordEdit,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Correct cells of one statement row.

    `data` changes; `raw_data` — the line as the vendor sent it, on the tables that keep
    one — never does, so what the consolidator actually printed stays recoverable. Third
    Party GDS keeps no such copy; its original is the uploaded file. Each edited field is listed in
    `data["__edited__"]` so the view can mark it.

    After the edit the row is re-derived exactly as at upload (bare numbers, ISO dates, the
    joined passenger name, a derived total fare), and its commission figure is CLEARED:
    a figure priced off the old values must not keep being summed. Re-running commission
    on the statement prices it again.
    """
    from app.services import flat_statement

    slug, model = _resolve(slug)
    if slug not in _EDITABLE_SLUGS:
        raise HTTPException(status_code=400, detail="Rows of this statement type can't be edited — re-upload the file instead.")
    obj = (await db.execute(
        select(model).where(model.id == record_id, *_scope(model, current_user))
    )).scalar_one_or_none()
    if not obj:
        raise HTTPException(status_code=404, detail="Record not found.")

    editable = {c["field"] for c in spec.columns(slug) if not str(c["field"]).startswith("__")}
    unknown = sorted(set(payload.fields) - editable)
    if unknown:
        raise HTTPException(status_code=400, detail=f"Not an editable column: {', '.join(unknown)}.")

    data = dict(obj.data or {})
    old_joined = " ".join(p for p in (data.get("first_name"), data.get("last_name")) if p)
    edited = {f for f in str(data.get("__edited__") or "").split(",") if f}
    for field, value in payload.fields.items():
        value = (value or "").strip()
        if value:
            data[field] = value
        else:
            data.pop(field, None)
        edited.add(field)

    # Values the upload DERIVED must follow the edit, or they keep describing the old row:
    # the passenger name when it was only First + Last joined, and a computed total fare.
    if ({"first_name", "last_name"} & payload.fields.keys()
            and data.get("passenger_name") in (old_joined, None) and "passenger_name" not in payload.fields):
        data.pop("passenger_name", None)
    if data.get("total_fare_source") == "derived" and "total_fare" not in payload.fields:
        data.pop("total_fare", None)
        data.pop("total_fare_source", None)
    builder = flat_statement.get(spec.parser(slug) or "")
    derive = getattr(builder, "_derive", None) if builder else None
    if derive:
        derive(data)

    # A comma string, not a list: every other value in `data` is text, and readers
    # (exports, facets) treat it that way.
    data["__edited__"] = ",".join(sorted(edited))
    obj.data = data
    # A typed table stores an amount as a number and a date as a date. A cell edited to
    # something that is neither would save as a blank, so it is refused instead — before
    # anything is committed.
    if is_typed(model):
        misfit = set(payload.fields) & set(str(obj.data.get(UNPARSED_KEY) or "").split(","))
        if misfit:
            headers = {c["field"]: c["header"] for c in spec.columns(slug)}
            names = ", ".join(headers.get(f, f) for f in sorted(misfit))
            raise HTTPException(
                status_code=400,
                detail=f"{names}: enter a number for an amount and a date for a date.")
    await _clear_commission(db, slug, current_user, row_ids=[obj.id])
    await db.commit()
    await db.refresh(obj)

    out = flatten_record(slug, obj)
    if slug in _SLUG_CHANNEL:
        await _annotate_service_fees(db, slug, current_user, [obj], [out], None)
    await _annotate_duplicates(db, slug, current_user, [obj], [out])
    return out


async def _unflag_repeats(db: AsyncSession, slug: str, model, user: User, gone) -> None:
    """Third Party GDS: lines saved as repeating one of the lines `gone` selects (their ids)
    are duplicates of nothing once those go — their flag goes too. Run BEFORE the delete:
    the foreign key would only null the pointer and leave the flag claiming a duplicate."""
    if slug != tp_gds_reupload.SLUG:
        return
    await db.execute(
        update(model)
        .where(model.duplicate_of_id.in_(gone), *_scope(model, user))
        .values(duplicate_kind=None, duplicate_of_id=None, duplicate_fields=None)
        .execution_options(synchronize_session=False))


@router.delete("/{slug}/records/{record_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_record(slug: str, record_id: int, db: AsyncSession = Depends(get_db), current_user: User = Depends(get_current_user)):
    """Delete a record — and, for a split ticket, all of its sibling legs.

    A leg is a slice of one ticket, not a standalone entry: dropping one would leave a
    part-ticket behind and silently skew every column total.
    """
    slug, model = _resolve(slug)
    obj = (await db.execute(
        select(model).where(model.id == record_id, *_scope(model, current_user))
    )).scalar_one_or_none()
    if not obj:
        raise HTTPException(status_code=404, detail="Record not found.")
    # Read before the row goes: a balance line is not an entry, so it is not in the tally.
    is_entry = not (obj.data or {}).get(statement_balance.ROW_KIND_KEY)

    if _splits(model) and (obj.sector_count or 1) > 1 and obj.row_seq is not None:
        sibling_ids = list((await db.execute(
            select(model.id).where(model.batch_id == obj.batch_id, model.row_seq == obj.row_seq,
                                   *_scope(model, current_user))
        )).scalars().all())
        await _clear_commission(db, slug, current_user, row_ids=sibling_ids)
        await db.execute(delete(model).where(
            model.batch_id == obj.batch_id, model.row_seq == obj.row_seq, *_scope(model, current_user)
        ))
    else:
        # Its commission figure goes with it — nothing else points at this row id, and a
        # figure describing a row that no longer exists would still be summed.
        await _clear_commission(db, slug, current_user, row_ids=[obj.id])
        await _unflag_repeats(db, slug, model, current_user, [obj.id])
        # And its projected ticket, on the two types that have one. `projected_ticket_id` is
        # ON DELETE SET NULL, so without this the ticket survives in Sold Tickets with no
        # statement row behind it and no way to reach it — the same orphan `delete_batch`
        # below already prevents for a whole upload. Refused once invoiced, for that
        # endpoint's reason: an invoice may not be left quoting a ticket whose provenance is
        # gone.
        ticket_id = getattr(obj, "projected_ticket_id", None)
        if ticket_id is not None:
            ticket = await db.get(UploadedTicket, ticket_id)
            if ticket is not None:
                if ticket.billing_id is not None:
                    raise HTTPException(
                        status_code=409,
                        detail=("This row is on an invoice. Delete that billing before "
                                "removing the row it was billed from."),
                    )
                await db.delete(ticket)
        await db.delete(obj)
    # A third-party statement's last line takes its header with it, so an emptied upload
    # leaves the uploads list exactly as it did when the lines were the list. Counted
    # UNSCOPED: deleting the header cascades to every line of the batch, whoever's it is.
    if slug in THIRD_PARTY_SLUGS:
        await db.flush()
        left = await db.scalar(
            select(func.count()).select_from(model).where(model.batch_id == obj.batch_id))
        if not left:
            await db.execute(delete(ThirdPartyStatement).where(
                ThirdPartyStatement.batch_id == obj.batch_id,
                *_scope(ThirdPartyStatement, current_user)))
        elif is_entry:
            # One entry fewer on the statement: its tally (models/third_party_statement.py)
            # keeps total = new + duplicate.
            H = ThirdPartyStatement
            await db.execute(update(H).where(H.batch_id == obj.batch_id, *_scope(H, current_user))
                             .values(new_entries=func.greatest(H.new_entries - 1, 0),
                                     total_entries=func.greatest(H.total_entries - 1, 0)))
    # The batch survives, so this is a re-project rather than a clear: the orphan sweep
    # is what removes the row that just went. Without it the deleted line keeps its sale
    # on the revenue board until somebody presses Rebuild.
    await _project(db, slug, current_user, obj.batch_id)
    await db.commit()


@router.delete("/{slug}/batches/{batch_id}")
async def delete_batch(slug: str, batch_id: str, db: AsyncSession = Depends(get_db), current_user: User = Depends(get_current_user)):
    """Delete an entire upload (all rows sharing the batch_id).

    REFUSED once any of the upload's rows is on an invoice. Deleting them would leave the
    invoice quoting a ticket whose provenance no longer exists, which is not a state a
    billed figure may be put into — the same guard as lcc_detailed.py::delete_batch.
    """
    slug, model = _resolve(slug)
    tickets: list[int] = []
    if spec.supports_billing(slug):
        tickets = list((await db.execute(
            select(model.projected_ticket_id).where(
                model.batch_id == batch_id, *_scope(model, current_user),
                model.projected_ticket_id.isnot(None),
            ).distinct()
        )).scalars().all())
        if tickets:
            invoiced = await db.scalar(
                select(func.count()).select_from(UploadedTicket).where(
                    UploadedTicket.id.in_(tickets), UploadedTicket.billing_id.isnot(None))
            ) or 0
            if invoiced:
                raise HTTPException(
                    status_code=409,
                    detail=(f"{invoiced} of this upload's tickets are already on an invoice. "
                            f"Delete those billings first."),
                )

    # Third Party GDS: a later upload that left out lines as already saved in this one keeps
    # holding only what it saved — nothing is moved into it. The uploads list warns before
    # the delete; uploading that statement again brings the rest back (services/
    # tp_gds_reupload.py compares it with what is left). Its lines flagged as repeating one
    # of this upload's are duplicates of nothing once this goes, and lose the flag.
    await _unflag_repeats(db, slug, model, current_user,
                          select(model.id).where(model.batch_id == batch_id, *_scope(model, current_user)))
    res = await db.execute(delete(model).where(model.batch_id == batch_id, *_scope(model, current_user)))
    # The statement itself, for the types that have a header. Its lines went just above, so
    # the CASCADE is the backstop rather than the mechanism, and `res` still counts them.
    header_gone = 0
    if slug in THIRD_PARTY_SLUGS:
        header_gone = (await db.execute(delete(ThirdPartyStatement).where(
            ThirdPartyStatement.batch_id == batch_id, ThirdPartyStatement.slug == slug,
            *_scope(ThirdPartyStatement, current_user)))).rowcount or 0
    # Explicitly: this table's batch_id carries no FK — every spec-driven type shares it,
    # and most of them keep no header row for one to point at. Left behind, these links
    # would keep an airline id looking in-use forever and block its deletion from the
    # Airline Master.
    await db.execute(
        delete(StatementBatchAirlineId).where(
            StatementBatchAirlineId.slug == slug,
            StatementBatchAirlineId.batch_id == batch_id,
            StatementBatchAirlineId.tenant_id == current_user.tenant_id,
        )
    )
    await _clear_commission(db, slug, current_user, batch_id=batch_id)
    # Same story for the consolidator link — no FK on batch_id, so nothing removes it for
    # us, and left behind it would keep the supplier looking in-use.
    await db.execute(
        delete(StatementBatchSupplier).where(
            StatementBatchSupplier.slug == slug,
            StatementBatchSupplier.batch_id == batch_id,
            StatementBatchSupplier.tenant_id == current_user.tenant_id,
        )
    )
    # The control figures and the Payment Module's derived state for this upload — none of
    # it has an FK to rows that no longer exist. Paid payment items are KEPT: a payment that
    # was made does not un-happen because its statement was deleted.
    await db.execute(
        delete(StatementBatchControl).where(
            StatementBatchControl.slug == slug,
            StatementBatchControl.batch_id == batch_id,
            StatementBatchControl.tenant_id == current_user.tenant_id,
        )
    )
    from app.services import payment_ledger
    await payment_ledger.forget_batch(
        db, slug=slug, tenant_id=current_user.tenant_id, user_id=current_user.id,
        batch_id=batch_id)
    # And the billing side: the projected tickets, their statement header, and the billing
    # header itself. The rows are already gone above, so nothing is left pointing at these
    # — and an UploadedTicket left behind would show in Sold Tickets as a billable ticket
    # with no statement, which is worse than absent.
    if spec.supports_billing(slug):
        link = await db.scalar(
            select(StatementBatchBilling).where(
                StatementBatchBilling.slug == slug,
                StatementBatchBilling.batch_id == batch_id,
                StatementBatchBilling.tenant_id == current_user.tenant_id,
                StatementBatchBilling.created_by_id == current_user.id,
            )
        )
        if tickets:
            await db.execute(delete(UploadedTicket).where(UploadedTicket.id.in_(tickets)))
        if link is not None:
            if link.billing_batch_id:
                await db.execute(delete(TicketStatement).where(
                    TicketStatement.batch_id == link.billing_batch_id))
            await db.delete(link)
    # `source_row_id` carries no foreign key, so nothing cascades and the projected rows
    # have to be dropped by hand. Before this call `clear_batch` had no callers at all:
    # a deleted upload left its sale and its income on the board as figures nothing
    # would ever restate.
    await income_board.forget(
        db, tenant_id=current_user.tenant_id, user_id=current_user.id,
        source=slug, batch_id=batch_id,
    )
    await db.commit()
    if not res.rowcount and not header_gone:
        raise HTTPException(status_code=404, detail="Upload not found.")
    return {"deleted": res.rowcount}


@router.get("/{slug}/template")
async def download_template(slug: str, current_user: User = Depends(get_current_user)):
    """Blank .xlsx with the exact headers (fixed columns + Tax_Type1/Tax1 … pairs).

    A file built from this maps itself: every header is the canonical one, so the wizard's
    auto-map fills in and reports a template match. The formatting is not decoration — NDC
    is 69 columns, and a frozen, bold header row is what makes the far end of that legible
    enough to fill in by hand.
    """
    slug, _ = _resolve(slug)
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    headers = spec.template_headers(slug)
    wb = Workbook()
    ws = wb.active
    ws.title = f"{spec.spec_for(slug)['label']} Template"[:31]
    ws.append(headers)

    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1E40AF")
    for i, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = head
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")
        # Wide enough to read the header, capped so one long name can't push the rest of
        # a 69-column sheet off screen.
        ws.column_dimensions[get_column_letter(i)].width = min(max(len(header) + 3, 12), 32)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 20

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{slug}_template.xlsx"'},
    )
